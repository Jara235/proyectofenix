import openpyxl

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)

for sname in wb.sheetnames:
    ws = wb[sname]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val is not None:
                s = str(val).strip()
                if any(k in s for k in ['1288816', '1286993', '1289138', '1290569', '10922', '4184.1', '4184.10']):
                    print(f"Sheet: {sname} | Row {r}, Col {c} | Val: {val}")
