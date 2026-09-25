import openpyxl
import re
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and 'ISOWEEKNUM' in cell.value:
                new_val = re.sub(r'ISOWEEKNUM\(([^)]+)\)', r'WEEKNUM(\1, 2)', cell.value)
                if new_val != cell.value:
                    cell.value = new_val

wb.save('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')
print('Reverted to WEEKNUM(..., 2)')

