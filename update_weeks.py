import openpyxl
import re
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and ('WEEKNUM' in cell.value or 'NUM.DE.SEMANA' in cell.value):
                # replace WEEKNUM(ref, 2) with ISOWEEKNUM(ref)
                new_val = re.sub(r'WEEKNUM\(([^,]+),\s*2\)', r'ISOWEEKNUM(\1)', cell.value)
                new_val = re.sub(r'NUM.DE.SEMANA\(([^,]+),\s*2\)', r'ISOWEEKNUM(\1)', new_val)
                if new_val != cell.value:
                    cell.value = new_val

# Arreglar la columna Semana de Captura_Facturas_Mobil
ws = wb['Captura_Facturas_Mobil']
factura_col = None
fecha_col = None
semana_col = None
for col in range(1, 20):
    val = ws.cell(row=4, column=col).value
    if val == 'No. Factura': factura_col = col
    elif val == 'Fecha': fecha_col = col
    elif val == 'Semana (Helper)': semana_col = col

if fecha_col and semana_col:
    col_letter = openpyxl.utils.get_column_letter(fecha_col)
    for r in range(5, ws.max_row + 1):
        if ws.cell(row=r, column=factura_col).value is not None:
            # Ensure it is a formula
            ws.cell(row=r, column=semana_col).value = f'=IF({col_letter}{r}="","", ISOWEEKNUM({col_letter}{r}))'

wb.save('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')
print('Success!')

