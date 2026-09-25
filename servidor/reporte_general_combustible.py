import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import psycopg2
import io
import os
import datetime as dt_mod

# Paleta corporativa
NAVY = '0F172A'
DARK_SLATE = '1E293B'
SLATE_BLUE = '1E3A5F'
HEADER_BLUE = '2563EB'
LIGHT_BLUE = 'E0F2FE'
ALT_ROW_1 = 'FFFFFF'
ALT_ROW_2 = 'F8FAFC'
ACCENT_GREEN = '059669'
LIGHT_GREEN = 'D1FAE5'
ACCENT_RED = 'DC2626'
LIGHT_RED = 'FEE2E2'
WHITE = 'FFFFFF'
BORDER_COLOR = 'CBD5E1'

def fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type="solid")

def font(bold=False, italic=False, color="000000", size=10):
    return Font(name="Calibri", bold=bold, italic=italic, color=color, size=size)

def align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def thin_border():
    s = Side(style='thin', color=BORDER_COLOR)
    return Border(left=s, right=s, top=s, bottom=s)

def header_border():
    s = Side(style='thin', color='475569')
    return Border(left=s, right=s, top=s, bottom=s)

def generar_libro_maestro_combustible(db_conn=None, ruta_guardado="Reportes/Reporte_General_Combustible.xlsx"):
    """
    Genera el Libro Maestro Consolidado de Combustible (Diésel y Gasolina) - Opción 1.
    Pestaña 1: RESUMEN GENERAL (Matriz histórica semanal S24-S38)
    Pestaña 2: CONSOLIDADO POR OBRA (35 Frentes y Centros de Trabajo)
    Pestañas 3..N: SEMANA XX (Detalle de Diésel y Gasolina de cada semana)
    """
    own_conn = False
    if db_conn is None:
        db_conn = psycopg2.connect(dbname='fenix_db', user='postgres', password='Bupito*268', host='localhost')
        own_conn = True

    cur = db_conn.cursor()

    wb = openpyxl.Workbook()
    # Remover la hoja por defecto al final o renombrarla
    ws_resumen = wb.active
    ws_resumen.title = "RESUMEN GENERAL"
    ws_resumen.views.sheetView[0].showGridLines = True

    # ─────────────────────────────────────────────────────────────────────────────
    # 1. OBTENER DATOS CONSOLIDADOS
    # ─────────────────────────────────────────────────────────────────────────────
    
    # Diésel por semana
    cur.execute("""
        SELECT 
            semana,
            COUNT(*) as cargas,
            COALESCE(SUM(litros), 0) as litros,
            COALESCE(SUM(importe_total), 0) as importe
        FROM diesel.consumos
        WHERE semana IS NOT NULL
        GROUP BY semana
        ORDER BY NULLIF(regexp_replace(semana, '[^0-9]', '', 'g'), '')::integer
    """)
    diesel_semanal = {}
    for r in cur.fetchall():
        digits = ''.join(filter(str.isdigit, str(r[0])))
        if digits:
            diesel_semanal[int(digits)] = {'cargas': int(r[1]), 'litros': float(r[2]), 'importe': float(r[3])}

    # Gasolina por semana
    cur.execute("""
        SELECT 
            semana,
            COUNT(*) as cargas,
            COALESCE(SUM(litros), 0) as litros,
            COALESCE(SUM(importe_total), 0) as importe
        FROM gasolina.consumos
        WHERE semana IS NOT NULL
        GROUP BY semana
        ORDER BY NULLIF(regexp_replace(semana, '[^0-9]', '', 'g'), '')::integer
    """)
    gasolina_semanal = {}
    for r in cur.fetchall():
        digits = ''.join(filter(str.isdigit, str(r[0])))
        if digits:
            gasolina_semanal[int(digits)] = {'cargas': int(r[1]), 'litros': float(r[2]), 'importe': float(r[3])}

    # Autorizaciones de gasolina por semana
    cur.execute("""
        SELECT semana, COALESCE(SUM(importe_semanal), 0) as autorizado
        FROM gasolina.autorizaciones_semanal
        GROUP BY semana
        ORDER BY semana
    """)
    aut_gasolina_semanal = {}
    for r in cur.fetchall():
        aut_gasolina_semanal[int(r[0])] = float(r[1])

    # Lista ordenada de todas las semanas
    all_weeks = sorted(set(list(diesel_semanal.keys()) + list(gasolina_semanal.keys())))

    # Totales acumulados
    tot_global_d_lts = sum(d['litros'] for d in diesel_semanal.values())
    tot_global_d_imp = sum(d['importe'] for d in diesel_semanal.values())
    tot_global_d_cargas = sum(d['cargas'] for d in diesel_semanal.values())

    tot_global_g_lts = sum(g['litros'] for g in gasolina_semanal.values())
    tot_global_g_imp = sum(g['importe'] for g in gasolina_semanal.values())
    tot_global_g_cargas = sum(g['cargas'] for g in gasolina_semanal.values())

    tot_global_comb_lts = tot_global_d_lts + tot_global_g_lts
    tot_global_comb_imp = tot_global_d_imp + tot_global_g_imp
    tot_global_cargas = tot_global_d_cargas + tot_global_g_cargas

    # ─────────────────────────────────────────────────────────────────────────────
    # HOJA 1: RESUMEN GENERAL (DASHBOARD EJECUTIVO)
    # ─────────────────────────────────────────────────────────────────────────────
    # Encabezado Principal
    ws_resumen.merge_cells("A1:K1")
    t1 = ws_resumen.cell(row=1, column=1, value="GRUPO FÉNIX — REPORTE GENERAL CONSOLIDADO DE COMBUSTIBLE (DIÉSEL Y GASOLINA)")
    t1.fill = fill(NAVY); t1.font = font(bold=True, color=WHITE, size=13); t1.alignment = align('center', 'center')
    ws_resumen.row_dimensions[1].height = 28

    ws_resumen.merge_cells("A2:K2")
    t2 = ws_resumen.cell(row=2, column=1, value=f"LIBRO MAESTRO CENTRALIZADO (OPCIÓN 1) | Última actualización: {dt_mod.datetime.now().strftime('%d/%m/%Y %H:%M')} | Dirección de Operaciones y Gobierno Corporativo")
    t2.fill = fill(DARK_SLATE); t2.font = font(italic=True, color='94A3B8', size=9); t2.alignment = align('center', 'center')
    ws_resumen.row_dimensions[2].height = 18

    # Tarjetas KPI (Fila 4 a 6)
    kpis = [
        ("INVERSIÓN TOTAL COMBUSTIBLE", f"${tot_global_comb_imp:,.2f}", "Diésel + Gasolina acumulado", "2563EB", 1, 3),
        ("TOTAL LITROS SUMINISTRADOS", f"{tot_global_comb_lts:,.2f} L", f"{tot_global_cargas:,} cargas totales", "0D9488", 4, 5),
        ("DIÉSEL ACUMULADO", f"${tot_global_d_imp:,.2f}", f"{tot_global_d_lts:,.2f} L ({(tot_global_d_imp/tot_global_comb_imp*100):.1f}% del total)", "4F46E5", 6, 8),
        ("GASOLINA ACUMULADA", f"${tot_global_g_imp:,.2f}", f"{tot_global_g_lts:,.2f} L ({(tot_global_g_imp/tot_global_comb_imp*100):.1f}% del total)", "059669", 9, 11),
    ]

    for title, val, sub, color_h, c_start, c_end in kpis:
        ws_resumen.merge_cells(start_row=4, start_column=c_start, end_row=4, end_column=c_end)
        c_t = ws_resumen.cell(row=4, column=c_start, value=title)
        c_t.fill = fill(color_h); c_t.font = font(bold=True, color=WHITE, size=9); c_t.alignment = align('center', 'center')
        
        ws_resumen.merge_cells(start_row=5, start_column=c_start, end_row=5, end_column=c_end)
        c_v = ws_resumen.cell(row=5, column=c_start, value=val)
        c_v.fill = fill('F1F5F9'); c_v.font = font(bold=True, color='0F172A', size=13); c_v.alignment = align('center', 'center')
        
        ws_resumen.merge_cells(start_row=6, start_column=c_start, end_row=6, end_column=c_end)
        c_s = ws_resumen.cell(row=6, column=c_start, value=sub)
        c_s.fill = fill('F1F5F9'); c_s.font = font(italic=True, color='64748B', size=8); c_s.alignment = align('center', 'center')

        for r_k in range(4, 7):
            for c_k in range(c_start, c_end + 1):
                ws_resumen.cell(row=r_k, column=c_k).border = thin_border()

    ws_resumen.row_dimensions[4].height = 18
    ws_resumen.row_dimensions[5].height = 24
    ws_resumen.row_dimensions[6].height = 16

    # Tabla Semanal (Inicia en fila 8)
    ws_resumen.merge_cells("A8:K8")
    h_tab = ws_resumen.cell(row=8, column=1, value="📅 MATRIZ SEMANAL CONSOLIDADA DE CONSUMO Y COSTOS")
    h_tab.fill = fill(SLATE_BLUE); h_tab.font = font(bold=True, color=WHITE, size=10); h_tab.alignment = align('left', 'center')
    ws_resumen.row_dimensions[8].height = 20

    headers_matriz = [
        "SEMANA", "CARGAS DIÉSEL", "LITROS DIÉSEL", "IMPORTE DIÉSEL ($)",
        "CARGAS GAS.", "LITROS GAS.", "IMPORTE GAS. ($)", "PRESUP. GAS. ($)",
        "SALDO GAS. ($)", "TOTAL LITROS COMB.", "TOTAL INVERSIÓN ($)"
    ]
    ws_resumen.row_dimensions[9].height = 22
    for col_idx, h_text in enumerate(headers_matriz, start=1):
        c = ws_resumen.cell(row=9, column=col_idx, value=h_text)
        c.fill = fill(DARK_SLATE); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center', wrap=True)
        c.border = header_border()

    r_row = 10
    sum_aut_gas = 0.0
    for idx, w in enumerate(all_weeks):
        d = diesel_semanal.get(w, {'cargas': 0, 'litros': 0.0, 'importe': 0.0})
        g = gasolina_semanal.get(w, {'cargas': 0, 'litros': 0.0, 'importe': 0.0})
        aut_g = aut_gasolina_semanal.get(w, 0.0)
        saldo_g = aut_g - g['importe'] if aut_g > 0 else 0.0
        sum_aut_gas += aut_g

        tot_lts_sem = d['litros'] + g['litros']
        tot_imp_sem = d['importe'] + g['importe']

        bg = ALT_ROW_1 if idx % 2 == 0 else ALT_ROW_2

        ws_resumen.cell(row=r_row, column=1, value=f"SEMANA {w}").fill = fill(bg)
        ws_resumen.cell(row=r_row, column=1).font = font(bold=True, size=9); ws_resumen.cell(row=r_row, column=1).alignment = align('center', 'center'); ws_resumen.cell(row=r_row, column=1).border = thin_border()

        c2 = ws_resumen.cell(row=r_row, column=2, value=d['cargas']); c2.fill = fill(bg); c2.font = font(size=9); c2.alignment = align('center', 'center'); c2.border = thin_border(); c2.number_format = '#,##0'
        c3 = ws_resumen.cell(row=r_row, column=3, value=d['litros']); c3.fill = fill(bg); c3.font = font(size=9); c3.alignment = align('right', 'center'); c3.border = thin_border(); c3.number_format = '#,##0.00 "L"'
        c4 = ws_resumen.cell(row=r_row, column=4, value=d['importe']); c4.fill = fill(bg); c4.font = font(bold=True, size=9); c4.alignment = align('right', 'center'); c4.border = thin_border(); c4.number_format = '"$"#,##0.00'

        c5 = ws_resumen.cell(row=r_row, column=5, value=g['cargas']); c5.fill = fill(bg); c5.font = font(size=9); c5.alignment = align('center', 'center'); c5.border = thin_border(); c5.number_format = '#,##0'
        c6 = ws_resumen.cell(row=r_row, column=6, value=g['litros']); c6.fill = fill(bg); c6.font = font(size=9); c6.alignment = align('right', 'center'); c6.border = thin_border(); c6.number_format = '#,##0.00 "L"'
        c7 = ws_resumen.cell(row=r_row, column=7, value=g['importe']); c7.fill = fill(bg); c7.font = font(bold=True, size=9); c7.alignment = align('right', 'center'); c7.border = thin_border(); c7.number_format = '"$"#,##0.00'

        c8 = ws_resumen.cell(row=r_row, column=8, value=aut_g); c8.fill = fill(bg); c8.font = font(size=9); c8.alignment = align('right', 'center'); c8.border = thin_border(); c8.number_format = '"$"#,##0.00'
        
        c9 = ws_resumen.cell(row=r_row, column=9, value=saldo_g)
        c9.fill = fill(LIGHT_RED if saldo_g < 0 else (LIGHT_GREEN if saldo_g > 0 else bg))
        c9.font = font(bold=True, color=ACCENT_RED if saldo_g < 0 else (ACCENT_GREEN if saldo_g > 0 else '000000'), size=9)
        c9.alignment = align('right', 'center'); c9.border = thin_border(); c9.number_format = '"$"#,##0.00;("-""$"#,##0.00);"-"'

        c10 = ws_resumen.cell(row=r_row, column=10, value=tot_lts_sem); c10.fill = fill(bg); c10.font = font(size=9); c10.alignment = align('right', 'center'); c10.border = thin_border(); c10.number_format = '#,##0.00 "L"'
        c11 = ws_resumen.cell(row=r_row, column=11, value=tot_imp_sem); c11.fill = fill(bg); c11.font = font(bold=True, color='1E3A5F', size=9); c11.alignment = align('right', 'center'); c11.border = thin_border(); c11.number_format = '"$"#,##0.00'

        ws_resumen.row_dimensions[r_row].height = 19
        r_row += 1

    # Fila de Totales
    ws_resumen.cell(row=r_row, column=1, value="TOTAL GENERAL").fill = fill(NAVY)
    ws_resumen.cell(row=r_row, column=1).font = font(bold=True, color=WHITE, size=9); ws_resumen.cell(row=r_row, column=1).alignment = align('center', 'center'); ws_resumen.cell(row=r_row, column=1).border = thin_border()

    c = ws_resumen.cell(row=r_row, column=2, value=tot_global_d_cargas); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center'); c.border = thin_border(); c.number_format = '#,##0'
    c = ws_resumen.cell(row=r_row, column=3, value=tot_global_d_lts); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
    c = ws_resumen.cell(row=r_row, column=4, value=tot_global_d_imp); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

    c = ws_resumen.cell(row=r_row, column=5, value=tot_global_g_cargas); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center'); c.border = thin_border(); c.number_format = '#,##0'
    c = ws_resumen.cell(row=r_row, column=6, value=tot_global_g_lts); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
    c = ws_resumen.cell(row=r_row, column=7, value=tot_global_g_imp); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

    c = ws_resumen.cell(row=r_row, column=8, value=sum_aut_gas); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'
    c = ws_resumen.cell(row=r_row, column=9, value=sum_aut_gas - tot_global_g_imp); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00;("-""$"#,##0.00);"-"'

    c = ws_resumen.cell(row=r_row, column=10, value=tot_global_comb_lts); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
    c = ws_resumen.cell(row=r_row, column=11, value=tot_global_comb_imp); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=10); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'
    ws_resumen.row_dimensions[r_row].height = 22

    # Anchos de columna
    col_widths_res = [14, 15, 17, 19, 14, 15, 18, 18, 16, 19, 21]
    for ci, w in enumerate(col_widths_res, start=1):
        ws_resumen.column_dimensions[get_column_letter(ci)].width = w

    # ─────────────────────────────────────────────────────────────────────────────
    # HOJA 2: CONSOLIDADO POR OBRA
    # ─────────────────────────────────────────────────────────────────────────────
    ws_obra = wb.create_sheet(title="CONSOLIDADO POR OBRA")
    ws_obra.views.sheetView[0].showGridLines = True

    ws_obra.merge_cells("A1:J1")
    t_ob1 = ws_obra.cell(row=1, column=1, value="GRUPO FÉNIX — CONSOLIDADO DE COMBUSTIBLE POR CENTRO DE TRABAJO Y OBRA")
    t_ob1.fill = fill(NAVY); t_ob1.font = font(bold=True, color=WHITE, size=12); t_ob1.alignment = align('center', 'center')
    ws_obra.row_dimensions[1].height = 26

    ws_obra.merge_cells("A2:J2")
    t_ob2 = ws_obra.cell(row=2, column=1, value=f"Desglose acumulado por frente de trabajo | {dt_mod.datetime.now().strftime('%d/%m/%Y %H:%M')}")
    t_ob2.fill = fill(DARK_SLATE); t_ob2.font = font(italic=True, color='94A3B8', size=9); t_ob2.alignment = align('center', 'center')
    ws_obra.row_dimensions[2].height = 18

    headers_ob = [
        "CENTRO DE TRABAJO / OBRA", "TIPO", "CARGAS DIÉSEL", "LITROS DIÉSEL", "IMPORTE DIÉSEL ($)",
        "CARGAS GAS.", "LITROS GAS.", "IMPORTE GAS. ($)", "TOTAL LITROS COMB.", "TOTAL INVERSIÓN ($)"
    ]
    ws_obra.row_dimensions[4].height = 22
    for c_i, h_txt in enumerate(headers_ob, start=1):
        c = ws_obra.cell(row=4, column=c_i, value=h_txt)
        c.fill = fill(DARK_SLATE); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center', wrap=True)
        c.border = header_border()

    cur.execute("""
        WITH d AS (
            SELECT obra_destino, COUNT(*) as cargas_d, SUM(litros) as lts_d, SUM(importe_total) as imp_d
            FROM diesel.consumos
            GROUP BY obra_destino
        ),
        g AS (
            SELECT obra_destino, COUNT(*) as cargas_g, SUM(litros) as lts_g, SUM(importe_total) as imp_g
            FROM gasolina.consumos
            GROUP BY obra_destino
        ),
        todas_obras AS (
            SELECT DISTINCT nombre, tipo FROM catalogos.obras
        )
        SELECT 
            o.nombre,
            o.tipo,
            COALESCE(d.cargas_d, 0) as cargas_d,
            COALESCE(d.lts_d, 0) as lts_d,
            COALESCE(d.imp_d, 0) as imp_d,
            COALESCE(g.cargas_g, 0) as cargas_g,
            COALESCE(g.lts_g, 0) as lts_g,
            COALESCE(g.imp_g, 0) as imp_g,
            (COALESCE(d.lts_d, 0) + COALESCE(g.lts_g, 0)) as total_lts,
            (COALESCE(d.imp_d, 0) + COALESCE(g.imp_g, 0)) as total_imp
        FROM todas_obras o
        LEFT JOIN d ON d.obra_destino = o.nombre
        LEFT JOIN g ON g.obra_destino = o.nombre
        WHERE (COALESCE(d.imp_d, 0) + COALESCE(g.imp_g, 0)) > 0
        ORDER BY total_imp DESC;
    """)
    obras_rows = cur.fetchall()

    r_ob = 5
    tot_cd_ob, tot_ld_ob, tot_id_ob = 0, 0.0, 0.0
    tot_cg_ob, tot_lg_ob, tot_ig_ob = 0, 0.0, 0.0

    for idx, ro in enumerate(obras_rows):
        nom_ob, tipo_ob, cd, ld, id_val, cg, lg, ig_val, tl, ti = ro
        cd, ld, id_val = int(cd), float(ld), float(id_val)
        cg, lg, ig_val = int(cg), float(lg), float(ig_val)
        tl, ti = float(tl), float(ti)

        tot_cd_ob += cd; tot_ld_ob += ld; tot_id_ob += id_val
        tot_cg_ob += cg; tot_lg_ob += lg; tot_ig_ob += ig_val

        bg = ALT_ROW_1 if idx % 2 == 0 else ALT_ROW_2

        ws_obra.cell(row=r_ob, column=1, value=nom_ob).fill = fill(bg)
        ws_obra.cell(row=r_ob, column=1).font = font(bold=True, size=9); ws_obra.cell(row=r_ob, column=1).border = thin_border()

        ws_obra.cell(row=r_ob, column=2, value=tipo_ob).fill = fill(bg)
        ws_obra.cell(row=r_ob, column=2).font = font(size=9); ws_obra.cell(row=r_ob, column=2).alignment = align('center', 'center'); ws_obra.cell(row=r_ob, column=2).border = thin_border()

        c = ws_obra.cell(row=r_ob, column=3, value=cd); c.fill = fill(bg); c.font = font(size=9); c.alignment = align('center', 'center'); c.border = thin_border(); c.number_format = '#,##0'
        c = ws_obra.cell(row=r_ob, column=4, value=ld); c.fill = fill(bg); c.font = font(size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
        c = ws_obra.cell(row=r_ob, column=5, value=id_val); c.fill = fill(bg); c.font = font(bold=True, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

        c = ws_obra.cell(row=r_ob, column=6, value=cg); c.fill = fill(bg); c.font = font(size=9); c.alignment = align('center', 'center'); c.border = thin_border(); c.number_format = '#,##0'
        c = ws_obra.cell(row=r_ob, column=7, value=lg); c.fill = fill(bg); c.font = font(size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
        c = ws_obra.cell(row=r_ob, column=8, value=ig_val); c.fill = fill(bg); c.font = font(bold=True, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

        c = ws_obra.cell(row=r_ob, column=9, value=tl); c.fill = fill(bg); c.font = font(size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
        c = ws_obra.cell(row=r_ob, column=10, value=ti); c.fill = fill(bg); c.font = font(bold=True, color='1E3A5F', size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

        ws_obra.row_dimensions[r_ob].height = 18
        r_ob += 1

    # Totales Obra
    ws_obra.cell(row=r_ob, column=1, value="TOTAL GENERAL OBRAS").fill = fill(NAVY)
    ws_obra.cell(row=r_ob, column=1).font = font(bold=True, color=WHITE, size=9); ws_obra.cell(row=r_ob, column=1).border = thin_border()
    ws_obra.cell(row=r_ob, column=2, value="").fill = fill(NAVY); ws_obra.cell(row=r_ob, column=2).border = thin_border()

    c = ws_obra.cell(row=r_ob, column=3, value=tot_cd_ob); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center'); c.border = thin_border(); c.number_format = '#,##0'
    c = ws_obra.cell(row=r_ob, column=4, value=tot_ld_ob); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
    c = ws_obra.cell(row=r_ob, column=5, value=tot_id_ob); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

    c = ws_obra.cell(row=r_ob, column=6, value=tot_cg_ob); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center'); c.border = thin_border(); c.number_format = '#,##0'
    c = ws_obra.cell(row=r_ob, column=7, value=tot_lg_ob); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
    c = ws_obra.cell(row=r_ob, column=8, value=tot_ig_ob); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

    c = ws_obra.cell(row=r_ob, column=9, value=tot_ld_ob + tot_lg_ob); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
    c = ws_obra.cell(row=r_ob, column=10, value=tot_id_ob + tot_ig_ob); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=10); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'
    ws_obra.row_dimensions[r_ob].height = 22

    col_widths_ob = [32, 18, 15, 17, 19, 14, 16, 18, 19, 21]
    for ci, w in enumerate(col_widths_ob, start=1):
        ws_obra.column_dimensions[get_column_letter(ci)].width = w

    # ─────────────────────────────────────────────────────────────────────────────
    # HOJAS 3..N: PESTAÑAS SEMANALES (SEMANA 24 a SEMANA 38)
    # ─────────────────────────────────────────────────────────────────────────────
    # Cache maestro autorizaciones para gasolina
    cur.execute("""
        SELECT responsable, placas, importe_semanal 
        FROM gasolina.autorizaciones_maestro 
        WHERE activo = TRUE;
    """)
    maestro_aut_map = {}
    for r in cur.fetchall():
        resp_clean = str(r[0] or '').strip().upper()
        placa_clean = str(r[1] or '').strip().upper()
        if placa_clean and placa_clean not in ('SP', 'SN', 'NA', 'XXXXXXX'):
            maestro_aut_map[placa_clean] = float(r[2] or 0)
        if resp_clean:
            maestro_aut_map[resp_clean] = float(r[2] or 0)

    for w in all_weeks:
        sem_str = str(w)
        ws_sem = wb.create_sheet(title=f"SEMANA {sem_str}")
        ws_sem.views.sheetView[0].showGridLines = True

        # Header de Semana
        ws_sem.merge_cells("A1:K1")
        t_s1 = ws_sem.cell(row=1, column=1, value=f"REPORTE DE OPERACIÓN Y COMBUSTIBLE — SEMANA {sem_str}")
        t_s1.fill = fill(NAVY); t_s1.font = font(bold=True, color=WHITE, size=12); t_s1.alignment = align('center', 'center')
        ws_sem.row_dimensions[1].height = 25

        ws_sem.merge_cells("A2:K2")
        t_s2 = ws_sem.cell(row=2, column=1, value=f"Grupo Fénix — Detalle Integrado de Suministros Diésel y Consumo de Gasolina | Semana Operativa {sem_str}")
        t_s2.fill = fill(DARK_SLATE); t_s2.font = font(italic=True, color='94A3B8', size=9); t_s2.alignment = align('center', 'center')
        ws_sem.row_dimensions[2].height = 18

        # --- SECCIÓN 1: DIÉSEL OPERATIVO ---
        ws_sem.merge_cells("A4:K4")
        sec_d = ws_sem.cell(row=4, column=1, value="🛢️ SECCIÓN I: CONSUMO DE DIÉSEL (MAQUINARIA, FRENTES DE OBRA Y PIPAS)")
        sec_d.fill = fill('1E3A8A'); sec_d.font = font(bold=True, color=WHITE, size=10); sec_d.alignment = align('left', 'center')
        ws_sem.row_dimensions[4].height = 22

        headers_d = [
            "FOLIO", "FECHA", "OBRA DESTINO", "EQUIPO", "NO. ECONÓMICO",
            "OPERADOR / RESPONSABLE", "LITROS", "COSTO/LT ($)", "IMPORTE TOTAL ($)", "TIPO CARGA", "ESTATUS"
        ]
        ws_sem.row_dimensions[5].height = 20
        for c_i, h_txt in enumerate(headers_d, start=1):
            c = ws_sem.cell(row=5, column=c_i, value=h_txt)
            c.fill = fill(DARK_SLATE); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center')
            c.border = header_border()

        cur.execute("""
            SELECT folio_conciliacion, fecha, obra_destino, equipo, equipo_economico,
                   COALESCE(operador, responsable, 'SIN REGISTRO') as operador,
                   litros, costo_por_litro, importe_total, tipo_movimiento, estatus_revision
            FROM diesel.consumos
            WHERE semana = %s OR semana = %s
            ORDER BY fecha, id;
        """, (sem_str, f"SEMANA {sem_str}"))
        diesel_movs = cur.fetchall()

        curr_r = 6
        sub_d_lts, sub_d_imp = 0.0, 0.0

        for r_i, dm in enumerate(diesel_movs):
            fol, fch, ob, eq, ec, op, lts, clt, imp, mov, est = dm
            lts = float(lts or 0); clt = float(clt or 0); imp = float(imp or 0)
            sub_d_lts += lts; sub_d_imp += imp

            bg = ALT_ROW_1 if r_i % 2 == 0 else ALT_ROW_2

            ws_sem.cell(row=curr_r, column=1, value=fol or f"D-{r_i+1}").fill = fill(bg)
            ws_sem.cell(row=curr_r, column=1).font = font(bold=True, size=8); ws_sem.cell(row=curr_r, column=1).alignment = align('center', 'center'); ws_sem.cell(row=curr_r, column=1).border = thin_border()

            ws_sem.cell(row=curr_r, column=2, value=str(fch or '')).fill = fill(bg)
            ws_sem.cell(row=curr_r, column=2).font = font(size=8); ws_sem.cell(row=curr_r, column=2).alignment = align('center', 'center'); ws_sem.cell(row=curr_r, column=2).border = thin_border()

            ws_sem.cell(row=curr_r, column=3, value=ob or 'CORPORATIVO').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=3).font = font(bold=True, size=8); ws_sem.cell(row=curr_r, column=3).border = thin_border()

            ws_sem.cell(row=curr_r, column=4, value=eq or '').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=4).font = font(size=8); ws_sem.cell(row=curr_r, column=4).border = thin_border()

            ws_sem.cell(row=curr_r, column=5, value=ec or '').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=5).font = font(size=8); ws_sem.cell(row=curr_r, column=5).alignment = align('center', 'center'); ws_sem.cell(row=curr_r, column=5).border = thin_border()

            ws_sem.cell(row=curr_r, column=6, value=op or '').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=6).font = font(size=8); ws_sem.cell(row=curr_r, column=6).border = thin_border()

            c = ws_sem.cell(row=curr_r, column=7, value=lts); c.fill = fill(bg); c.font = font(size=8); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
            c = ws_sem.cell(row=curr_r, column=8, value=clt); c.fill = fill(bg); c.font = font(size=8); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'
            c = ws_sem.cell(row=curr_r, column=9, value=imp); c.fill = fill(bg); c.font = font(bold=True, size=8); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

            ws_sem.cell(row=curr_r, column=10, value=mov or '').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=10).font = font(size=8); ws_sem.cell(row=curr_r, column=10).alignment = align('center', 'center'); ws_sem.cell(row=curr_r, column=10).border = thin_border()

            ws_sem.cell(row=curr_r, column=11, value=est or 'APROBADO').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=11).font = font(size=8); ws_sem.cell(row=curr_r, column=11).alignment = align('center', 'center'); ws_sem.cell(row=curr_r, column=11).border = thin_border()

            ws_sem.row_dimensions[curr_r].height = 17
            curr_r += 1

        # Subtotal Diésel
        ws_sem.cell(row=curr_r, column=1, value=f"SUBTOTAL DIÉSEL (SEM {sem_str})").fill = fill(SLATE_BLUE)
        ws_sem.cell(row=curr_r, column=1).font = font(bold=True, color=WHITE, size=8); ws_sem.cell(row=curr_r, column=1).border = thin_border()
        for ci_d in range(2, 7):
            ws_sem.cell(row=curr_r, column=ci_d, value="").fill = fill(SLATE_BLUE); ws_sem.cell(row=curr_r, column=ci_d).border = thin_border()

        c = ws_sem.cell(row=curr_r, column=7, value=sub_d_lts); c.fill = fill(SLATE_BLUE); c.font = font(bold=True, color=WHITE, size=8); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
        ws_sem.cell(row=curr_r, column=8, value="").fill = fill(SLATE_BLUE); ws_sem.cell(row=curr_r, column=8).border = thin_border()
        c = ws_sem.cell(row=curr_r, column=9, value=sub_d_imp); c.fill = fill(SLATE_BLUE); c.font = font(bold=True, color=WHITE, size=8); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'
        for ci_d in range(10, 12):
            ws_sem.cell(row=curr_r, column=ci_d, value="").fill = fill(SLATE_BLUE); ws_sem.cell(row=curr_r, column=ci_d).border = thin_border()

        ws_sem.row_dimensions[curr_r].height = 20
        curr_r += 2

        # --- SECCIÓN 2: GASOLINA Y VEHÍCULOS ---
        ws_sem.merge_cells(start_row=curr_r, start_column=1, end_row=curr_r, end_column=11)
        sec_g = ws_sem.cell(row=curr_r, column=1, value="🚗 SECCIÓN II: CONSUMO DE GASOLINA (VEHÍCULOS Y CENTROS DE TRABAJO)")
        sec_g.fill = fill('047857'); sec_g.font = font(bold=True, color=WHITE, size=10); sec_g.alignment = align('left', 'center')
        ws_sem.row_dimensions[curr_r].height = 22
        curr_r += 1

        headers_g = [
            "FOLIO", "FECHA", "OBRA / CENTRO", "RESPONSABLE", "VEHÍCULO",
            "PLACA", "PROVEEDOR", "LITROS", "COSTO/LT ($)", "IMPORTE ($)", "ESTATUS"
        ]
        ws_sem.row_dimensions[curr_r].height = 20
        for c_i, h_txt in enumerate(headers_g, start=1):
            c = ws_sem.cell(row=curr_r, column=c_i, value=h_txt)
            c.fill = fill(DARK_SLATE); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center')
            c.border = header_border()
        curr_r += 1

        cur.execute("""
            SELECT folio_conciliacion, fecha, obra_destino, conductor, vehiculo, placa,
                   gasolineria, litros, costo_por_litro, importe_total, estatus_revision
            FROM gasolina.consumos
            WHERE semana = %s OR semana = %s
            ORDER BY fecha, id;
        """, (sem_str, f"SEMANA {sem_str}"))
        gasolina_movs = cur.fetchall()

        sub_g_lts, sub_g_imp = 0.0, 0.0

        for r_i, gm in enumerate(gasolina_movs):
            fol, fch, ob, cond, veh, plc, gas_prov, lts, clt, imp, est = gm
            lts = float(lts or 0); clt = float(clt or 0); imp = float(imp or 0)
            sub_g_lts += lts; sub_g_imp += imp

            bg = ALT_ROW_1 if r_i % 2 == 0 else ALT_ROW_2

            ws_sem.cell(row=curr_r, column=1, value=fol or f"G-{r_i+1}").fill = fill(bg)
            ws_sem.cell(row=curr_r, column=1).font = font(bold=True, size=8); ws_sem.cell(row=curr_r, column=1).alignment = align('center', 'center'); ws_sem.cell(row=curr_r, column=1).border = thin_border()

            ws_sem.cell(row=curr_r, column=2, value=str(fch or '')).fill = fill(bg)
            ws_sem.cell(row=curr_r, column=2).font = font(size=8); ws_sem.cell(row=curr_r, column=2).alignment = align('center', 'center'); ws_sem.cell(row=curr_r, column=2).border = thin_border()

            ws_sem.cell(row=curr_r, column=3, value=ob or 'CORPORATIVO').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=3).font = font(bold=True, size=8); ws_sem.cell(row=curr_r, column=3).border = thin_border()

            ws_sem.cell(row=curr_r, column=4, value=cond or '').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=4).font = font(size=8); ws_sem.cell(row=curr_r, column=4).border = thin_border()

            ws_sem.cell(row=curr_r, column=5, value=veh or '').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=5).font = font(size=8); ws_sem.cell(row=curr_r, column=5).border = thin_border()

            ws_sem.cell(row=curr_r, column=6, value=plc or 'S/P').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=6).font = font(size=8); ws_sem.cell(row=curr_r, column=6).alignment = align('center', 'center'); ws_sem.cell(row=curr_r, column=6).border = thin_border()

            ws_sem.cell(row=curr_r, column=7, value=gas_prov or 'LEVET').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=7).font = font(size=8); ws_sem.cell(row=curr_r, column=7).alignment = align('center', 'center'); ws_sem.cell(row=curr_r, column=7).border = thin_border()

            c = ws_sem.cell(row=curr_r, column=8, value=lts); c.fill = fill(bg); c.font = font(size=8); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
            c = ws_sem.cell(row=curr_r, column=9, value=clt); c.fill = fill(bg); c.font = font(size=8); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'
            c = ws_sem.cell(row=curr_r, column=10, value=imp); c.fill = fill(bg); c.font = font(bold=True, size=8); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

            ws_sem.cell(row=curr_r, column=11, value=est or 'APROBADO').fill = fill(bg)
            ws_sem.cell(row=curr_r, column=11).font = font(size=8); ws_sem.cell(row=curr_r, column=11).alignment = align('center', 'center'); ws_sem.cell(row=curr_r, column=11).border = thin_border()

            ws_sem.row_dimensions[curr_r].height = 17
            curr_r += 1

        # Subtotal Gasolina
        ws_sem.cell(row=curr_r, column=1, value=f"SUBTOTAL GASOLINA (SEM {sem_str})").fill = fill('065F46')
        ws_sem.cell(row=curr_r, column=1).font = font(bold=True, color=WHITE, size=8); ws_sem.cell(row=curr_r, column=1).border = thin_border()
        for ci_g in range(2, 8):
            ws_sem.cell(row=curr_r, column=ci_g, value="").fill = fill('065F46'); ws_sem.cell(row=curr_r, column=ci_g).border = thin_border()

        c = ws_sem.cell(row=curr_r, column=8, value=sub_g_lts); c.fill = fill('065F46'); c.font = font(bold=True, color=WHITE, size=8); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
        ws_sem.cell(row=curr_r, column=9, value="").fill = fill('065F46'); ws_sem.cell(row=curr_r, column=9).border = thin_border()
        c = ws_sem.cell(row=curr_r, column=10, value=sub_g_imp); c.fill = fill('065F46'); c.font = font(bold=True, color=WHITE, size=8); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'
        ws_sem.cell(row=curr_r, column=11, value="").fill = fill('065F46'); ws_sem.cell(row=curr_r, column=11).border = thin_border()

        ws_sem.row_dimensions[curr_r].height = 20
        curr_r += 2

        # TOTAL CONSOLIDADO SEMANAL
        ws_sem.cell(row=curr_r, column=1, value=f"TOTAL CONSOLIDADO SEMANA {sem_str} (DIÉSEL + GASOLINA)").fill = fill(NAVY)
        ws_sem.cell(row=curr_r, column=1).font = font(bold=True, color=WHITE, size=9); ws_sem.cell(row=curr_r, column=1).border = thin_border()
        for ci_t in range(2, 8):
            ws_sem.cell(row=curr_r, column=ci_t, value="").fill = fill(NAVY); ws_sem.cell(row=curr_r, column=ci_t).border = thin_border()

        c = ws_sem.cell(row=curr_r, column=8, value=sub_d_lts + sub_g_lts); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '#,##0.00 "L"'
        ws_sem.cell(row=curr_r, column=9, value="").fill = fill(NAVY); ws_sem.cell(row=curr_r, column=9).border = thin_border()
        c = ws_sem.cell(row=curr_r, column=10, value=sub_d_imp + sub_g_imp); c.fill = fill(NAVY); c.font = font(bold=True, color=WHITE, size=10); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'
        ws_sem.cell(row=curr_r, column=11, value="").fill = fill(NAVY); ws_sem.cell(row=curr_r, column=11).border = thin_border()

        ws_sem.row_dimensions[curr_r].height = 22

        # Anchos de columna pestaña semanal
        col_widths_sem = [18, 13, 26, 26, 18, 14, 15, 15, 14, 17, 14]
        for ci, w_w in enumerate(col_widths_sem, start=1):
            ws_sem.column_dimensions[get_column_letter(ci)].width = w_w

    cur.close()
    if own_conn:
        db_conn.close()

    # Guardar en archivo local
    if ruta_guardado:
        os.makedirs(os.path.dirname(ruta_guardado), exist_ok=True)
        wb.save(ruta_guardado)
        print(f"Libro Maestro guardado con éxito en: {ruta_guardado}")

    # Retornar buffer para streaming en Flask
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

if __name__ == '__main__':
    buf = generar_libro_maestro_combustible(ruta_guardado="Reportes/Reporte_General_Combustible.xlsx")
    print("Generación finalizada correctamente.")
