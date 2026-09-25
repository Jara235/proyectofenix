import openpyxl

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)
print("Sheet names:", wb.sheetnames)

targets = ['10881', '10922', '10929', '10943', 'A10881', 'A10922', 'A10929', 'A10943', '4184', '4,184']

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val is not None:
                s_val = str(val).strip()
                for t in targets:
                    if t.lower() in s_val.lower():
                        # Print context: whole row
                        row_vals = [ws.cell(r, col).value for col in range(1, ws.max_column + 1) if ws.cell(r, col).value is not None]
                        print(f"\n[FOUND in '{sheetname}' Row {r}, Col {c}] Target: {t} | Value: {val}")
                        print(f"  Row content: {row_vals}")
                        break
