import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import psycopg2
from psycopg2.extras import DictCursor
import datetime

# 35 Exact Rows from the user's template image (Columns A to E)
TEMPLATE_ROWS = [
    ("ING. ARMANDO COLIN", "BACHEO TOLUCA", "CHEVROLET S10", "PAT8298", 1000.00),
    ("ING. ARMANDO COLIN", "BACHEO TOLUCA", "EQUIPO MENOR", "S/P", 500.00),
    ("ING. ANTONIO IBAÑEZ", "MAQUINARIA", "EQUIPO MENOR", "S/P", 500.00),
    ("ING. ANTONIO IBAÑEZ", "MAQUINARIA", "EQUIPO MENOR", "S/P", 500.00),
    ("ALBERTO HERNANDEZ DE JESUS", "MAQUINARIA", "DODGE RAM 700", "PCW9391", 1500.00),
    ("JUAN CARLOS NAZAR CHAVEZ", "MAQUINARIA", "DODGE RAM 700", "PCW9238", 1500.00),
    ("EDGAR VELAZQUEZ ESQUIVEL", "MAQUINARIA", "DODGE RAM 700", "S/P", 1500.00),
    ("ROBERTO CABRERA GARCIA", "H. COLEGIO MILITAR", "CHANGAN ALSVIN", "LMH893A", 500.00),
    ("OSCAR ESTRADA ALVAREZ", "EX HACIENDA DE SAN ENCARNACION", "NISSAN NP-300", "MAH844C", 1000.00),
    ("DIEGO CARREOLA", "CHEDRAUI TOLUCA / T. G. VICENTE LOMBARDO", "MITSUBISHI L200", "NYC829C", 1500.00),
    ("DIEGO CARREOLA", "BACHEO TOLUCA", "EQUIPO MENOR", "S/P", 500.00),
    ("TOMAS ARANDA MONRROY", "CHEDRAUI TOLUCA / T. G. VICENTE LOMBARDO", "DODGE RAM 4000 (3 1/2)", "LF05470", 2500.00),
    ("JAVIER PEREZ DÍAZ", "CHEDRAUI TOLUCA / T. G. VICENTE LOMBARDO / OBRA MÉXICO-TOLUCA", "MITSUBISHI L200", "MHL758A", 2500.00),
    ("ERIK CAMILO RENDON DOMINGUEZ", "OBRA MÉXICO-TOLUCA", "TOYOTA HIACE (URBAN)", "PBT1229", 2500.00),
    ("CRISTIAN REYES GAMORA", "OBRA MÉXICO-TOLUCA", "DODGE RAM 4000 (3 1/2)", "LH49730", 4000.00),
    ("CRISTIAN REYES GAMORA", "OBRA MÉXICO-TOLUCA", "MITSUBISHI L200", "NZT266B", 1500.00),
    ("APOLINAR REYES BOLAINA", "LERMA TENANGO", "MITSUBISHI L200", "NYZ971C", 1500.00),
    ("APOLINAR REYES BOLAINA", "LERMA TENANGO", "CORTADORA DE CONCRETO / MOTOR AUXILIAR PETROL 2 DORS / MOTOR AUXILIAR REPACEDOR", "S/P", 1000.00),
    ("LEONCIO MARTINEZ PASCACIO", "LERMA TENANGO", "MITSUBISHI L200", "NYZ790C", 1500.00),
    ("CARMELO ALVAREZ ANICETO", "LERMA TENANGO", "DODGE RAM 4000 (3 1/2)", "LH49746", 3000.00),
    ("DIEGO FERNANDEZ SANTIAGO", "LERMA TENANGO", "TOYOTA HIACE (URBAN)", "PBT1230", 2500.00),
    ("DAMIAN ANTONIO PUINI", "P. ASFALTO HUIXQUILUCAN", "FORD RANGER", "PCU7482", 2500.00),
    ("LUIS VALDEZ", "P. ASFALTO HUIXQUILUCAN", "EQUIPO MENOR", "S/P", 1000.00),
    ("CLEMENTE SANABRIA", "TRANSPORTES FLOTILLA", "RAM 1200", "LHB188D", 2500.00),
    ("CLEMENTE SANABRIA", "TRANSPORTES FLOTILLA", "CHEVROLET TORNADO", "MNX601B", 1000.00),
    ("BRYAN GABRIEL CHAVEZ ALVAREZ", "TRANSPORTES FLOTILLA", "MOTOCICLETA 1", "NUZ948C", 500.00),
    ("BRYAN GABRIEL CHAVEZ ALVAREZ", "TRANSPORTES FLOTILLA", "MOTOCICLETA 2", "S/P", 500.00),
    ("ROBERTO ORTEGA", "ESTIMACIONES / PLANTA PEGASO", "CHANGAN ALSVIN", "MNX601B", 1000.00),
    ("ROBERTO ORTEGA", "ESTIMACIONES / PLANTA PEGASO", "EQUIPO MENOR", "S/P", 500.00),
    ("JOSE RAUL LUJANO", "ESTIMACIONES / PLANTA PEGASO", "RAM 1200", "LHB184D", 1500.00),
    ("PATRICIO MARTINEZ VILLALOBOS", "JALISCO", "RAM 1500", "LHB182D", 1500.00),
    ("PAOLA JARAMILLO", "COACALCO", "CHANGAN ALSVIN", "LMH893A", 1000.00),
    ("SAMUEL ORTEGA SILVA", "LERMA", "MITSUBISHI L200", "NYZ829C", 2500.00),
    ("CHOFER SILVA", "LERMA", "RAM 1200", "LMH893A", 1500.00),
    ("CHOFER SILVA", "LERMA", "RAM 1200", "LMH893A", 1500.00)
]

def main():
    conn = psycopg2.connect(dbname='fenix_db', user='postgres', password='Bupito*268', host='localhost')
    cur = conn.cursor(cursor_factory=DictCursor)

    # Fetch Week 30 Gasoline Consumos
    cur.execute('''
        SELECT id, fecha, conductor, obra_destino, vehiculo, placa, litros, costo_por_litro, importe_total
        FROM gasolina.consumos
        WHERE semana = '30'
        ORDER BY fecha, conductor, placa
    ''')
    rows_w30 = cur.fetchall()

    print(f"Total Week 30 Gasolina consumos fetched: {len(rows_w30)}")
    
    # Dates present in Week 30 (20/07/2026 to 26/07/2026)
    dates_list = ['2026-07-20', '2026-07-21', '2026-07-22', '2026-07-23', '2026-07-24', '2026-07-25', '2026-07-26']

    # Map loads by (row_index, date, station)
    # Station rule: price == 22.89 -> LEVET, else MOBILE
    load_matrix = {} # (row_idx, date, station) -> sum_importe

    for r in rows_w30:
        cond = (r['conductor'] or '').strip().upper()
        plc = (r['placa'] or '').strip().upper()
        veh = (r['vehiculo'] or '').strip().upper()
        fch = str(r['fecha'])
        cost = float(r['costo_por_litro'] or 0)
        imp = float(r['importe_total'] or 0)

        station = 'LEVET' if abs(cost - 22.89) < 0.01 else 'MOBILE'

        # Match to template row index
        matched_idx = None

        # 1. Try matching by placa first
        if plc and plc != 'S/P':
            for idx, trow in enumerate(TEMPLATE_ROWS):
                t_plc = trow[3].strip().upper()
                if t_plc and t_plc != 'S/P' and (plc in t_plc or t_plc in plc):
                    matched_idx = idx
                    break

        # 2. If no placa match, match by conductor name
        if matched_idx is None and cond:
            for idx, trow in enumerate(TEMPLATE_ROWS):
                t_resp = trow[0].strip().upper()
                # Check direct containment
                cond_clean = cond.replace('ING.', '').replace('CHAVEZ', '').strip()
                t_resp_clean = t_resp.replace('ING.', '').replace('CHAVEZ', '').strip()
                
                parts_cond = cond_clean.split()
                if any(p in t_resp for p in parts_cond if len(p) > 3):
                    matched_idx = idx
                    break

        # Fallback to row 0 if no match
        if matched_idx is None:
            matched_idx = 0

        key = (matched_idx, fch, station)
        load_matrix[key] = load_matrix.get(key, 0.0) + imp

    # Create Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SEMANA 30 GASOLINA"
    ws.sheet_view.showGridLines = True

    # Palette & Styles
    GREEN_HEADER_BG = '6EE7B7' # light emerald green header
    BLUE_SUBHDR_BG  = '38BDF8' # light sky blue
    GREEN_DAY_BG    = 'A7F3D0' # pastel green
    TOTALS_ROW_BG   = 'FED7AA' # warm peach
    BORDER_THIN     = Border(left=Side(style='thin', color='CBD5E1'),
                             right=Side(style='thin', color='CBD5E1'),
                             top=Side(style='thin', color='CBD5E1'),
                             bottom=Side(style='thin', color='CBD5E1'))

    def font_style(bold=False, size=9, color='000000'):
        return Font(name='Arial', size=size, bold=bold, color=color)

    # Row 1: Title
    num_days = len(dates_list)
    total_cols = 5 + (num_days * 3) + 2 # 5 fixed + days*3 + Total Consumido + Remanente

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = "GASOLINA - J.D.J. EQUIPO Y CONSTRUCCIONES S.A. DE C.V. (SEMANA 30)"
    title_cell.fill = PatternFill('solid', fgColor='10B981')
    title_cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')

    # Row 2 & 3 Headers
    headers_fixed = ["RESPONSABLE", "CENTRO DE TRABAJO", "UNIDAD / EQUIPO", "PLACAS", "IMPORTE SEMANAL AUTORIZADO"]
    for col_i, htxt in enumerate(headers_fixed, start=1):
        ws.merge_cells(start_row=2, start_column=col_i, end_row=3, end_column=col_i)
        c = ws.cell(row=2, column=col_i)
        c.value = htxt
        c.fill = PatternFill('solid', fgColor=BLUE_SUBHDR_BG)
        c.font = Font(name='Arial', size=8, bold=True)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDER_THIN

    # Date headers
    col_curr = 6
    for dt_str in dates_list:
        # Format date header DD/MM/YYYY
        dt_obj = datetime.datetime.strptime(dt_str, '%Y-%m-%d')
        lbl_dt = dt_obj.strftime('%d/%m/%Y')

        ws.merge_cells(start_row=2, start_column=col_curr, end_row=2, end_column=col_curr+2)
        c_dt = ws.cell(row=2, column=col_curr)
        c_dt.value = lbl_dt
        c_dt.fill = PatternFill('solid', fgColor=GREEN_DAY_BG)
        c_dt.font = Font(name='Arial', size=8, bold=True)
        c_dt.alignment = Alignment(horizontal='center', vertical='center')
        c_dt.border = BORDER_THIN

        # Sub-headers row 3
        sub_hdrs = ["LEVET", "MOBILE", "31 VALL"]
        for sh_i, sh_txt in enumerate(sub_hdrs):
            c_sh = ws.cell(row=3, column=col_curr + sh_i)
            c_sh.value = sh_txt
            c_sh.fill = PatternFill('solid', fgColor='E2E8F0')
            c_sh.font = Font(name='Arial', size=7, bold=True)
            c_sh.alignment = Alignment(horizontal='center', vertical='center')
            c_sh.border = BORDER_THIN

        col_curr += 3

    # Total Consumido & Remanente headers
    ws.merge_cells(start_row=2, start_column=col_curr, end_row=3, end_column=col_curr)
    c_tot = ws.cell(row=2, column=col_curr)
    c_tot.value = "TOTAL CONSUMIDO"
    c_tot.fill = PatternFill('solid', fgColor='38BDF8')
    c_tot.font = Font(name='Arial', size=8, bold=True)
    c_tot.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c_tot.border = BORDER_THIN

    ws.merge_cells(start_row=2, start_column=col_curr+1, end_row=3, end_column=col_curr+1)
    c_rem = ws.cell(row=2, column=col_curr+1)
    c_rem.value = "REMANENTE / SALDO RESTANTE"
    c_rem.fill = PatternFill('solid', fgColor='38BDF8')
    c_rem.font = Font(name='Arial', size=8, bold=True)
    c_rem.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c_rem.border = BORDER_THIN

    # Populate Data Rows (Row 4 to Row 38)
    row_start = 4
    for idx, trow in enumerate(TEMPLATE_ROWS):
        r_num = row_start + idx
        ws.cell(row=r_num, column=1, value=trow[0]).alignment = Alignment(horizontal='left', vertical='center')
        ws.cell(row=r_num, column=2, value=trow[1]).alignment = Alignment(horizontal='left', vertical='center')
        ws.cell(row=r_num, column=3, value=trow[2]).alignment = Alignment(horizontal='left', vertical='center')
        ws.cell(row=r_num, column=4, value=trow[3]).alignment = Alignment(horizontal='center', vertical='center')

        c_auth = ws.cell(row=r_num, column=5, value=trow[4])
        c_auth.number_format = '"$"#,##0.00'
        c_auth.alignment = Alignment(horizontal='right', vertical='center')

        # Fill cols A-E border & font
        for ci in range(1, 6):
            c_cell = ws.cell(row=r_num, column=ci)
            c_cell.font = font_style(size=8)
            c_cell.border = BORDER_THIN

        row_total_consumed = 0.0

        # Fill Daily columns
        col_c = 6
        for dt_str in dates_list:
            v_levet = load_matrix.get((idx, dt_str, 'LEVET'), 0.0)
            v_mobile = load_matrix.get((idx, dt_str, 'MOBILE'), 0.0)
            v_vall = 0.0 # 31 vall

            # LEVET
            c_lev = ws.cell(row=r_num, column=col_c)
            if v_levet > 0:
                c_lev.value = round(v_levet, 2)
                c_lev.number_format = '"$"#,##0.00'
                c_lev.fill = PatternFill('solid', fgColor='BBF7D0') # green highlight
                c_lev.font = Font(name='Arial', size=8, bold=True, color='166534')
            c_lev.border = BORDER_THIN
            c_lev.alignment = Alignment(horizontal='right', vertical='center')

            # MOBILE
            c_mob = ws.cell(row=r_num, column=col_c+1)
            if v_mobile > 0:
                c_mob.value = round(v_mobile, 2)
                c_mob.number_format = '"$"#,##0.00'
                c_mob.fill = PatternFill('solid', fgColor='FECACA') # red/pink highlight
                c_mob.font = Font(name='Arial', size=8, bold=True, color='991B1B')
            c_mob.border = BORDER_THIN
            c_mob.alignment = Alignment(horizontal='right', vertical='center')

            # 31 VALL
            c_val = ws.cell(row=r_num, column=col_c+2)
            if v_vall > 0:
                c_val.value = round(v_vall, 2)
                c_val.number_format = '"$"#,##0.00'
                c_val.fill = PatternFill('solid', fgColor='FEF08A') # yellow
            c_val.border = BORDER_THIN
            c_val.alignment = Alignment(horizontal='right', vertical='center')

            row_total_consumed += (v_levet + v_mobile + v_vall)
            col_c += 3

        # Total Consumido
        c_tc = ws.cell(row=r_num, column=col_c)
        if row_total_consumed > 0:
            c_tc.value = round(row_total_consumed, 2)
        c_tc.number_format = '"$"#,##0.00'
        c_tc.font = font_style(bold=True, size=8)
        c_tc.alignment = Alignment(horizontal='right', vertical='center')
        c_tc.border = BORDER_THIN

        # Remanente
        remanente = trow[4] - row_total_consumed
        c_rem_cell = ws.cell(row=r_num, column=col_c+1)
        c_rem_cell.value = round(remanente, 2)
        c_rem_cell.number_format = '"$"#,##0.00'
        rem_bg = 'E6F4EA' if remanente >= 0 else 'FCE4E4'
        rem_fg = '166534' if remanente >= 0 else 'DC2626'
        c_rem_cell.fill = PatternFill('solid', fgColor=rem_bg)
        c_rem_cell.font = Font(name='Arial', size=8, bold=True, color=rem_fg)
        c_rem_cell.alignment = Alignment(horizontal='right', vertical='center')
        c_rem_cell.border = BORDER_THIN

    # Totals Row (Row 39)
    last_r = row_start + len(TEMPLATE_ROWS)
    ws.merge_cells(start_row=last_r, start_column=1, end_row=last_r, end_column=4)
    tot_label = ws.cell(row=last_r, column=1)
    tot_label.value = "TOTALES"
    tot_label.fill = PatternFill('solid', fgColor=TOTALS_ROW_BG)
    tot_label.font = Font(name='Arial', size=9, bold=True)
    tot_label.alignment = Alignment(horizontal='center', vertical='center')
    tot_label.border = BORDER_THIN

    for ci in range(2, 5):
        c_dummy = ws.cell(row=last_r, column=ci)
        c_dummy.fill = PatternFill('solid', fgColor=TOTALS_ROW_BG)
        c_dummy.border = BORDER_THIN

    # Sum E
    sum_auth = sum(tr[4] for tr in TEMPLATE_ROWS)
    c_tot_e = ws.cell(row=last_r, column=5)
    c_tot_e.value = sum_auth
    c_tot_e.number_format = '"$"#,##0.00'
    c_tot_e.fill = PatternFill('solid', fgColor=TOTALS_ROW_BG)
    c_tot_e.font = Font(name='Arial', size=9, bold=True)
    c_tot_e.alignment = Alignment(horizontal='right', vertical='center')
    c_tot_e.border = BORDER_THIN

    # Column sum formulas for daily columns & total consumed & remanente
    for col_idx in range(6, total_cols + 1):
        col_let = get_column_letter(col_idx)
        c_sum = ws.cell(row=last_r, column=col_idx)
        c_sum.value = f"=SUM({col_let}4:{col_let}{last_r-1})"
        c_sum.number_format = '"$"#,##0.00'
        c_sum.fill = PatternFill('solid', fgColor=TOTALS_ROW_BG)
        c_sum.font = Font(name='Arial', size=9, bold=True)
        c_sum.alignment = Alignment(horizontal='right', vertical='center')
        c_sum.border = BORDER_THIN

    # Column widths
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 32
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 18

    for cidx in range(6, total_cols + 1):
        ws.column_dimensions[get_column_letter(cidx)].width = 14

    out_file = "c:\\Users\\JOSE\\Desktop\\Control_Gasolina_Semana_30_Oficial.xlsx"
    wb.save(out_file)
    conn.close()
    print(f"=== File saved successfully at: {out_file} ===")

if __name__ == '__main__':
    main()
