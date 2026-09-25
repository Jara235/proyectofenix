import openpyxl
import psycopg2
from psycopg2.extras import DictCursor
import re

def analyze_and_compare():
    path = r'c:\Users\JOSE\Desktop\Proyecto fenix\CONCILIACIÓN COMBUSTIBLE GASOLINERAS.xlsx'
    wb = openpyxl.load_workbook(path, data_only=True)

    catalog = []
    sheet_names = ['COPIA BASE', 'SEMANA 29', 'SEMANA 28', 'SEMANA 27', 'SEMANA 26']

    # Regex for valid vehicle plates (5 to 8 alphanumeric characters, not pure numbers)
    plate_regex = re.compile(r'^(?=.*[A-Z])[A-Z0-9]{5,8}$')

    for sname in sheet_names:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        curr_resp = None
        curr_centro = None
        for r in range(4, ws.max_row + 1):
            resp = ws.cell(r, 2).value
            centro = ws.cell(r, 3).value
            unidad = ws.cell(r, 4).value
            placa = ws.cell(r, 5).value
            monto = ws.cell(r, 6).value

            if resp is not None and str(resp).strip() != '':
                curr_resp = str(resp).strip()
            if centro is not None and str(centro).strip() != '':
                curr_centro = str(centro).strip()

            monto_val = 0.0
            if monto is not None:
                try:
                    monto_val = float(monto)
                except (ValueError, TypeError):
                    monto_val = 0.0

            plc_clean = str(placa).strip().upper() if placa else None
            
            # Validate plate format
            if plc_clean and not plate_regex.match(plc_clean):
                plc_clean = None

            und_str = str(unidad).strip() if unidad else None
            if und_str and (und_str.upper() in ('UNIDAD / EQUIPO', 'TOTAL', 'CONSUMO', 'NONE', 'N/A', '') or 'FACTURA' in und_str.upper() or 'SEMANA' in und_str.upper()):
                und_str = None

            if und_str or plc_clean:
                catalog.append({
                    'sheet': sname,
                    'responsable': curr_resp,
                    'centro_trabajo': curr_centro,
                    'unidad_equipo': und_str,
                    'placa': plc_clean,
                    'monto_autorizado': monto_val
                })

    # Filter unique plates
    by_placa = {}
    by_vehiculo = {}

    for c in catalog:
        plc = c['placa']
        if plc and (plc not in by_placa or (c['monto_autorizado'] > by_placa[plc]['monto_autorizado'])):
            by_placa[plc] = c
        veh = c['unidad_equipo']
        if veh and (veh.upper() not in by_vehiculo):
            by_vehiculo[veh.upper()] = c

    print("=== CATALOGO DE PLACAS Y VEHICULOS REALES ===")
    print(f"Placas únicas válidas catalogadas: {len(by_placa)}")

    # Fetch DB records from gasolina.consumos
    conn = psycopg2.connect(dbname='fenix_db', user='postgres', password='Bupito*268', host='localhost')
    cur = conn.cursor(cursor_factory=DictCursor)

    cur.execute('''
        SELECT id, folio_conciliacion, fecha, semana, origen, obra_destino, vehiculo, placa, conductor, litros, costo_por_litro, importe_total, estatus_revision, observaciones
        FROM gasolina.consumos
        ORDER BY semana DESC, fecha, id
    ''')
    db_rows = [dict(r) for r in cur.fetchall()]
    conn.close()

    proposed_updates = []

    for r in db_rows:
        r_id = r['id']
        r_folio = r['folio_conciliacion'] or f"ID-{r_id}"
        r_semana = r['semana']
        r_placa = (r['placa'] or '').strip().upper()
        r_veh = (r['vehiculo'] or '').strip().upper()
        r_cond = (r['conductor'] or '').strip()
        r_obra = (r['obra_destino'] or '').strip()

        matched_cat = None
        match_by = None

        if r_placa and r_placa not in ('S/P', 'N/A', 'NONE', '') and r_placa in by_placa:
            matched_cat = by_placa[r_placa]
            match_by = 'PLACA'
        elif r_veh and r_veh in by_vehiculo:
            matched_cat = by_vehiculo[r_veh]
            match_by = 'VEHICULO'

        prop_cond = matched_cat['responsable'] if (matched_cat and (not r_cond or r_cond in ('No Identificado', 'N/I', ''))) else r_cond
        prop_obra = matched_cat['centro_trabajo'] if (matched_cat and (not r_obra or r_obra in ('SIN OBRA', 'N/A', ''))) else r_obra
        prop_veh = matched_cat['unidad_equipo'] if (matched_cat and (not r['vehiculo'] or r['vehiculo'] in ('N/A', 'S/P', ''))) else r['vehiculo']
        prop_placa = matched_cat['placa'] if (matched_cat and (not r['placa'] or r['placa'] in ('S/P', 'N/A', ''))) else r['placa']
        prop_monto = matched_cat['monto_autorizado'] if matched_cat else 0.0

        has_change = (
            (prop_cond and prop_cond != r_cond) or 
            (prop_obra and prop_obra != r_obra) or 
            (prop_veh and prop_veh != r['vehiculo']) or 
            (prop_placa and prop_placa != r['placa'])
        )

        proposed_updates.append({
            'id': r_id,
            'folio': r_folio,
            'semana': r_semana,
            'fecha': str(r['fecha']),
            'actual_placa': r['placa'] or 'S/P',
            'prop_placa': prop_placa or 'S/P',
            'actual_vehiculo': r['vehiculo'] or 'N/A',
            'prop_vehiculo': prop_veh or 'N/A',
            'actual_conductor': r_cond or 'N/I',
            'prop_conductor': prop_cond or 'N/I',
            'actual_obra': r_obra or 'SIN OBRA',
            'prop_obra': prop_obra or 'SIN OBRA',
            'monto_autorizado': prop_monto,
            'match_by': match_by or 'SIN COINCIDENCIA',
            'has_change': has_change
        })

    changed_count = sum(1 for p in proposed_updates if p['has_change'])

    # Build markdown report
    md = []
    md.append("# ⛽ Conciliación y Complementación de Gasolina\n")
    md.append("## Matriz Extraída de `CONCILIACIÓN COMBUSTIBLE GASOLINERAS.xlsx` vs Base de Datos\n")
    md.append(f"- **Total Registros en BD:** `{len(db_rows)} cargas`")
    md.append(f"- **Registros a Complementar con el Catálogo:** `{changed_count} cargas`")
    md.append(f"- **Registros Ya Completos:** `{len(db_rows) - changed_count} cargas`\n")

    md.append("### 1. Catálogo de Referencia Extraído del Archivo Excel")
    md.append("| # | Responsable | Centro de Trabajo / Obra | Unidad / Equipo | Placa | Importe Semanal Autorizado ($) |")
    md.append("|---|-------------|--------------------------|-----------------|-------|-------------------------------:|")

    for idx, (plc, c) in enumerate(sorted(by_placa.items()), 1):
        md.append(f"| {idx} | {c['responsable']} | {c['centro_trabajo']} | {c['unidad_equipo']} | **`{plc}`** | ${c['monto_autorizado']:,.2f} |")

    md.append("\n---\n")

    md.append("### 2. Propuesta de Complementación de Registros (Pre-Inyección BD)")
    md.append("> [!IMPORTANT]")
    md.append("> **Revisión antes de inyectar:** A continuación se detalla cada registro de la BD. Las celdas marcadas con **➡️** resaltan los campos incompletos que se van a llenar automáticamente con los datos oficiales del Excel.")

    md.append("\n| Folio | Sem. | Placa (Actual ➡️ Propuesta) | Vehículo (Actual ➡️ Propuesta) | Conductor/Responsable (Actual ➡️ Propuesta) | Obra/Centro (Actual ➡️ Propuesta) | Estatus |")
    md.append("|-------|-----:|-----------------------------|--------------------------------|----------------------------------------------|-----------------------------------|---------|")

    for p in proposed_updates:
        status_tag = "✏️ **Complementar**" if p['has_change'] else "✅ **Sin Cambios**"

        placa_str = f"`{p['actual_placa']}`" if p['actual_placa'] == p['prop_placa'] else f"`{p['actual_placa']}` ➡️ **`{p['prop_placa']}`**"
        veh_str = f"{p['actual_vehiculo']}" if p['actual_vehiculo'] == p['prop_vehiculo'] else f"{p['actual_vehiculo']} ➡️ **{p['prop_vehiculo']}**"
        cond_str = f"{p['actual_conductor']}" if p['actual_conductor'] == p['prop_conductor'] else f"{p['actual_conductor']} ➡️ **{p['prop_conductor']}**"
        obra_str = f"{p['actual_obra']}" if p['actual_obra'] == p['prop_obra'] else f"{p['actual_obra']} ➡️ **{p['prop_obra']}**"

        md.append(f"| `{p['folio']}` | {p['semana']} | {placa_str} | {veh_str} | {cond_str} | {obra_str} | {status_tag} |")

    art_path = r'C:\Users\JOSE\.gemini\antigravity\brain\f3511922-47fb-4fdf-9cb3-6eae6dade26d\conciliacion_gasolina_propuesta.md'
    with open(art_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(md))

    print(f"\nReport written successfully to artifact: {art_path}")

if __name__ == '__main__':
    analyze_and_compare()
