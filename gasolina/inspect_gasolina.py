import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

import openpyxl, os

BASE = r'c:\Users\JOSE\Desktop\Proyecto fenix\gasolina'

files = {
    'SEMANA26': 'SEMANA 26 (1).xlsx',
    'LEVET':    'CONTROL JDJ PROVISIONAL semana 27.xlsx',
    'HUIX':     'CONSUMOS  DE GASOLINA SEMANALES GT.xlsx',
    'MAESTRO':  'Maestro_Conciliacion_Gasolina.xlsx',
}

for key, fname in files.items():
    path = os.path.join(BASE, fname)
    print(f"\n{'='*60}")
    print(f"ARCHIVO: {fname}")
    print(f"{'='*60}")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        print(f"Hojas: {wb.sheetnames}")
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            print(f"\n  --- Hoja: '{sheet}' ({ws.max_row} filas x {ws.max_column} cols) ---")
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True)):
                if any(v is not None for v in row):
                    print(f"  F{i+1}: {list(row)}")
    except Exception as e:
        print(f"ERROR: {e}")
