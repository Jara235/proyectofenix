import psycopg2
from psycopg2.extras import DictCursor
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_excel_report():
    # 1. Fetch DB records
    conn = psycopg2.connect(dbname='fenix_db', user='postgres', password='Bupito*268', host='localhost')
    cur = conn.cursor(cursor_factory=DictCursor)

    cur.execute('''
        SELECT id, folio_conciliacion, fecha, obra_destino, vehiculo, placa, conductor, litros, costo_por_litro, importe_total, observaciones
        FROM gasolina.consumos
        WHERE semana::text = '30'
        ORDER BY fecha, id
    ''')
    db_rows = [dict(r) for r in cur.fetchall()]
    conn.close()

    # 2. Fetch authorizations from official template
    wb_auth = openpyxl.load_workbook(r'c:\Users\JOSE\Desktop\Control_Gasolina_Semana_30_Oficial.xlsx', data_only=True)
    ws_auth = wb_auth['SEMANA 30 GASOLINA']

    auth_records = []
    for r in range(4, 200):
        resp = ws_auth.cell(r, 1).value
        centro = ws_auth.cell(r, 2).value
        veh = ws_auth.cell(r, 3).value
        placa = str(ws_auth.cell(r, 4).value or '').strip().upper()
        monto_auth = ws_auth.cell(r, 5).value
        if resp and (monto_auth is not None):
            auth_records.append({
                'responsable': str(resp).strip(),
                'centro': str(centro or 'N/A').strip(),
                'vehiculo': str(veh or 'N/A').strip(),
                'placa': placa if placa else 'S/P',
                'importe_autorizado': float(monto_auth or 0)
            })

    # Group consumos by Placa
    consumos_by_placa = {}
    for r in db_rows:
        plc = (r['placa'] or 'S/P').strip().upper()
        veh = (r['vehiculo'] or 'N/A').strip()
        if plc not in consumos_by_placa:
            consumos_by_placa[plc] = {
                'vehiculo': veh,
                'litros': 0.0,
                'importe': 0.0,
                'cargas': 0,
                'conductores': set(),
                'obras': set()
            }
        consumos_by_placa[plc]['litros'] += float(r['litros'] or 0)
        consumos_by_placa[plc]['importe'] += float(r['importe_total'] or 0)
        consumos_by_placa[plc]['cargas'] += 1
        if r['conductor']: consumos_by_placa[plc]['conductores'].add(r['conductor'].strip())
        if r['obra_destino']: consumos_by_placa[plc]['obras'].add(r['obra_destino'].strip())

    # Map authorizations by Placa
    auth_by_placa = {}
    for a in auth_records:
        plc = a['placa']
        if plc not in auth_by_placa:
            auth_by_placa[plc] = {
                'responsable': a['responsable'],
                'centro': a['centro'],
                'vehiculo': a['vehiculo'],
                'importe_autorizado': 0.0
            }
        auth_by_placa[plc]['importe_autorizado'] += a['importe_autorizado']

    all_placas = sorted(list(set(consumos_by_placa.keys()) | set(auth_by_placa.keys())))

    # Create new Excel workbook
    wb = openpyxl.Workbook()
    
    # Styling definitions
    font_title = Font(name='Segoe UI', size=14, bold=True, color='FFFFFF')
    font_subtitle = Font(name='Segoe UI', size=10, italic=True, color='E2E8F0')
    font_header = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    font_data = Font(name='Segoe UI', size=10)
    font_bold = Font(name='Segoe UI', size=10, bold=True)
    font_total = Font(name='Segoe UI', size=11, bold=True, color='0F172A')

    fill_title = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    fill_header = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_total = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')

    border_thin = Side(border_style='thin', color='CBD5E1')
    border_double = Side(border_style='double', color='0F172A')
    border_thick = Side(border_style='medium', color='334155')

    box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    total_border = Border(top=border_thick, bottom=border_double, left=border_thin, right=border_thin)

    # -------------------------------------------------------------------------
    # SHEET 1: Detalle de Consumos (Base de Datos S30)
    # -------------------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Detalle_Consumos_S30"
    ws1.views.sheetView[0].showGridLines = True

    # Title block
    ws1.merge_cells('A1:J1')
    ws1['A1'] = "REPORTES FÉNIX - DETALLE DE CONSUMOS DE GASOLINA (SEMANA 30)"
    ws1['A1'].font = font_title
    ws1['A1'].fill = fill_title
    ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[1].height = 35

    ws1.merge_cells('A2:J2')
    ws1['A2'] = f"Total de Registros: {len(db_rows)} cargas capturadas"
    ws1['A2'].font = font_subtitle
    ws1['A2'].fill = fill_title
    ws1['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[2].height = 20

    headers1 = ['#', 'Folio', 'Fecha', 'Obra Destino', 'Vehículo', 'Placa', 'Conductor / Responsable', 'Litros', 'P.U. ($)', 'Importe Total ($)']
    ws1.append([]) # blank row 3
    ws1.append(headers1) # row 4
    ws1.row_dimensions[4].height = 25

    for col_num, h in enumerate(headers1, 1):
        cell = ws1.cell(row=4, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = box_border

    tot_litros_1 = 0.0
    tot_importe_1 = 0.0

    for idx, r in enumerate(db_rows, 1):
        row_idx = idx + 4
        litros = float(r['litros'] or 0)
        pu = float(r['costo_por_litro'] or 0)
        importe = float(r['importe_total'] or 0)

        tot_litros_1 += litros
        tot_importe_1 += importe

        row_data = [
            idx,
            r['folio_conciliacion'] or '',
            str(r['fecha'] or ''),
            r['obra_destino'] or 'SIN OBRA',
            r['vehiculo'] or 'N/A',
            r['placa'] or 'S/P',
            r['conductor'] or 'No Identificado',
            litros,
            pu,
            importe
        ]
        ws1.append(row_data)
        ws1.row_dimensions[row_idx].height = 20

        # Style data row
        for c_idx in range(1, 11):
            cell = ws1.cell(row=row_idx, column=c_idx)
            cell.font = font_data
            cell.border = box_border
            if idx % 2 == 0:
                cell.fill = fill_zebra

            if c_idx == 1:
                cell.alignment = Alignment(horizontal='center')
            elif c_idx in (2, 3, 6):
                cell.alignment = Alignment(horizontal='center')
            elif c_idx in (8, 9, 10):
                cell.alignment = Alignment(horizontal='right')

        # Number formats
        ws1.cell(row=row_idx, column=8).number_format = '#,##0.00 "L"'
        ws1.cell(row=row_idx, column=9).number_format = '"$"#,##0.00'
        ws1.cell(row=row_idx, column=10).number_format = '"$"#,##0.00'

    # Total row sheet 1
    tot_row_1 = len(db_rows) + 5
    ws1.cell(row=tot_row_1, column=1, value="TOTAL").font = font_total
    ws1.cell(row=tot_row_1, column=8, value=tot_litros_1).number_format = '#,##0.00 "L"'
    ws1.cell(row=tot_row_1, column=10, value=tot_importe_1).number_format = '"$"#,##0.00'

    for c_idx in range(1, 11):
        cell = ws1.cell(row=tot_row_1, column=c_idx)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = total_border
        if c_idx in (8, 9, 10):
            cell.alignment = Alignment(horizontal='right')

    # -------------------------------------------------------------------------
    # SHEET 2: Resumen por Placa
    # -------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="Resumen_Por_Placa")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells('A1:G1')
    ws2['A1'] = "CONSUMO TOTAL DE GASOLINA AGRUPADO POR PLACA / VEHÍCULO (SEMANA 30)"
    ws2['A1'].font = font_title
    ws2['A1'].fill = fill_title
    ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws2.row_dimensions[1].height = 35

    headers2 = ['Placa', 'Vehículo', 'Cargas', 'Conductor(es)', 'Obra(s)', 'Total Litros (L)', 'Total Importe ($)']
    ws2.append([])
    ws2.append(headers2)
    ws2.row_dimensions[3].height = 25

    for col_num, h in enumerate(headers2, 1):
        cell = ws2.cell(row=3, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = box_border

    r_idx = 4
    for plc in sorted(consumos_by_placa.keys()):
        c = consumos_by_placa[plc]
        cond_str = ", ".join(c['conductores']) if c['conductores'] else "N/I"
        obra_str = ", ".join(c['obras']) if c['obras'] else "SIN OBRA"

        ws2.append([
            plc,
            c['vehiculo'],
            c['cargas'],
            cond_str,
            obra_str,
            c['litros'],
            c['importe']
        ])
        ws2.row_dimensions[r_idx].height = 20

        for c_idx in range(1, 8):
            cell = ws2.cell(row=r_idx, column=c_idx)
            cell.font = font_data
            cell.border = box_border
            if (r_idx % 2) == 0:
                cell.fill = fill_zebra

            if c_idx == 1:
                cell.font = font_bold
                cell.alignment = Alignment(horizontal='center')
            elif c_idx == 3:
                cell.alignment = Alignment(horizontal='center')
            elif c_idx in (6, 7):
                cell.alignment = Alignment(horizontal='right')

        ws2.cell(row=r_idx, column=6).number_format = '#,##0.00 "L"'
        ws2.cell(row=r_idx, column=7).number_format = '"$"#,##0.00'
        r_idx += 1

    # Total row sheet 2
    ws2.cell(row=r_idx, column=1, value="TOTAL").font = font_total
    ws2.cell(row=r_idx, column=3, value=sum(c['cargas'] for c in consumos_by_placa.values()))
    ws2.cell(row=r_idx, column=6, value=tot_litros_1).number_format = '#,##0.00 "L"'
    ws2.cell(row=r_idx, column=7, value=tot_importe_1).number_format = '"$"#,##0.00'

    for c_idx in range(1, 8):
        cell = ws2.cell(row=r_idx, column=c_idx)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = total_border
        if c_idx == 3:
            cell.alignment = Alignment(horizontal='center')
        elif c_idx in (6, 7):
            cell.alignment = Alignment(horizontal='right')

    # -------------------------------------------------------------------------
    # SHEET 3: Comparativo Autorizaciones
    # -------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="Comparativo_Autorizaciones")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells('A1:H1')
    ws3['A1'] = "COMPARATIVO DE CONSUMO VS. AUTORIZACIÓN SEMANAL (SEMANA 30)"
    ws3['A1'].font = font_title
    ws3['A1'].fill = fill_title
    ws3['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws3.row_dimensions[1].height = 35

    headers3 = ['Placa', 'Responsable / Centro', 'Vehículo', 'Importe Autorizado ($)', 'Consumo Real ($)', 'Consumo Real (L)', 'Diferencia / Remanente ($)', 'Estatus']
    ws3.append([])
    ws3.append(headers3)
    ws3.row_dimensions[3].height = 25

    for col_num, h in enumerate(headers3, 1):
        cell = ws3.cell(row=3, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = box_border

    tot_auth_3 = 0.0
    tot_cons_3 = 0.0
    tot_litros_3 = 0.0

    fill_green = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid') # En regla
    fill_red = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')   # Excedido
    fill_amber = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid') # Sin fondo
    fill_gray = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')  # Sin consumo

    font_green = Font(name='Segoe UI', size=10, bold=True, color='166534')
    font_red = Font(name='Segoe UI', size=10, bold=True, color='991B1B')
    font_amber = Font(name='Segoe UI', size=10, bold=True, color='92400E')
    font_gray = Font(name='Segoe UI', size=10, color='64748B')

    r_idx3 = 4
    for plc in all_placas:
        auth_info = auth_by_placa.get(plc, {'responsable': 'Sin Asignar', 'centro': 'N/A', 'vehiculo': 'N/A', 'importe_autorizado': 0.0})
        cons_info = consumos_by_placa.get(plc, {'vehiculo': auth_info['vehiculo'], 'litros': 0.0, 'importe': 0.0, 'cargas': 0})

        monto_auth = auth_info['importe_autorizado']
        monto_cons = cons_info['importe']
        litros_cons = cons_info['litros']
        diferencia = monto_auth - monto_cons

        tot_auth_3 += monto_auth
        tot_cons_3 += monto_cons
        tot_litros_3 += litros_cons

        if monto_auth == 0 and monto_cons > 0:
            estatus_txt = "⚠️ Sin Fondo"
            st_fill = fill_amber
            st_font = font_amber
        elif diferencia < -0.01:
            estatus_txt = "❌ Excedido"
            st_fill = fill_red
            st_font = font_red
        elif monto_cons == 0:
            estatus_txt = "⚪ Sin Consumo"
            st_fill = fill_gray
            st_font = font_gray
        else:
            estatus_txt = "✅ En Regla"
            st_fill = fill_green
            st_font = font_green

        veh_disp = cons_info['vehiculo'] if cons_info['vehiculo'] != 'N/A' else auth_info['vehiculo']
        resp_disp = f"{auth_info['responsable']} ({auth_info['centro']})" if auth_info['centro'] != 'N/A' else auth_info['responsable']

        ws3.append([
            plc,
            resp_disp,
            veh_disp,
            monto_auth,
            monto_cons,
            litros_cons,
            diferencia,
            estatus_txt
        ])
        ws3.row_dimensions[r_idx3].height = 20

        for c_idx in range(1, 9):
            cell = ws3.cell(row=r_idx3, column=c_idx)
            cell.font = font_data
            cell.border = box_border
            if (r_idx3 % 2) == 0:
                cell.fill = fill_zebra

            if c_idx == 1:
                cell.font = font_bold
                cell.alignment = Alignment(horizontal='center')
            elif c_idx in (4, 5, 6, 7):
                cell.alignment = Alignment(horizontal='right')
            elif c_idx == 8:
                cell.fill = st_fill
                cell.font = st_font
                cell.alignment = Alignment(horizontal='center')

        ws3.cell(row=r_idx3, column=4).number_format = '"$"#,##0.00'
        ws3.cell(row=r_idx3, column=5).number_format = '"$"#,##0.00'
        ws3.cell(row=r_idx3, column=6).number_format = '#,##0.00 "L"'
        ws3.cell(row=r_idx3, column=7).number_format = '"$"#,##0.00'

        r_idx3 += 1

    # Total row sheet 3
    tot_dif_3 = tot_auth_3 - tot_cons_3
    ws3.cell(row=r_idx3, column=1, value="TOTAL GLOBAL").font = font_total
    ws3.cell(row=r_idx3, column=4, value=tot_auth_3).number_format = '"$"#,##0.00'
    ws3.cell(row=r_idx3, column=5, value=tot_cons_3).number_format = '"$"#,##0.00'
    ws3.cell(row=r_idx3, column=6, value=tot_litros_3).number_format = '#,##0.00 "L"'
    ws3.cell(row=r_idx3, column=7, value=tot_dif_3).number_format = '"$"#,##0.00'
    ws3.cell(row=r_idx3, column=8, value="✅ En Regla" if tot_dif_3 >= 0 else "❌ Excedido Global").font = font_green if tot_dif_3 >= 0 else font_red

    for c_idx in range(1, 9):
        cell = ws3.cell(row=r_idx3, column=c_idx)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = total_border
        if c_idx in (4, 5, 6, 7):
            cell.alignment = Alignment(horizontal='right')
        elif c_idx == 8:
            cell.alignment = Alignment(horizontal='center')

    # -------------------------------------------------------------------------
    # Auto-adjust column widths for all sheets
    # -------------------------------------------------------------------------
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # ignore merged title rows
                if cell.row in (1, 2):
                    continue
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output_path = r'c:\Users\JOSE\Desktop\Control_Gasolina_Semana_30_Completo.xlsx'
    wb.save(output_path)
    print(f"Excel file created successfully at: {output_path}")

if __name__ == '__main__':
    create_excel_report()
