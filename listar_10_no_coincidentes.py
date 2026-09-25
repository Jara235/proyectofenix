import openpyxl
import psycopg2
from psycopg2.extras import DictCursor

def list_missing_charges():
    excel_path = r'c:\Users\JOSE\Desktop\Proyecto fenix\CONCILIACIÓN COMBUSTIBLE GASOLINERAS.xlsx'
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['SEMANA 30']

    excel_charges = []
    dates = [(7, '2026-07-20'), (10, '2026-07-21'), (13, '2026-07-22'), (16, '2026-07-23'), (19, '2026-07-24'), (22, '2026-07-25'), (25, '2026-07-26')]

    for r in range(4, ws.max_row + 1):
        placa = ws.cell(r, 5).value
        plc = str(placa).strip().upper() if placa else None
        if plc in ('NONE', 'S/P', '', 'N/A', 'PLACAS', 'AUTORIZADO') or (plc and len(plc) > 12):
            plc = None

        for col_start, d_str in dates:
            for val in [ws.cell(r, col_start).value, ws.cell(r, col_start + 1).value, ws.cell(r, col_start + 2).value]:
                if val is not None:
                    try:
                        amt = float(val)
                        if amt > 0:
                            excel_charges.append({'fecha': d_str, 'placa': plc, 'importe': amt})
                    except (ValueError, TypeError):
                        pass

    conn = psycopg2.connect(dbname='fenix_db', user='postgres', password='Bupito*268', host='localhost')
    cur = conn.cursor(cursor_factory=DictCursor)
    cur.execute('''
        SELECT id, folio_conciliacion, fecha, obra_destino, vehiculo, placa, conductor, litros, costo_por_litro, importe_total, observaciones
        FROM gasolina.consumos
        WHERE semana::text = '30'
        ORDER BY fecha, id
    ''')
    db_rows = [dict(r) for r in cur.fetchall()]
    conn.close()

    missing_in_excel = []
    for idx, r in enumerate(db_rows, 1):
        folio = r['folio_conciliacion'] or f"ID-{r['id']}"
        fecha = str(r['fecha'] or '')
        plc = (r['placa'] or 'S/P').strip().upper()
        veh = r['vehiculo'] or 'N/A'
        cond = r['conductor'] or 'No Identificado'
        obra = r['obra_destino'] or 'SIN OBRA'
        imp = float(r['importe_total'] or 0)
        litros = float(r['litros'] or 0)

        match = None
        for c in excel_charges:
            if c['fecha'] == fecha and (c['placa'] == plc or plc == 'S/P'):
                if abs(c['importe'] - imp) < 5.0:
                    match = c
                    break

        if not match:
            missing_in_excel.append({
                'num': idx,
                'folio': folio,
                'fecha': fecha,
                'obra': obra,
                'vehiculo': veh,
                'placa': plc,
                'conductor': cond,
                'litros': litros,
                'importe': imp,
                'obs': r['observaciones'] or ''
            })

    print(f"Total cargas en BD no encontradas en pestaña Excel Semana 30: {len(missing_in_excel)}")
    for m in missing_in_excel:
        print(f"#{m['num']:2d} | Folio: {m['folio']:<16} | Fecha: {m['fecha']} | Obra: {m['obra']:<20} | Veh: {m['vehiculo']:<22} | Placa: {m['placa']:<10} | Cond: {m['conductor']:<22} | Importe: ${m['importe']:,.2f}")

if __name__ == '__main__':
    list_missing_charges()
