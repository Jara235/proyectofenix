import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')
ws = wb['Captura_Facturas_Mobil']
print('--- Verificacion ---')
for r in range(40, 55):
    factura = ws.cell(row=r, column=1).value
    fecha = ws.cell(row=r, column=2).value
    lts = ws.cell(row=r, column=5).value
    semana = ws.cell(row=r, column=14).value
    if factura:
        print(f'R{r}: {factura} | {str(fecha)[:10]} | Lts:{lts} | Semana:{str(semana)[:50]}')

