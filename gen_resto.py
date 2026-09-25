import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def style_headers(ws):
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    align = Alignment(horizontal='center', vertical='center')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align
        ws.column_dimensions[cell.column_letter].width = 20

def create_acarreos():
    wb = openpyxl.Workbook()
    ws_cat = wb.active
    ws_cat.title = 'CATALOGOS'
    ws_cat.append(['OBRAS', 'CODIGO_OBRA', 'SINDICATOS', 'MATERIALES', 'CAMIONES', 'PLACA_CAMION', 'ORIGEN'])
    
    ws_mov = wb.create_sheet('BD_ACARREOS')
    ws_mov.append(['FOLIO_CONCILIACION', 'FECHA', 'SEMANA', 'SINDICATO', 'MATERIAL', 'ORIGEN', 'OBRA_DESTINO', 'PLACA_CAMION', 'CAPACIDAD_M3', 'PRECIO_UNITARIO', 'IMPORTE_TOTAL', 'ESTATUS_CONCILIACION', 'OBSERVACIONES'])
    
    ws_fac = wb.create_sheet('BD_FACTURAS')
    ws_fac.append(['FOLIO_CONCILIACION', 'FOLIO_FACTURA', 'FECHA_FACTURA', 'SEMANA', 'SINDICATO_PROVEEDOR', 'VIAJES_AMPARADOS', 'IMPORTE', 'I.V.A', 'IMPORTE_TOTAL', 'ESTATUS_CONCILIACION', 'OBSERVACIONES'])
    
    ws_td = wb.create_sheet('TABLAS_DINAMICAS')
    ws_td.append(['Aquí puedes insertar tus tablas dinámicas de resumen'])
    
    for sheet in wb.sheetnames:
        style_headers(wb[sheet])
    
    for i in range(2, 501):
        # Importe_Total = Capacidad_M3 * Precio_Unitario
        ws_mov[f'K{i}'] = f'=IF(OR(I{i}="", J{i}=""), "", I{i}*J{i})'
        # Estatus Conciliacion
        ws_mov[f'L{i}'] = f'=IF(A{i}="", "", IF(SUMIFS(BD_FACTURAS!I:I, BD_FACTURAS!A:A, A{i}) = K{i}, "✅ Cuadrado", "⚠️ Diferencia"))'

    wb.save('c:/Users/JOSE/Desktop/Proyecto fenix/MAESTRO_CONTROL_ACARREOS_NUEVO.xlsx')
    print('Created Acarreos')

def create_mezcla():
    wb = openpyxl.Workbook()
    ws_cat = wb.active
    ws_cat.title = 'CATALOGOS'
    ws_cat.append(['MINAS_ORIGEN', 'PLANTAS', 'OBRAS', 'MATERIALES', 'CAMIONES', 'PLACA'])
    
    ws_mp = wb.create_sheet('BD_MATERIA_PRIMA')
    ws_mp.append(['FOLIO_CONCILIACION', 'FECHA', 'SEMANA', 'MINA_ORIGEN', 'PLANTA_DESTINO', 'MATERIAL', 'CAMION', 'PLACA', 'TONELADAS', 'COSTO_MATERIAL', 'FLETE', 'IMPORTE_TOTAL', 'ESTATUS_CONCILIACION', 'OBSERVACIONES'])
    
    ws_prod = wb.create_sheet('BD_PRODUCCION')
    ws_prod.append(['FOLIO_CONCILIACION', 'FECHA', 'SEMANA', 'PLANTA', 'TIPO_MEZCLA', 'TONELADAS_PRODUCIDAS', 'EMULSION_CONSUMIDA_LTS', 'AGREGADOS_CONSUMIDOS_M3', 'COSTO_TOTAL_PRODUCCION', 'OBSERVACIONES'])
    
    ws_ten = wb.create_sheet('BD_TENDIDO')
    ws_ten.append(['FOLIO_CONCILIACION', 'FECHA', 'SEMANA', 'PLANTA_ORIGEN', 'OBRA_DESTINO', 'VIAJES_ENVIADOS', 'TONELADAS_TENDIDAS', 'METROS_CUADRADOS_TENDIDOS', 'RENDIMIENTO_OBRA', 'OBSERVACIONES'])
    
    ws_td = wb.create_sheet('TABLAS_DINAMICAS')
    ws_td.append(['Aquí puedes insertar tus tablas dinámicas de resumen'])
    
    for sheet in wb.sheetnames:
        style_headers(wb[sheet])
    
    wb.save('c:/Users/JOSE/Desktop/Proyecto fenix/MAESTRO_CONTROL_MEZCLA_NUEVO.xlsx')
    print('Created Mezcla')

create_acarreos()
create_mezcla()

