import openpyxl
import psycopg2
from psycopg2.extras import DictCursor
import datetime

def clean_str(s):
    if not s:
        return ""
    return str(s).strip()

def analyze_tables_1_and_2():
    wb = openpyxl.load_workbook(r'c:\Users\JOSE\Desktop\Proyecto fenix\CONCILIACIÓN COMBUSTIBLE GASOLINERAS.xlsx', data_only=True)
    ws = wb['SEMANA 30']

    # Dates header for Table 1 (Row 2):
    # Col 7 (G): 2026-07-20 | Col 10 (J): 2026-07-21 | Col 13 (M): 2026-07-22 | Col 16 (P): 2026-07-23 | Col 19 (S): 2026-07-24 | Col 22 (V): 2026-07-25 | Col 25 (Y): 2026-07-26
    t1_dates = []
    for c in range(7, 28, 3):
        val = ws.cell(2, c).value
        if isinstance(val, (datetime.datetime, datetime.date)):
            t1_dates.append((c, val.strftime('%Y-%m-%d')))

    # Dates header for Table 2 (Row 48):
    t2_dates = []
    for c in range(7, 28, 3):
        val = ws.cell(48, c).value
        if isinstance(val, (datetime.datetime, datetime.date)):
            t2_dates.append((c, val.strftime('%Y-%m-%d')))
        else:
            # Fallback to t1_dates if dates are relative or formatted
            if len(t1_dates) > (c-7)//3:
                t2_dates.append((c, t1_dates[(c-7)//3][1]))

    print("Table 1 Dates:", t1_dates)
    print("Table 2 Dates:", t2_dates)

    # Parse Table 1 (Rows 4 to 46)
    t1_charges = []
    curr_resp = None
    curr_centro = None

    for r in range(4, 47):
        resp = ws.cell(r, 2).value
        centro = ws.cell(r, 3).value
        unidad = ws.cell(r, 4).value
        placa = ws.cell(r, 5).value
        monto_auth = ws.cell(r, 6).value

        if resp is not None and str(resp).strip() != '':
            curr_resp = str(resp).strip()
        if centro is not None and str(centro).strip() != '':
            curr_centro = str(centro).strip()

        plc_str = clean_str(placa).upper()
        und_str = clean_str(unidad)

        if plc_str in ('NONE', 'S/P', 'N/A', 'PLACAS', 'AUTORIZADO') or len(plc_str) > 12:
            plc_str = ''

        if und_str.upper() in ('UNIDAD / EQUIPO', 'TOTAL', 'CONSUMO', 'NONE', 'N/A') or 'FACTURA' in und_str.upper() or 'SEMANA' in und_str.upper():
            und_str = ''

        try:
            monto_val = float(monto_auth or 0)
        except (ValueError, TypeError):
            monto_val = 0.0

        for col_start, d_str in t1_dates:
            for prov_offset, prov_name in enumerate(['LEVET', 'MOBILE', 'SI VALE']):
                val = ws.cell(r, col_start + prov_offset).value
                if val is not None:
                    try:
                        amt = float(val)
                        if amt > 0:
                            t1_charges.append({
                                'tabla': 'Tabla 1 (J.D.J.)',
                                'row': r,
                                'responsable': curr_resp,
                                'centro_trabajo': curr_centro,
                                'unidad': und_str,
                                'placa': plc_str,
                                'monto_auth': monto_val,
                                'fecha': d_str,
                                'proveedor': prov_name,
                                'importe': amt
                            })
                    except (ValueError, TypeError):
                        pass

    # Parse Table 2 (Rows 49 to 61)
    t2_charges = []
    curr_resp = None
    curr_centro = None

    for r in range(49, 62):
        resp = ws.cell(r, 2).value
        centro = ws.cell(r, 3).value
        unidad = ws.cell(r, 4).value
        placa = ws.cell(r, 5).value
        monto_auth = ws.cell(r, 6).value

        if resp is not None and str(resp).strip() != '':
            curr_resp = str(resp).strip()
        if centro is not None and str(centro).strip() != '':
            curr_centro = str(centro).strip()

        plc_str = clean_str(placa).upper()
        und_str = clean_str(unidad)

        if plc_str in ('NONE', 'S/P', 'N/A', 'PLACAS', 'AUTORIZADO') or len(plc_str) > 12:
            plc_str = ''

        if und_str.upper() in ('UNIDAD / EQUIPO', 'TOTAL', 'CONSUMO', 'NONE', 'N/A') or 'FACTURA' in und_str.upper() or 'SEMANA' in und_str.upper():
            und_str = ''

        try:
            monto_val = float(monto_auth or 0)
        except (ValueError, TypeError):
            monto_val = 0.0

        for col_start, d_str in t2_dates:
            for prov_offset, prov_name in enumerate(['LEVET', 'MOBILE', 'SI VALE']):
                val = ws.cell(r, col_start + prov_offset).value
                if val is not None:
                    try:
                        amt = float(val)
                        if amt > 0:
                            t2_charges.append({
                                'tabla': 'Tabla 2 (TRD SAN MIGUEL)',
                                'row': r,
                                'responsable': curr_resp,
                                'centro_trabajo': curr_centro,
                                'unidad': und_str,
                                'placa': plc_str,
                                'monto_auth': monto_val,
                                'fecha': d_str,
                                'proveedor': prov_name,
                                'importe': amt
                            })
                    except (ValueError, TypeError):
                        pass

    all_excel_t1_t2 = t1_charges + t2_charges

    print(f"Total cargas extraídas de Tabla 1: {len(t1_charges)} (Importe: ${sum(c['importe'] for c in t1_charges):,.2f})")
    print(f"Total cargas extraídas de Tabla 2: {len(t2_charges)} (Importe: ${sum(c['importe'] for c in t2_charges):,.2f})")
    print(f"Total cargas Tablas 1 y 2: {len(all_excel_t1_t2)} (Importe: ${sum(c['importe'] for c in all_excel_t1_t2):,.2f})")

    # Fetch DB records for Semana 30
    conn = psycopg2.connect(dbname='fenix_db', user='postgres', password='Bupito*268', host='localhost')
    cur = conn.cursor(cursor_factory=DictCursor)
    cur.execute('''
        SELECT id, folio_conciliacion, fecha, obra_destino, vehiculo, placa, conductor, litros, costo_por_litro, importe_total, observaciones
        FROM gasolina.consumos
        WHERE semana::text = '30'
        ORDER BY fecha, id
    ''')
    db_rows = [dict(r) for r in cur.fetchall()]
    conn.close()

    print(f"Total cargas en Base de Datos Semana 30: {len(db_rows)} (Importe: ${sum(float(r['importe_total'] or 0) for r in db_rows):,.2f})")

    # Compare Excel Tablas 1 & 2 charges against DB records
    used_db_ids = set()
    matched_excel = []
    unmatched_excel = []

    for ex_i, ex in enumerate(all_excel_t1_t2, 1):
        ex_fecha = ex['fecha']
        ex_plc = ex['placa']
        ex_resp = clean_str(ex['responsable']).upper()
        ex_und = clean_str(ex['unidad']).upper()
        ex_imp = ex['importe']

        best_db_match = None
        best_db_score = 0

        for r in db_rows:
            r_id = r['id']
            if r_id in used_db_ids:
                continue

            db_fecha = str(r['fecha'] or '')
            db_plc = clean_str(r['placa']).upper()
            db_cond = clean_str(r['conductor']).upper()
            db_veh = clean_str(r['vehiculo']).upper()
            db_imp = float(r['importe_total'] or 0)

            score = 0

            # Date match
            if db_fecha == ex_fecha:
                score += 10

            # Amount match (within $5)
            if abs(db_imp - ex_imp) < 5.0:
                score += 30
            elif abs(db_imp - ex_imp) < 50.0:
                score += 10

            # Placa match
            if ex_plc and db_plc and (ex_plc in db_plc or db_plc in ex_plc):
                score += 40

            # Responsable match
            if ex_resp and db_cond:
                resp_tokens = set(ex_resp.split())
                cond_tokens = set(db_cond.split())
                overlap = resp_tokens & cond_tokens
                if overlap:
                    score += 20 * len(overlap)

            # Unidad match
            if ex_und and db_veh and (ex_und in db_veh or db_veh in ex_und):
                score += 15

            if score > best_db_score and score >= 30:
                best_db_score = score
                best_db_match = r

        if best_db_match is not None:
            used_db_ids.add(best_db_match['id'])
            matched_excel.append({
                'excel': ex,
                'db': best_db_match,
                'score': best_db_score
            })
        else:
            unmatched_excel.append(ex)

    print(f"\nCargas de Excel (Tablas 1 y 2) que COINCIDEN con la BD: {len(matched_excel)}")
    print(f"Cargas de Excel (Tablas 1 y 2) que NO ESTÁN en la BD (FALTANTES EN BD): {len(unmatched_excel)}")

    # Build Markdown Report
    md = []
    md.append("# ⛽ Cargas de Excel (Tablas 1 y 2 - Semana 30) vs. Base de Datos")
    md.append("## Análisis de Cargas Presentes en el Archivo Excel que Faltan o Coinciden en el Sistema\n")

    md.append("### 📊 Resumen Estadístico")
    md.append(f"- **Total Cargas en Tabla 1 (J.D.J. Gasolina):** `{len(t1_charges)} cargas` | `${sum(c['importe'] for c in t1_charges):,.2f} MXN`")
    md.append(f"- **Total Cargas en Tabla 2 (TRD San Miguel Gasolina):** `{len(t2_charges)} cargas` | `${sum(c['importe'] for c in t2_charges):,.2f} MXN`")
    md.append(f"- **Total Cargas Tablas 1 y 2 Combinadas:** `{len(all_excel_t1_t2)} cargas` | `${sum(c['importe'] for c in all_excel_t1_t2):,.2f} MXN`")
    md.append(f"- **Cargas Coincidentes con la Base de Datos:** `{len(matched_excel)} cargas`")
    md.append(f"- ⚠️ **Cargas Presentes en Excel pero FALTANTES en la Base de Datos:** `{len(unmatched_excel)} cargas` | **Monto Faltante: `${sum(u['importe'] for u in unmatched_excel):,.2f} MXN`**\n")

    md.append("---")
    md.append("### 🔴 1. Cargas del Excel (Tablas 1 y 2) que NO ESTÁN en la Base de Datos")
    md.append("> [!WARNING]")
    md.append("> Estas son las cargas registradas en la plantilla oficial de Excel de la Semana 30 que **aún no han sido capturadas/inyectadas en la Base de Datos**:")

    md.append("\n| # | Tabla Excel | Fecha | Responsable / Conductor | Centro de Trabajo / Obra | Unidad / Equipo | Placa | Proveedor | Importe Excel ($) |")
    md.append("|--:|:------------|:------|:------------------------|:-------------------------|:----------------|:------|:----------|------------------:|")

    for idx, u in enumerate(unmatched_excel, 1):
        md.append(f"| {idx} | {u['tabla']} | {u['fecha']} | {u['responsable']} | {u['centro_trabajo']} | {u['unidad']} | **`{u['placa'] or 'S/P'}`** | {u['proveedor']} | **${u['importe']:,.2f}** |")

    md.append("\n---\n")

    md.append("### 🟢 2. Cargas del Excel (Tablas 1 y 2) que SÍ COINCIDEN con la Base de Datos")
    md.append("| # | Tabla Excel | Fecha | Responsable (Excel ➡️ BD) | Placa (Excel ➡️ BD) | Importe Excel ($) | Importe BD ($) | Folio BD |")
    md.append("|--:|:------------|:------|:--------------------------|:--------------------|------------------:|---------------:|:---------|")

    for idx, m in enumerate(matched_excel, 1):
        ex = m['excel']
        db = m['db']
        folio = db['folio_conciliacion'] or f"ID-{db['id']}"
        imp_db = float(db['importe_total'] or 0)

        md.append(f"| {idx} | {ex['tabla']} | {ex['fecha']} | {ex['responsable']} ➡️ **{db['conductor']}** | `{ex['placa']}` ➡️ **`{db['placa']}`** | **${ex['importe']:,.2f}** | **${imp_db:,.2f}** | `{folio}` |")

    art_path = r'C:\Users\JOSE\.gemini\antigravity\brain\f3511922-47fb-4fdf-9cb3-6eae6dade26d\cargas_faltantes_excel_tablas1_2.md'
    with open(art_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(md))

    print(f"\nReport generated at: {art_path}")

if __name__ == '__main__':
    analyze_tables_1_and_2()
