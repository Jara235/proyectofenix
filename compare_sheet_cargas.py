import openpyxl
from app_admin import get_db

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)
ws = wb['SEMANA 33']

# Extract all individual charges from Google Sheet Table 1 and Table 2
# Day columns mapping in Table 1 (Row 4 to 43)
# Dates:
# Col 7: Lun Levet, Col 8: Lun Mobile, Col 10: Lun Si Vale (10/08/2026)
# Col 11: Mar Levet, Col 12: Mar Mobile, Col 14: Mar Si Vale (11/08/2026)
# Col 15: Mie Levet, Col 16: Mie Mobile, Col 18: Mie Si Vale (12/08/2026)
# Col 19: Jue Levet, Col 20: Jue Mobile, Col 22: Jue Si Vale (13/08/2026)
# Col 23: Vie Levet, Col 24: Vie Mobile, Col 26: Vie Si Vale (14/08/2026)
# Col 27: Sab Levet, Col 28: Sab Mobile, Col 30: Sab Si Vale (15/08/2026)

day_cols = [
    ('2026-08-10', 7, 'LEVET', None),
    ('2026-08-10', 8, 'MOBILE', 9),
    ('2026-08-10', 10, 'SI VALE', None),
    ('2026-08-11', 11, 'LEVET', None),
    ('2026-08-11', 12, 'MOBILE', 13),
    ('2026-08-11', 14, 'SI VALE', None),
    ('2026-08-12', 15, 'LEVET', None),
    ('2026-08-12', 16, 'MOBILE', 17),
    ('2026-08-12', 18, 'SI VALE', None),
    ('2026-08-13', 19, 'LEVET', None),
    ('2026-08-13', 20, 'MOBILE', 21),
    ('2026-08-13', 22, 'SI VALE', None),
    ('2026-08-14', 23, 'LEVET', None),
    ('2026-08-14', 24, 'MOBILE', 25),
    ('2026-08-14', 26, 'SI VALE', None),
    ('2026-08-15', 27, 'LEVET', None),
    ('2026-08-15', 28, 'MOBILE', 29),
    ('2026-08-15', 30, 'SI VALE', None),
]

sheet_cargas = []

# Table 1: JDJ
for r in range(4, 44):
    resp = ws.cell(r, 2).value
    obra = ws.cell(r, 3).value
    unidad = ws.cell(r, 4).value
    placa = ws.cell(r, 5).value
    
    resp_str = str(resp).strip() if resp is not None else ''
    obra_str = str(obra).strip() if obra is not None else ''
    unidad_str = str(unidad).strip() if unidad is not None else ''
    placa_str = str(placa).strip().upper() if placa is not None else ''
    if placa_str in ['PLACAS', 'S/P', 'NONE']: placa_str = ''
    
    conductor = resp_str or unidad_str
    if not conductor or 'TOTAL' in conductor.upper():
        continue
        
    for fecha, col_val, gas, col_folio in day_cols:
        val = ws.cell(r, col_val).value
        try:
            val_f = float(val or 0)
        except:
            val_f = 0.0
        if val_f > 0:
            folio = ''
            if col_folio:
                f_val = ws.cell(r, col_folio).value
                if f_val: folio = str(f_val).strip().upper()
            sheet_cargas.append({
                'row': r,
                'empresa': 'JDJ',
                'conductor': conductor,
                'obra': obra_str,
                'unidad': unidad_str,
                'placa': placa_str,
                'fecha': fecha,
                'gasolineria': gas,
                'importe': val_f,
                'folio_sheet': folio
            })

# Table 2: TRD
for r in range(50, 58):
    resp = ws.cell(r, 2).value
    obra = ws.cell(r, 3).value
    unidad = ws.cell(r, 4).value
    placa = ws.cell(r, 5).value
    
    resp_str = str(resp).strip() if resp is not None else ''
    obra_str = str(obra).strip() if obra is not None else ''
    unidad_str = str(unidad).strip() if unidad is not None else ''
    placa_str = str(placa).strip().upper() if placa is not None else ''
    if placa_str in ['PLACAS', 'S/P', 'NONE']: placa_str = ''
    
    conductor = resp_str or unidad_str
    if not conductor or 'TOTAL' in conductor.upper():
        continue
        
    for fecha, col_val, gas, col_folio in day_cols:
        val = ws.cell(r, col_val).value
        try:
            val_f = float(val or 0)
        except:
            val_f = 0.0
        if val_f > 0:
            folio = ''
            if col_folio:
                f_val = ws.cell(r, col_folio).value
                if f_val: folio = str(f_val).strip().upper()
            sheet_cargas.append({
                'row': r,
                'empresa': 'TRD',
                'conductor': conductor,
                'obra': obra_str,
                'unidad': unidad_str,
                'placa': placa_str,
                'fecha': fecha,
                'gasolineria': gas,
                'importe': val_f,
                'folio_sheet': folio
            })

print(f"Total cargas found in Google Sheet S33: {len(sheet_cargas)}")
total_sheet_imp = sum(c['importe'] for c in sheet_cargas)
print(f"Total importe in Google Sheet S33: ${total_sheet_imp:,.2f}")

print("\n--- MOBIL CARGAS IN GOOGLE SHEET S33 ---")
mobil_sheet = [c for c in sheet_cargas if c['gasolineria'] == 'MOBILE']
for m in mobil_sheet:
    print(f"  Row {m['row']:2d} | {m['fecha']} | {m['conductor'][:22]:22s} | Placa: {m['placa']:8s} | ${m['importe']:8.2f} | Folio: {m['folio_sheet']}")

print(f"Mobil Total in Sheet: ${sum(m['importe'] for m in mobil_sheet):,.2f} ({len(mobil_sheet)} cargas)")
