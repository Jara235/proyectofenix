import os
import xml.etree.ElementTree as ET
import PyPDF2
import re
import openpyxl

base_dir = r"c:\Users\JOSE\Desktop\Proyecto fenix\facturas\Diesel\Semana_28"
excel_path = r"c:\Users\JOSE\Desktop\Proyecto fenix\MAESTRO_CONTROL_DIESEL_NUEVO.xlsx"

# Helper for destination matching
def normalizar_destino(texto):
    if not texto: return "No aplica"
    texto = str(texto).upper()
    if "MEXICO-TOLUCA" in texto or "MÉXICO-TOLUCA" in texto or "MEXICO TOLUCA" in texto:
        return "México-Toluca"
    elif "TRES MARIAS" in texto or "TRES MARÍAS" in texto or "LERMA" in texto:
        return "Lerma - Tres Marías"
    elif "BACHEO" in texto:
        return "Bacheo Toluca"
    elif "HUIXQUILUCAN" in texto:
        return "Planta Asflato Huixquilucan"
    elif "DRAGONES" in texto:
        return "Dragones"
    elif "JALISCO" in texto:
        return "Jalisco"
    elif "PROVIDENCIA" in texto:
        return "Providencia"
    elif "PEGASO" in texto:
        if "MAQUINARIA" in texto:
            return "Maquinaria Pegaso"
        return "Planta Asfalto Pegaso"
    return "No aplica"

# Extract from XML
def parse_xml(xml_path):
    try:
        tree = ET.parse(xml_path)
        root = tree.getroot()
        namespaces = dict([node for _, node in ET.iterparse(xml_path, events=['start-ns'])])
        cfdi_ns = namespaces.get('cfdi', 'http://www.sat.gob.mx/cfd/4')
        if not cfdi_ns: cfdi_ns = 'http://www.sat.gob.mx/cfd/3'
        ns_map = {'cfdi': cfdi_ns}
        
        folio = root.get('Folio', '')
        fecha = root.get('Fecha', '').split('T')[0]
        
        proveedor = ""
        emisor = root.find('cfdi:Emisor', ns_map)
        if emisor is not None:
            proveedor = emisor.get('Nombre', '')
            
        total = float(root.get('Total', 0))
        subtotal = float(root.get('SubTotal', 0))
        iva = total - subtotal # Aprox
        
        is_gasolina = False
        litros = 0.0
        precio = 0.0
        conceptos = root.find('cfdi:Conceptos', ns_map)
        if conceptos is not None:
            for concepto in conceptos.findall('cfdi:Concepto', ns_map):
                desc = str(concepto.get('Descripcion', '')).upper()
                if "GASOLINA" in desc or "MAGNA" in desc or "PREMIUM" in desc:
                    is_gasolina = True
                litros += float(concepto.get('Cantidad', 0))
                precio = float(concepto.get('ValorUnitario', 0))
                
        tipo = "Gasolina" if is_gasolina else "Diésel"
        return {
            'folio': folio, 'fecha': fecha, 'proveedor': proveedor,
            'litros': litros, 'precio': precio, 'importe': subtotal,
            'iva': iva, 'total': total, 'tipo': tipo
        }
    except Exception as e:
        return None

# Get Destination strictly from PDF
def get_destino_from_pdf(pdf_path):
    try:
        text = ""
        with open(pdf_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() + "\n"
        return normalizar_destino(text)
    except:
        return "No aplica"


def main():
    if not os.path.exists(base_dir):
        print(f"La carpeta {base_dir} no existe.")
        return
        
    archivos = os.listdir(base_dir)
    file_map = {}
    for f in archivos:
        if f.lower().endswith('.xml') or f.lower().endswith('.pdf'):
            bname = os.path.splitext(f)[0].lower()
            if bname not in file_map: file_map[bname] = {}
            if f.lower().endswith('.xml'): file_map[bname]['xml'] = f
            if f.lower().endswith('.pdf'): file_map[bname]['pdf'] = f
            
    facturas_extraidas = []
    
    for bname, exts in file_map.items():
        data = None
        destino = "No aplica"
        
        # Read XML if exists
        if 'xml' in exts:
            data = parse_xml(os.path.join(base_dir, exts['xml']))
        
        # Read PDF
        if 'pdf' in exts:
            pdf_path = os.path.join(base_dir, exts['pdf'])
            destino = get_destino_from_pdf(pdf_path)
            
        if not data: continue
        
        data['destino'] = destino
        facturas_extraidas.append(data)
        
    print(f"Se extrajeron {len(facturas_extraidas)} facturas de la semana 28.")
    
    # Anexar al Excel
    print(f"Anexando {len(facturas_extraidas)} facturas de Diesel al Excel...")
    try:
        wb = openpyxl.load_workbook(excel_path)
        sheet = wb["BD_FACTURAS"]
        
        # Verificar duplicados en Excel para no volver a agregarlas si ya se guardó alguna
        folios_existentes = set()
        for row in range(2, sheet.max_row + 1):
            val = sheet.cell(row=row, column=2).value
            if val:
                folios_existentes.add(str(val).strip())
                
        agregados = 0
        for d in facturas_extraidas:
            if str(d['folio']).strip() in folios_existentes:
                continue
                
            row = sheet.max_row + 1
            # Col 2 (B): FOLIO
            sheet.cell(row=row, column=2).value = d['folio']
            # Col 3 (C): FECHA
            sheet.cell(row=row, column=3).value = d['fecha']
            # Col 5 (E): PROVEEDOR
            sheet.cell(row=row, column=5).value = d['proveedor']
            # Col 6 (F): PUNTO DE CARGA
            sheet.cell(row=row, column=6).value = d['destino']
            # Col 7 (G): LITROS
            sheet.cell(row=row, column=7).value = d['litros']
            # Col 8 (H): PRECIO
            sheet.cell(row=row, column=8).value = d['precio']
            # Col 9 (I): IMPORTE
            sheet.cell(row=row, column=9).value = d['importe']
            # Col 10 (J): IVA
            sheet.cell(row=row, column=10).value = d['iva']
            # Col 11 (K): TOTAL
            sheet.cell(row=row, column=11).value = d['total']
            # Col 12 (L): TIPO
            sheet.cell(row=row, column=12).value = "Diésel"
            
            # Formulas
            sheet.cell(row=row, column=4).value = f'=IF(C{row}="","", "Semana "&WEEKNUM(C{row},2))'
            sheet.cell(row=row, column=1).value = f'=IF(C{row}="","", "FA" & "-" & IFERROR(VLOOKUP(F{row},CATALOGOS!$A$2:$B$50,2,0),"XX") & "-" & WEEKNUM(C{row},2) & "-" & TEXT(ROW()-1,"000"))'
            agregados += 1
            
        if agregados > 0:
            wb.save(excel_path)
            print(f"Excel actualizado exitosamente. Se agregaron {agregados} nuevas facturas.")
        else:
            print("No se encontraron nuevas facturas para agregar (todas ya existían en el Excel).")
        
    except Exception as e:
        print(f"Error actualizando Excel: {e}")

if __name__ == "__main__":
    main()
