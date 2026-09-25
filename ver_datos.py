import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')
ws = wb['Captura_Facturas_Mobil']
print('Encabezados:')
for c in range(1, 16): print(f' Col{c}: {ws.cell(row=4, column=c).value}')
print('')
print('Datos de filas:')
for r in range(5, ws.max_row + 1):
    factura = ws.cell(row=r, column=1).value
    fecha = ws.cell(row=r, column=2).value
    destino = ws.cell(row=r, column=4).value
    lts = ws.cell(row=r, column=5).value
    semana_n = ws.cell(row=r, column=14).value
    if factura or fecha:
        print(f'R{r}: {factura} | {str(fecha)[:10]} | {destino} | Lts:{lts} | ColN:{semana_n}')

