import openpyxl

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)

def parse_sheet_semana(ws, sem_num):
    # Tabla 1: JDJ
    personas = []
    # Find JDJ table
    for r in range(4, 44):
        resp = ws.cell(r, 2).value
        obra = ws.cell(r, 3).value
        unidad = ws.cell(r, 4).value
        placas = ws.cell(r, 5).value
        auth = ws.cell(r, 6).value
        
        # Clean strings
        resp = str(resp).strip() if resp is not None else ''
        obra = str(obra).strip() if obra is not None else ''
        unidad = str(unidad).strip() if unidad is not None else ''
        placas = str(placas).strip().upper() if placas is not None else ''
        if placas in ['PLACAS', 'S/P', 'NONE']: placas = ''
        
        try:
            auth_val = float(auth or 0)
        except:
            auth_val = 0.0
            
        if resp or unidad or placas:
            if 'TOTAL' not in resp.upper():
                personas.append({
                    'empresa': 'JDJ',
                    'num_renglon': len(personas) + 1,
                    'responsable': resp or unidad,
                    'centro_trabajo': obra or 'General',
                    'unidad_equipo': unidad or 'Vehículo',
                    'placas': placas,
                    'importe_semanal': auth_val
                })
                
    # Tabla 2: TRD
    for r in range(50, 58):
        resp = ws.cell(r, 2).value
        obra = ws.cell(r, 3).value
        unidad = ws.cell(r, 4).value
        placas = ws.cell(r, 5).value
        auth = ws.cell(r, 6).value
        
        resp = str(resp).strip() if resp is not None else ''
        obra = str(obra).strip() if obra is not None else ''
        unidad = str(unidad).strip() if unidad is not None else ''
        placas = str(placas).strip().upper() if placas is not None else ''
        if placas in ['PLACAS', 'S/P', 'NONE']: placas = ''
        
        try:
            auth_val = float(auth or 0)
        except:
            auth_val = 0.0
            
        if resp or unidad or placas:
            if 'TOTAL' not in resp.upper():
                personas.append({
                    'empresa': 'TRD',
                    'num_renglon': len(personas) + 1,
                    'responsable': resp or unidad,
                    'centro_trabajo': obra or 'General',
                    'unidad_equipo': unidad or 'Vehículo',
                    'placas': placas,
                    'importe_semanal': auth_val
                })
    return personas

s33_personas = parse_sheet_semana(wb['SEMANA 33'], 33)
print(f"Extracted {len(s33_personas)} personas for Semana 33:")
for p in s33_personas:
    print(f"  {p['num_renglon']:2d} | {p['empresa']} | {p['responsable'][:25]:25s} | {p['centro_trabajo'][:20]:20s} | {p['unidad_equipo'][:15]:15s} | {p['placas'][:8]:8s} | ${p['importe_semanal']:,.2f}")
