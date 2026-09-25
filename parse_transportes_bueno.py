import openpyxl
import pandas as pd

excel_path = r"c:\Users\JOSE\Desktop\Proyecto fenix\TRANSPORTES\TRANSPORTES BUENO.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['TRANSPORTES BUENO']

print("=== INSPECCIÓN COMPLETA DE FILAS EN TRANSPORTES BUENO ===")
rows = []
for r in range(3, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 42)]
    if any(v is not None for v in vals):
        rows.append((r, vals))

print(f"Total filas con algún dato: {len(rows)}")

units = []
current_unit = None

for r, v in rows:
    num = v[0]
    eco = v[1]
    unidad = v[2]
    tipo_desc = v[3]
    serie = v[4]
    
    # If num or eco is present, it's a primary unit record
    if num is not None or (eco is not None and str(eco).strip() != ""):
        current_unit = {
            'row': r,
            'num': num,
            'eco': str(eco).strip() if eco is not None else '',
            'unidad': str(unidad).strip() if unidad is not None else '',
            'descripcion': str(tipo_desc).strip() if tipo_desc is not None else '',
            'serie': str(serie).strip() if serie is not None else '',
            'factura_tipo': v[5],
            'factura_autenticidad': v[6],
            'factura_uuid': v[7],
            'factura_num': v[8],
            'factura_fecha': v[9],
            'factura_receptor': v[10],
            'factura_emisor': v[11],
            'placas_entidad': v[12],
            'tarjeta_circulacion': v[13],
            'placas': str(v[14]).strip() if v[14] is not None else '',
            'vigencia_placas': v[15],
            'pago_derechos_2025': v[16],
            'cotejo_estatal': v[17],
            'permiso_carga': v[18],
            'cotejo_federal': v[19],
            'fisico_mecanica': v[20],
            'cotejo_fisico_mecanica': v[21],
            'ambiental': v[22],
            'poliza_interna': v[23],
            'poliza_cis': v[24],
            'poliza_vencimiento': v[25],
            'poliza_pago': v[26],
            'poliza_flotilla': v[27],
            'poliza_propietario': v[28],
            'observaciones_1': v[29],
            'observaciones_2': v[30],
            'observaciones_3': v[31],
            'ubicacion_ct': [x for x in v[32:] if x is not None],
            'sub_rows': []
        }
        units.append(current_unit)
    else:
        # Secondary row (e.g. secondary plates/permits)
        if current_unit is not None:
            current_unit['sub_rows'].append({
                'row': r,
                'entidad': v[12],
                'tarjeta_circulacion': v[13],
                'placas': v[14],
                'vigencia': v[15],
                'permiso_carga': v[18],
                'vals': [x for x in v if x is not None]
            })

print(f"\nTotal unidades principales identificadas: {len(units)}")
print("\nLISTADO DE UNIDADES DETECTADAS:")
for u in units:
    sub_info = f" (+{len(u['sub_rows'])} fila placa secundaria)" if u['sub_rows'] else ""
    print(f"  #{u['num']} | Eco: {u['eco']:12s} | Tipo: {u['unidad']:25s} | Placas: {u['placas']:10s} | Serie: {u['serie'][:18]} | Desc: {u['descripcion'][:30]}{sub_info}")
    if u['sub_rows']:
        for sr in u['sub_rows']:
            print(f"      -> Placa extra: {sr.get('placas')} ({sr.get('entidad')}) | Tarjeta: {sr.get('tarjeta_circulacion')}")
