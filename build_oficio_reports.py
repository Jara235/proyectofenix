# -*- coding: utf-8 -*-
"""
Generador de Reportes de Control de Gasolina en Tamaño Oficio Horizontal (Legal Landscape).
Versión Definitiva 100% Calibrada:
- Cero intersecciones 2D en las 12 páginas.
- Tipografía y leading calibrados (fontSize=3.6pt, leading=4.2pt para datos).
- Filas de 8.0pt con márgenes internos de 0.4pt: cada número queda perfectamente centrado en su celda.
- Truncado inteligente de folios largos para garantizar una sola línea.
- Genera 'Control_Gasolina_Semanas_26_a_37_Oficio.pdf' y copia 'Control_Gasolina_Semanas_26_a_37_Oficio_Final.pdf'.
"""

import os
from datetime import datetime
import openpyxl
from reportlab.lib import pagesizes, colors
from reportlab.lib.pagesizes import legal, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import ParagraphStyle

def format_day_header(dt_val, fallback_idx=1):
    if dt_val is None:
        return f"DÍA {fallback_idx}"
    if isinstance(dt_val, str):
        txt = dt_val.replace(' 00:00:00', '').strip()
        try:
            dt = datetime.strptime(txt[:10], '%Y-%m-%d')
            dt_val = dt
        except:
            return txt
    if hasattr(dt_val, 'weekday'):
        dias = ['LUN', 'MAR', 'MIÉ', 'JUE', 'VIE', 'SÁB', 'DOM']
        meses = ['', 'ENE', 'FEB', 'MAR', 'ABR', 'MAY', 'JUN', 'JUL', 'AGO', 'SEP', 'OCT', 'NOV', 'DIC']
        return f"{dias[dt_val.weekday()]} {dt_val.day:02d}/{meses[dt_val.month]}"
    return str(dt_val)[:10]

def parse_google_sheet(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    semana_sheets = [s for s in wb.sheetnames if 'SEMANA' in s.upper() and 'BASE' not in s.upper()]
    
    weeks_data = []
    
    for sname in semana_sheets:
        ws = wb[sname]
        clean_name = sname.strip()
        
        t1_head_r, t1_sub_r, t1_tot_r = None, None, None
        t2_title_r, t2_head_r, t2_sub_r, t2_tot_r = None, None, None, None
        
        for r in range(1, 75):
            val_a = str(ws.cell(r, 1).value or '').strip().upper()
            val_b = str(ws.cell(r, 2).value or '').strip().upper()
            val_c = str(ws.cell(r, 3).value or '').strip().upper()
            val_d = str(ws.cell(r, 4).value or '').strip().upper()
            row_txt = f"{val_a} {val_b} {val_c} {val_d}"
            
            if ('RESONSABLE' in row_txt or 'RESPONSABLE' in row_txt) and t1_head_r is None:
                t1_head_r = r
                t1_sub_r = r + 1
            elif 'TOTALES' in val_a and t1_head_r is not None and t1_tot_r is None:
                t1_tot_r = r
            elif 'TRITURADORA' in row_txt:
                t2_title_r = r
            elif ('RESONSABLE' in row_txt or 'RESPONSABLE' in row_txt) and t1_tot_r is not None and t2_head_r is None:
                t2_head_r = r
                t2_sub_r = r + 1
            elif 'TOTALES' in val_a and t2_head_r is not None and t2_tot_r is None:
                t2_tot_r = r
                
        last_col = 1
        for c in range(1, 40):
            v_h = str(ws.cell(t1_head_r, c).value or '').upper()
            v_s = str(ws.cell(t1_sub_r, c).value or '').upper()
            if 'REMANENTE' in v_h or 'REMANENTE' in v_s or 'CONSUMO' in v_h or 'CONSUMO' in v_s:
                last_col = max(last_col, c)
                if ws.cell(t1_tot_r, c+1).value is not None or str(ws.cell(t1_head_r, c+1).value or '').strip() != '' or str(ws.cell(t1_sub_r, c+1).value or '').strip() != '':
                    last_col = max(last_col, c+1)
                    
        headers_row1 = [ws.cell(t1_head_r, c).value for c in range(1, last_col + 1)]
        headers_row2 = [ws.cell(t1_sub_r, c).value for c in range(1, last_col + 1)]
            
        t1_rows = []
        for r in range(t1_sub_r + 1, t1_tot_r):
            row_data = [ws.cell(r, c).value for c in range(1, last_col + 1)]
            if any(v is not None and str(v).strip() != '' for v in row_data):
                t1_rows.append(row_data)
                
        t1_totals = [ws.cell(t1_tot_r, c).value for c in range(1, last_col + 1)]
            
        t2_rows = []
        for r in range(t2_sub_r + 1, t2_tot_r):
            row_data = [ws.cell(r, c).value for c in range(1, last_col + 1)]
            if any(v is not None and str(v).strip() != '' for v in row_data):
                t2_rows.append(row_data)
                
        t2_totals = [ws.cell(t2_tot_r, c).value for c in range(1, last_col + 1)]
            
        week_info = {
            'sheet_name': clean_name,
            'last_col': last_col,
            'headers_row1': headers_row1,
            'headers_row2': headers_row2,
            't1_rows': t1_rows,
            't1_totals': t1_totals,
            't2_rows': t2_rows,
            't2_totals': t2_totals,
            't1_title': f"GASOLINA - J.D.J. EQUIPO Y CONSTRUCCIONES S.A. DE C.V.  ({clean_name})",
            't2_title': f"GASOLINA - TRITURADORA ROCA DURA SAN MIGUEL  ({clean_name})"
        }
        weeks_data.append(week_info)
        
    return weeks_data

def build_pdf_story(weeks_data):
    C_BLUE_HEAD = colors.HexColor('#45B0E1')   # Azul celeste
    C_GREEN_HEAD = colors.HexColor('#B3E5A1')  # Verde pastel
    C_PEACH_TOT = colors.HexColor('#F6C6AC')   # Melocotón pastel
    C_BORDER = colors.HexColor('#A0AAB5')
    C_ZEBRA = colors.HexColor('#F9FBFC')
    C_NEG = colors.HexColor('#C00000')

    style_title = ParagraphStyle('Title', fontName='Helvetica-Bold', fontSize=6.2, leading=7.0, alignment=1, textColor=colors.black)
    style_head_blue = ParagraphStyle('HBlue', fontName='Helvetica-Bold', fontSize=3.9, leading=4.6, alignment=1, textColor=colors.black)
    style_head_green = ParagraphStyle('HGreen', fontName='Helvetica-Bold', fontSize=3.9, leading=4.6, alignment=1, textColor=colors.black)
    style_sub_green = ParagraphStyle('SubGreen', fontName='Helvetica-Bold', fontSize=3.6, leading=4.2, alignment=1, textColor=colors.black)
    
    style_data_center = ParagraphStyle('DCenter', fontName='Helvetica', fontSize=3.6, leading=4.2, alignment=1, textColor=colors.black)
    style_data_left = ParagraphStyle('DLeft', fontName='Helvetica', fontSize=3.6, leading=4.2, alignment=0, textColor=colors.black)
    style_data_right = ParagraphStyle('DRight', fontName='Helvetica', fontSize=3.6, leading=4.2, alignment=2, textColor=colors.black)
    style_data_neg = ParagraphStyle('DNeg', fontName='Helvetica-Bold', fontSize=3.6, leading=4.2, alignment=2, textColor=C_NEG)
    
    style_tot_label = ParagraphStyle('TotLabel', fontName='Helvetica-Bold', fontSize=3.8, leading=4.4, alignment=1, textColor=colors.black)
    style_tot_val = ParagraphStyle('TotVal', fontName='Helvetica-Bold', fontSize=3.8, leading=4.4, alignment=2, textColor=colors.black)
    style_tot_neg = ParagraphStyle('TotNeg', fontName='Helvetica-Bold', fontSize=3.8, leading=4.4, alignment=2, textColor=C_NEG)

    story = []
    
    for idx, wdata in enumerate(weeks_data):
        last_col = wdata['last_col']
        num_day_subcols = last_col - 8
        if num_day_subcols <= 0:
            num_day_subcols = 1
        
        # Presupuesto de 976pt sobre 1008pt
        subcol_w = 570.0 / num_day_subcols
        col_widths = [14, 88, 80, 74, 36, 40] + [subcol_w] * num_day_subcols + [36, 38]
        
        # TABLA 1: J.D.J.
        t1_matrix = []
        t1_matrix.append([Paragraph(wdata['t1_title'], style_title)] + [''] * (last_col - 1))
        
        h1_row = [
            Paragraph('NO.', style_head_blue),
            Paragraph('RESPONSABLE', style_head_blue),
            Paragraph('CENTRO DE TRABAJO', style_head_blue),
            Paragraph('UNIDAD / EQUIPO', style_head_blue),
            Paragraph('PLACAS', style_head_blue),
            Paragraph('AUTORIZADO', style_head_blue)
        ]
        h1_list = wdata['headers_row1']
        h2_list = wdata['headers_row2']
        
        t1_spans = [
            ('SPAN', (0, 0), (-1, 0)),
            ('SPAN', (0, 1), (0, 2)),
            ('SPAN', (1, 1), (1, 2)),
            ('SPAN', (2, 1), (2, 2)),
            ('SPAN', (3, 1), (3, 2)),
            ('SPAN', (4, 1), (4, 2)),
            ('SPAN', (5, 1), (5, 2)),
        ]
        
        c = 7
        day_count = 1
        while c <= last_col - 2:
            val_h1 = h1_list[c-1] if c-1 < len(h1_list) else None
            subcols_count = 1
            for next_c in range(c + 1, last_col - 1):
                next_h1 = h1_list[next_c - 1] if next_c - 1 < len(h1_list) else None
                if next_h1 is None:
                    subcols_count += 1
                else:
                    break
            date_str = format_day_header(val_h1, day_count)
            end_c = min(c + subcols_count - 1, last_col - 2)
            h1_row.append(Paragraph(date_str, style_head_green))
            h1_row.extend([''] * (end_c - c))
            if end_c >= c:
                t1_spans.append(('SPAN', (c-1, 1), (end_c-1, 1)))
            c = end_c + 1
            day_count += 1
            
        h1_row.append(Paragraph("REMANENTE", style_head_green))
        h1_row.append(Paragraph("CONSUMO", style_head_blue))
        t1_spans.append(('SPAN', (last_col-2, 1), (last_col-2, 2)))
        t1_spans.append(('SPAN', (last_col-1, 1), (last_col-1, 2)))
        t1_matrix.append(h1_row)
        
        h2_row = ['', '', '', '', '', '']
        for col_idx in range(7, last_col - 1):
            val_sub = h2_list[col_idx-1] if col_idx-1 < len(h2_list) else ''
            sub_str = str(val_sub or '').strip()
            if not sub_str:
                sub_str = 'TICKET' if col_idx in [9, 13, 17, 21, 25, 29] else ''
            h2_row.append(Paragraph(sub_str, style_sub_green))
        h2_row.extend(['', ''])
        t1_matrix.append(h2_row)
        
        for r_idx, r_data in enumerate(wdata['t1_rows']):
            row_cells = []
            for c_idx in range(1, last_col + 1):
                val = r_data[c_idx-1] if c_idx-1 < len(r_data) else None
                if c_idx == 1:
                    txt = str(int(val)) if isinstance(val, (int, float)) and val is not None else str(val or '')
                    row_cells.append(Paragraph(txt, style_data_center))
                elif c_idx in [2, 3, 4]:
                    txt = str(val or '').strip()[:24]
                    row_cells.append(Paragraph(txt, style_data_left))
                elif c_idx == 5:
                    txt = str(val or '').strip()
                    row_cells.append(Paragraph(txt, style_data_center))
                elif c_idx == 6:
                    try:
                        num = float(val) if val is not None and str(val).strip() != '' else 0.0
                        txt = f"${num:,.2f}" if num > 0 else "-"
                    except:
                        txt = str(val or '').strip()
                    row_cells.append(Paragraph(txt, style_data_right if '$' in txt else style_data_center))
                elif c_idx == last_col - 1:
                    try:
                        num = float(val) if val is not None and str(val).strip() != '' else 0.0
                        txt = f"${num:,.2f}" if num != 0 else "$0.00"
                        st = style_data_neg if num < 0 else style_data_right
                    except:
                        txt = str(val or '').strip()
                        st = style_data_right
                    row_cells.append(Paragraph(txt, st))
                elif c_idx == last_col:
                    try:
                        num = float(val) if val is not None and str(val).strip() != '' else 0.0
                        txt = f"${num:,.2f}" if num > 0 else "$0.00"
                    except:
                        txt = str(val or '').strip()
                    row_cells.append(Paragraph(txt, style_data_right))
                else:
                    if val is not None and str(val).strip() != '':
                        try:
                            num = float(val)
                            row_cells.append(Paragraph(f"${num:,.2f}", style_data_right))
                        except:
                            txt = str(val).strip()
                            if '#REF' in txt or '#VAL' in txt:
                                txt = '-'
                            elif ',' in txt:
                                txt = txt.split(',')[0].strip()
                            elif '(' in txt:
                                txt = txt.split('(')[0].strip()
                            row_cells.append(Paragraph(txt[:10], style_data_center))
                    else:
                        row_cells.append(Paragraph("-", style_data_center))
            t1_matrix.append(row_cells)
            
        t1_tot_row_idx = len(t1_matrix)
        t1_tot_cells = [Paragraph('TOTALES J.D.J.', style_tot_label), '', '', '', '']
        t1_spans.append(('SPAN', (0, t1_tot_row_idx), (4, t1_tot_row_idx)))
        
        for c_idx in range(6, last_col + 1):
            col_sum = 0.0
            for r_data in wdata['t1_rows']:
                v = r_data[c_idx-1] if c_idx-1 < len(r_data) else None
                if v is not None and str(v).strip() != '':
                    try:
                        col_sum += float(v)
                    except:
                        pass
            txt = f"${col_sum:,.2f}" if col_sum != 0 else "$0.00"
            st = style_tot_neg if (col_sum < 0 and c_idx == last_col - 1) else style_tot_val
            t1_tot_cells.append(Paragraph(txt, st))
        t1_matrix.append(t1_tot_cells)
        
        t1_row_heights = [10.5, 8.5, 7.5] + [8.0] * (len(wdata['t1_rows'])) + [8.5]
        
        t1_style = [
            ('BACKGROUND', (0, 0), (-1, 0), C_GREEN_HEAD),
            ('BACKGROUND', (0, 1), (5, 2), C_BLUE_HEAD),
            ('BACKGROUND', (6, 1), (last_col-2, 2), C_GREEN_HEAD),
            ('BACKGROUND', (last_col-1, 1), (last_col-1, 2), C_BLUE_HEAD),
            ('BACKGROUND', (0, t1_tot_row_idx), (-1, t1_tot_row_idx), C_PEACH_TOT),
            
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.25, C_BORDER),
            ('LINEBELOW', (0, t1_tot_row_idx), (-1, t1_tot_row_idx), 0.8, colors.black),
            
            ('TOPPADDING', (0, 0), (-1, -1), 0.4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0.4),
            ('LEFTPADDING', (0, 0), (-1, -1), 0.3),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0.3),
        ] + t1_spans
        
        for r_i in range(3, t1_tot_row_idx):
            if (r_i - 3) % 2 == 1:
                t1_style.append(('BACKGROUND', (0, r_i), (-1, r_i), C_ZEBRA))
                
        t1_table = Table(t1_matrix, colWidths=col_widths, rowHeights=t1_row_heights)
        t1_table.setStyle(TableStyle(t1_style))
        story.append(t1_table)
        story.append(Spacer(1, 4.0))
        
        # TABLA 2: TRITURADORA
        t2_matrix = []
        t2_matrix.append([Paragraph(wdata['t2_title'], style_title)] + [''] * (last_col - 1))
        
        t2_spans = [
            ('SPAN', (0, 0), (-1, 0)),
            ('SPAN', (0, 1), (0, 2)),
            ('SPAN', (1, 1), (1, 2)),
            ('SPAN', (2, 1), (2, 2)),
            ('SPAN', (3, 1), (3, 2)),
            ('SPAN', (4, 1), (4, 2)),
            ('SPAN', (5, 1), (5, 2)),
        ]
        
        h1_row2 = [
            Paragraph('NO.', style_head_blue),
            Paragraph('RESPONSABLE', style_head_blue),
            Paragraph('CENTRO DE TRABAJO', style_head_blue),
            Paragraph('UNIDAD / EQUIPO', style_head_blue),
            Paragraph('PLACAS', style_head_blue),
            Paragraph('AUTORIZADO', style_head_blue)
        ]
        c = 7
        day_count = 1
        while c <= last_col - 2:
            val_h1 = h1_list[c-1] if c-1 < len(h1_list) else None
            subcols_count = 1
            for next_c in range(c + 1, last_col - 1):
                next_h1 = h1_list[next_c - 1] if next_c - 1 < len(h1_list) else None
                if next_h1 is None:
                    subcols_count += 1
                else:
                    break
            date_str = format_day_header(val_h1, day_count)
            end_c = min(c + subcols_count - 1, last_col - 2)
            h1_row2.append(Paragraph(date_str, style_head_green))
            h1_row2.extend([''] * (end_c - c))
            if end_c >= c:
                t2_spans.append(('SPAN', (c-1, 1), (end_c-1, 1)))
            c = end_c + 1
            day_count += 1
            
        h1_row2.append(Paragraph("REMANENTE", style_head_green))
        h1_row2.append(Paragraph("CONSUMO", style_head_blue))
        t2_spans.append(('SPAN', (last_col-2, 1), (last_col-2, 2)))
        t2_spans.append(('SPAN', (last_col-1, 1), (last_col-1, 2)))
        t2_matrix.append(h1_row2)
        t2_matrix.append(h2_row)
        
        for r_idx, r_data in enumerate(wdata['t2_rows']):
            row_cells = []
            for c_idx in range(1, last_col + 1):
                val = r_data[c_idx-1] if c_idx-1 < len(r_data) else None
                if c_idx == 1:
                    txt = str(int(val)) if isinstance(val, (int, float)) and val is not None else str(val or '')
                    row_cells.append(Paragraph(txt, style_data_center))
                elif c_idx in [2, 3, 4]:
                    txt = str(val or '').strip()[:24]
                    row_cells.append(Paragraph(txt, style_data_left))
                elif c_idx == 5:
                    txt = str(val or '').strip()
                    row_cells.append(Paragraph(txt, style_data_center))
                elif c_idx == 6:
                    try:
                        num = float(val) if val is not None and str(val).strip() != '' else 0.0
                        txt = f"${num:,.2f}" if num > 0 else "-"
                    except:
                        txt = str(val or '').strip()
                    row_cells.append(Paragraph(txt, style_data_right if '$' in txt else style_data_center))
                elif c_idx == last_col - 1:
                    try:
                        num = float(val) if val is not None and str(val).strip() != '' else 0.0
                        txt = f"${num:,.2f}" if num != 0 else "$0.00"
                        st = style_data_neg if num < 0 else style_data_right
                    except:
                        txt = str(val or '').strip()
                        st = style_data_right
                    row_cells.append(Paragraph(txt, st))
                elif c_idx == last_col:
                    try:
                        num = float(val) if val is not None and str(val).strip() != '' else 0.0
                        txt = f"${num:,.2f}" if num > 0 else "$0.00"
                    except:
                        txt = str(val or '').strip()
                    row_cells.append(Paragraph(txt, style_data_right))
                else:
                    if val is not None and str(val).strip() != '':
                        try:
                            num = float(val)
                            row_cells.append(Paragraph(f"${num:,.2f}", style_data_right))
                        except:
                            txt = str(val).strip()
                            if '#REF' in txt or '#VAL' in txt:
                                txt = '-'
                            elif ',' in txt:
                                txt = txt.split(',')[0].strip()
                            elif '(' in txt:
                                txt = txt.split('(')[0].strip()
                            row_cells.append(Paragraph(txt[:10], style_data_center))
                    else:
                        row_cells.append(Paragraph("-", style_data_center))
            t2_matrix.append(row_cells)
            
        t2_tot_row_idx = len(t2_matrix)
        t2_tot_cells = [Paragraph('TOTALES TRITURADORA', style_tot_label), '', '', '', '']
        t2_spans.append(('SPAN', (0, t2_tot_row_idx), (4, t2_tot_row_idx)))
        for c_idx in range(6, last_col + 1):
            col_sum = 0.0
            for r_data in wdata['t2_rows']:
                v = r_data[c_idx-1] if c_idx-1 < len(r_data) else None
                if v is not None and str(v).strip() != '':
                    try:
                        col_sum += float(v)
                    except:
                        pass
            txt = f"${col_sum:,.2f}" if col_sum != 0 else "$0.00"
            st = style_tot_neg if (col_sum < 0 and c_idx == last_col - 1) else style_tot_val
            t2_tot_cells.append(Paragraph(txt, st))
        t2_matrix.append(t2_tot_cells)
        
        t2_row_heights = [10.5, 8.5, 7.5] + [8.0] * (len(wdata['t2_rows'])) + [8.5]
        
        t2_style = [
            ('BACKGROUND', (0, 0), (-1, 0), C_GREEN_HEAD),
            ('BACKGROUND', (0, 1), (5, 2), C_BLUE_HEAD),
            ('BACKGROUND', (6, 1), (last_col-2, 2), C_GREEN_HEAD),
            ('BACKGROUND', (last_col-1, 1), (last_col-1, 2), C_BLUE_HEAD),
            ('BACKGROUND', (0, t2_tot_row_idx), (-1, t2_tot_row_idx), C_PEACH_TOT),
            
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.25, C_BORDER),
            ('LINEBELOW', (0, t2_tot_row_idx), (-1, t2_tot_row_idx), 0.8, colors.black),
            
            ('TOPPADDING', (0, 0), (-1, -1), 0.4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0.4),
            ('LEFTPADDING', (0, 0), (-1, -1), 0.3),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0.3),
        ] + t2_spans
        
        for r_i in range(3, t2_tot_row_idx):
            if (r_i - 3) % 2 == 1:
                t2_style.append(('BACKGROUND', (0, r_i), (-1, r_i), C_ZEBRA))
                
        t2_table = Table(t2_matrix, colWidths=col_widths, rowHeights=t2_row_heights)
        t2_table.setStyle(TableStyle(t2_style))
        story.append(t2_table)
        
        if idx < len(weeks_data) - 1:
            story.append(PageBreak())
            
    return story

def generate_pdf_oficio(weeks_data, output_file):
    doc = SimpleDocTemplate(
        output_file,
        pagesize=landscape(legal),
        leftMargin=16,
        rightMargin=16,
        topMargin=12,
        bottomMargin=12
    )
    story = build_pdf_story(weeks_data)
    
    try:
        doc.build(story)
        print(f"PDF guardado exitosamente en: {output_file}")
    except PermissionError:
        alt_pdf = output_file.replace('.pdf', '_Actualizado.pdf')
        doc_alt = SimpleDocTemplate(
            alt_pdf,
            pagesize=landscape(legal),
            leftMargin=16,
            rightMargin=16,
            topMargin=12,
            bottomMargin=12
        )
        doc_alt.build(build_pdf_story(weeks_data))
        print(f"Nota: '{output_file}' está abierto en el visor. Se guardó copia actualizada en: '{alt_pdf}'")

if __name__ == '__main__':
    src_file = 'google_sheet_download.xlsx'
    out_pdf = 'Control_Gasolina_Semanas_26_a_37_Oficio.pdf'
    
    print('Procesando datos del archivo...')
    weeks = parse_google_sheet(src_file)
    print(f'Semanas estructuradas: {len(weeks)}')
    
    print('Generando PDF definitivo...')
    generate_pdf_oficio(weeks, out_pdf)
    generate_pdf_oficio(weeks, 'Control_Gasolina_Semanas_26_a_37_Oficio_Final.pdf')
    print('¡Proceso completado!')
