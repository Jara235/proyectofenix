import openpyxl

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)

print("Searching across all sheets for 10922 / 4184 / 175.07 / A10922:")
found = False
for sname in wb.sheetnames:
    ws = wb[sname]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is not None:
                sv = str(v).lower()
                if '10922' in sv or '4184' in sv or '175.07' in sv:
                    print(f"FOUND in Sheet '{sname}' [Row {r}, Col {c}]: {v}")
                    row_vals = {col: ws.cell(r, col).value for col in range(1, ws.max_column + 1) if ws.cell(r, col).value is not None}
                    print(f"  Row {r}: {row_vals}")
                    found = True

if not found:
    print("Not found with direct text. Let's check formulas in the raw workbook...")
    wb_raw = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=False)
    for sname in wb_raw.sheetnames:
        ws = wb_raw[sname]
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is not None:
                    sv = str(v).lower()
                    if '10922' in sv or '4184' in sv or '175.07' in sv:
                        print(f"FOUND FORMULA in Sheet '{sname}' [Row {r}, Col {c}]: {v}")
                        found = True

print("\nDone searching.")
