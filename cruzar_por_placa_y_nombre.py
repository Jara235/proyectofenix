import openpyxl
import psycopg2
from psycopg2.extras import DictCursor
import re

def clean_str(s):
    if not s:
        return ""
    return str(s).strip().upper()

def compare_by_placa_and_name():
    excel_path = r'c:\Users\JOSE\Desktop\Proyecto fenix\CONCILIACIÓN COMBUSTIBLE GASOLINERAS.xlsx'
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['SEMANA 30']

    # 1. Parse Excel Charges from SEMANA 30
    excel_charges = []
    dates = [(7, '2026-07-20'), (10, '2026-07-21'), (13, '2026-07-22'), (16, '2026-07-23'), (19, '2026-07-24'), (22, '2026-07-25'), (25, '2026-07-26')]

    for r in range(4, ws.max_row + 1):
        resp = ws.cell(r, 2).value
        centro = ws.cell(r, 3).value
        unidad = ws.cell(r, 4).value
        placa = ws.cell(r, 5).value

        resp_clean = clean_str(resp)
        centro_clean = clean_str(centro)
        und_clean = clean_str(unidad)
        plc_clean = clean_str(placa)

        if plc_clean in ('NONE', 'S/P', 'N/A', 'PLACAS', 'AUTORIZADO') or len(plc_clean) > 12:
            plc_clean = ''

        if und_clean in ('UNIDAD / EQUIPO', 'TOTAL', 'CONSUMO', 'NONE', 'N/A') or 'FACTURA' in und_clean or 'SEMANA' in und_clean:
            und_clean = ''

        for col_start, d_str in dates:
            for prov_idx, prov_name in enumerate(['LEVET', 'MOBILE', 'SI VALE']):
                val = ws.cell(r, col_start + prov_idx).value
                if val is not None:
                    try:
                        amt = float(val)
                        if amt > 0:
                            excel_charges.append({
                                'ex_row': r,
                                'fecha': d_str,
                                'proveedor': prov_name,
                                'responsable': resp_clean,
                                'centro_trabajo': centro_clean,
                                'unidad': und_clean,
                                'placa': plc_clean,
                                'importe': amt
                            })
                    except (ValueError, TypeError):
                        pass

    print(f"Total charges extracted from Excel SEMANA 30: {len(excel_charges)}")

    # 2. Fetch DB records for Semana 30
    conn = psycopg2.connect(dbname='fenix_db', user='postgres', password='Bupito*268', host='localhost')
    cur = conn.cursor(cursor_factory=DictCursor)
    cur.execute('''
        SELECT id, folio_conciliacion, fecha, origen, obra_destino, vehiculo, placa, conductor, litros, costo_por_litro, importe_total, observaciones
        FROM gasolina.consumos
        WHERE semana::text = '30'
        ORDER BY fecha, id
    ''')
    db_rows = [dict(r) for r in cur.fetchall()]
    conn.close()

    print(f"Total DB records for Semana 30: {len(db_rows)}")

    # 3. Match DB records to Excel charges by PLACA or RESPONSABLE / CONDUCTOR + AMOUNT / DATE
    used_excel_indices = set()
    matches = []
    unmatched_db = []

    for idx, r in enumerate(db_rows, 1):
        r_id = r['id']
        folio = r['folio_conciliacion'] or f"ID-{r_id}"
        fecha = str(r['fecha'] or '')
        plc = clean_str(r['placa'])
        veh = clean_str(r['vehiculo'])
        cond = clean_str(r['conductor'])
        obra = clean_str(r['obra_destino'])
        imp = float(r['importe_total'] or 0)

        best_match_idx = None
        best_match_score = 0

        for ex_i, c in enumerate(excel_charges):
            if ex_i in used_excel_indices:
                continue

            score = 0

            # Date match
            if c['fecha'] == fecha:
                score += 10

            # Amount match (within $5)
            if abs(c['importe'] - imp) < 5.0:
                score += 30
            elif abs(c['importe'] - imp) < 50.0:
                score += 10

            # Placa match
            if plc and c['placa'] and (plc in c['placa'] or c['placa'] in plc):
                score += 40

            # Conductor / Responsable match
            if cond and c['responsable']:
                # Token overlap
                cond_tokens = set(cond.split())
                resp_tokens = set(c['responsable'].split())
                overlap = cond_tokens & resp_tokens
                if overlap:
                    score += 20 * len(overlap)

            # Vehicle match
            if veh and c['unidad'] and (veh in c['unidad'] or c['unidad'] in veh):
                score += 15

            if score > best_match_score and score >= 30:
                best_match_score = score
                best_match_idx = ex_i

        if best_match_idx is not None:
            used_excel_indices.add(best_match_idx)
            ex_item = excel_charges[best_match_idx]
            matches.append({
                'db_num': idx,
                'db_folio': folio,
                'db_fecha': fecha,
                'db_placa': r['placa'] or 'S/P',
                'db_vehiculo': r['vehiculo'] or 'N/A',
                'db_conductor': r['conductor'] or 'No Identificado',
                'db_obra': r['obra_destino'] or 'SIN OBRA',
                'db_importe': imp,
                'ex_responsable': ex_item['responsable'],
                'ex_centro': ex_item['centro_trabajo'],
                'ex_unidad': ex_item['unidad'],
                'ex_placa': ex_item['placa'],
                'ex_importe': ex_item['importe'],
                'score': best_match_score
            })
        else:
            unmatched_db.append({
                'db_num': idx,
                'db_folio': folio,
                'db_fecha': fecha,
                'db_placa': r['placa'] or 'S/P',
                'db_vehiculo': r['vehiculo'] or 'N/A',
                'db_conductor': r['conductor'] or 'No Identificado',
                'db_obra': r['obra_destino'] or 'SIN OBRA',
                'db_importe': imp,
                'db_litros': float(r['litros'] or 0)
            })

    print(f"\nMatches found (ignoring obra): {len(matches)} de {len(db_rows)}")
    print(f"Registros en BD verdaderamente sin coincidencia en Excel: {len(unmatched_db)}")

    # Build Markdown Artifact
    md = []
    md.append("# ⛽ Cruce Directo por Placa y Nombre/Responsable (Semana 30)")
    md.append("## Ignorando la 'Obra' de Captura para evitar sesgo por error de ingreso en el portal\n")

    md.append(f"- **Total Registros en BD:** `{len(db_rows)} cargas`")
    md.append(f"- **Coincidencias encontradas por Placa / Conductor / Fecha / Monto:** `{len(matches)} cargas` (**{len(matches)/len(db_rows)*100:.1f}%**)")
    md.append(f"- **Cargas verdaderamente no encontradas en el Excel:** `{len(unmatched_db)} cargas`\n")

    md.append("### 1. Registros de BD Coincidentes con la Pestaña `SEMANA 30` del Excel")
    md.append("| # BD | Folio | Fecha | Placa BD ➡️ Excel | Conductor BD ➡️ Responsable Excel | Obra Capturada ➡️ Centro Trabajo Excel | Monto BD | Monto Excel | Coincidencia |")
    md.append("|-----:|-------|-------|-------------------|----------------------------------|---------------------------------------|---------:|------------:|:-------------|")

    for m in matches:
        plc_str = f"`{m['db_placa']}`" if m['db_placa'] == m['ex_placa'] else f"`{m['db_placa']}` ➡️ **`{m['ex_placa']}`**"
        cond_str = m['db_conductor'] if clean_str(m['db_conductor']) == m['ex_responsable'] else f"{m['db_conductor']} ➡️ **{m['ex_responsable']}**"
        obra_str = f"{m['db_obra']} ➡️ **{m['ex_centro']}**" if m['ex_centro'] else m['db_obra']

        md.append(f"| {m['db_num']} | `{m['db_folio']}` | {m['db_fecha']} | {plc_str} | {cond_str} | {obra_str} | ${m['db_importe']:,.2f} | ${m['ex_importe']:,.2f} | ✅ Score {m['score']} |")

    md.append("\n---\n")

    md.append("### 2. Cargas en BD que NO tienen Coincidencia en la Pestaña `SEMANA 30` del Excel")
    if unmatched_db:
        md.append("| # BD | Folio | Fecha | Obra Capturada | Vehículo | Placa | Conductor | Litros | Importe BD ($) |")
        md.append("|-----:|-------|-------|----------------|----------|-------|-----------|-------:|---------------:|")
        for u in unmatched_db:
            md.append(f"| {u['db_num']} | `{u['db_folio']}` | {u['db_fecha']} | {u['db_obra']} | {u['db_vehiculo']} | **`{u['db_placa']}`** | {u['db_conductor']} | {u['db_litros']:,.2f} L | **${u['db_importe']:,.2f}** |")
    else:
        md.append("¡Todas las cargas de la Base de Datos coinciden con el Excel!")

    art_path = r'C:\Users\JOSE\.gemini\antigravity\brain\f3511922-47fb-4fdf-9cb3-6eae6dade26d\cruce_semana30_placa_nombre.md'
    with open(art_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(md))

    print(f"\nReport written to: {art_path}")

if __name__ == '__main__':
    compare_by_placa_and_name()
