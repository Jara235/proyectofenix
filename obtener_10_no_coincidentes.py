import openpyxl
import psycopg2
from psycopg2.extras import DictCursor

def find_unmatched():
    wb = openpyxl.load_workbook(r'c:\Users\JOSE\Desktop\Proyecto fenix\CONCILIACIÓN COMBUSTIBLE GASOLINERAS.xlsx', data_only=True)
    ws = wb['SEMANA 30']

    excel_charges = []
    dates = [(7, '2026-07-20'), (10, '2026-07-21'), (13, '2026-07-22'), (16, '2026-07-23'), (19, '2026-07-24'), (22, '2026-07-25'), (25, '2026-07-26')]

    for r in range(4, ws.max_row + 1):
        placa = ws.cell(r, 5).value
        unidad = ws.cell(r, 4).value
        resp = ws.cell(r, 2).value
        plc = str(placa).strip().upper() if placa else None
        und = str(unidad).strip().upper() if unidad else None
        if plc in ('NONE', 'S/P', '', 'N/A', 'PLACAS', 'AUTORIZADO') or (plc and len(plc) > 12):
            plc = None

        for col_start, d_str in dates:
            for source_name, val in [('LEVET', ws.cell(r, col_start).value), ('MOBILE', ws.cell(r, col_start + 1).value), ('SI VALE', ws.cell(r, col_start + 2).value)]:
                if val is not None:
                    try:
                        amt = float(val)
                        if amt > 0:
                            excel_charges.append({'fecha': d_str, 'placa': plc, 'unidad': und, 'responsable': resp, 'importe': amt})
                    except (ValueError, TypeError):
                        pass

    conn = psycopg2.connect(dbname='fenix_db', user='postgres', password='Bupito*268', host='localhost')
    cur = conn.cursor(cursor_factory=DictCursor)
    cur.execute('''
        SELECT id, folio_conciliacion, fecha, obra_destino, vehiculo, placa, conductor, litros, costo_por_litro, importe_total, observaciones
        FROM gasolina.consumos
        WHERE semana::text = '30'
        ORDER BY fecha, id
    ''')
    db_rows = [dict(r) for r in cur.fetchall()]
    conn.close()

    no_match_db = []
    used_excel_indices = set()

    for idx, r in enumerate(db_rows, 1):
        folio = r['folio_conciliacion'] or f"ID-{r['id']}"
        fecha = str(r['fecha'] or '')
        plc = (r['placa'] or 'S/P').strip().upper()
        veh = (r['vehiculo'] or 'N/A').strip().upper()
        cond = r['conductor'] or 'No Identificado'
        obra = r['obra_destino'] or 'SIN OBRA'
        imp = float(r['importe_total'] or 0)

        matched_idx = None
        for ex_i, c in enumerate(excel_charges):
            if ex_i in used_excel_indices:
                continue
            if c['fecha'] == fecha:
                # Check amount match (within $15)
                if abs(c['importe'] - imp) < 15.0:
                    if (c['placa'] and c['placa'] in plc) or (c['unidad'] and c['unidad'] in veh) or (plc == 'S/P') or (c['placa'] is None):
                        matched_idx = ex_i
                        break

        if matched_idx is not None:
            used_excel_indices.add(matched_idx)
        else:
            no_match_db.append({
                'num': idx,
                'folio': folio,
                'fecha': fecha,
                'obra': obra,
                'vehiculo': r['vehiculo'] or 'N/A',
                'placa': plc,
                'conductor': cond,
                'litros': float(r['litros'] or 0),
                'importe': imp,
                'observaciones': r['observaciones'] or 'Sin observaciones'
            })

    print(f"Total cargas en BD sin registro coincidente en pestaña Excel SEMANA 30: {len(no_match_db)}")
    for m in no_match_db:
        print(f"#{m['num']:2d} | Folio: {m['folio']:<16} | Fecha: {m['fecha']} | Obra: {m['obra']:<18} | Veh: {m['vehiculo']:<22} | Placa: {m['placa']:<10} | Cond: {m['conductor']:<20} | Importe: ${m['importe']:,.2f}")

    # Write to a clean markdown report artifact
    md = []
    md.append("# ⛽ Cargas Registradas en Base de Datos No Encontradas en Excel (Semana 30)\n")
    md.append(f"Se identificaron **{len(no_match_db)} cargas** capturadas en el sistema de captura de gasolina para la **Semana 30** que **no tienen una fila coincidente de consumo diario en la pestaña `SEMANA 30`** del archivo `CONCILIACIÓN COMBUSTIBLE GASOLINERAS.xlsx`:\n")

    md.append("| # | Folio | Fecha | Obra Destino | Vehículo | Placa | Conductor | Litros | Importe ($) | Observaciones |")
    md.append("|--:|-------|-------|--------------|----------|-------|-----------|-------:|------------:|---------------|")

    for m in no_match_db:
        md.append(f"| {m['num']} | `{m['folio']}` | {m['fecha']} | {m['obra']} | {m['vehiculo']} | **`{m['placa']}`** | {m['conductor']} | {m['litros']:,.2f} L | **${m['importe']:,.2f}** | {m['obs'] if 'obs' in m else 'Sin observaciones'} |")

    art_path = r'C:\Users\JOSE\.gemini\antigravity\brain\f3511922-47fb-4fdf-9cb3-6eae6dade26d\cargas_sin_coincidencia_excel_s30.md'
    with open(art_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(md))

    print(f"Report artifact written to: {art_path}")

if __name__ == '__main__':
    find_unmatched()
