import openpyxl

def main():
    path = 'c:/Users/JOSE/Desktop/Proyecto fenix/MAESTRO_CONTROL_DIESEL_NUEVO.xlsx'
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    print('Sheets:', wb.sheetnames)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for r, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True)):
            if sum(1 for c in row if c is not None and str(c).strip() != '') > 3:
                headers = [str(c) if c else '' for c in row]
                print(f'\n--- Sheet: {sheet} (Row {r+1}) ---')
                print(headers)
                break

if __name__ == '__main__':
    main()
