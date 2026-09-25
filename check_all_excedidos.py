import openpyxl

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)

for w in ['SEMANA 26', 'SEMANA 27', 'SEMANA 28', 'SEMANA 29', 'SEMANA 30', 'SEMANA 31', 'SEMANA 32', 'SEMANA 33']:
    if w in wb.sheetnames:
        ws = wb[w]
        print(f"\n==================== {w} ====================")
        excedidos_rows = []
        for r in range(1, ws.max_row + 1):
            for c in range(1, 10):
                v = str(ws.cell(r, c).value or '')
                if 'CONSUMO EXCEDIDO' in v.upper():
                    print(f"Found Excedidos header at Row {r}")
                    # Read subsequent rows until empty or total
                    for r_sub in range(r+2, r+25):
                        row_v = [ws.cell(r_sub, col).value for col in range(1, 10)]
                        if not any(row_v) or 'TOTAL' in str(row_v):
                            if 'TOTAL' in str(row_v):
                                print(f"  {r_sub}: {row_v}")
                            break
                        excedidos_rows.append(row_v)
                        print(f"  {r_sub}: {row_v}")
                    break
