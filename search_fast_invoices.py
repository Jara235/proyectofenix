import openpyxl

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)
print("Sheet names:", wb.sheetnames)

targets = ['10881', '10922', '10929', '10943', 'A10881', 'A10922', 'A10929', 'A10943', '4184']

# Search specifically in SEMANA 33 first, and then other sheets
for sname in ['SEMANA 33', 'SEMANA 32', 'SEMANA 31', 'SEMANA 30', 'Estado de Cuenta']:
    if sname not in wb.sheetnames:
        continue
    ws = wb[sname]
    print(f"\n--- Searching in {sname} ---")
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        row_str = " | ".join([str(v) for v in row if v is not None])
        for t in targets:
            if t.lower() in row_str.lower():
                print(f"Row {r_idx}: {row_str}")
                break
