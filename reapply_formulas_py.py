import openpyxl

excel_path = r"c:\Users\JOSE\Desktop\Proyecto fenix\MAESTRO_CONTROL_DIESEL_NUEVO.xlsx"

try:
    print("Loading workbook...")
    wb = openpyxl.load_workbook(excel_path)
    
    print("Writing BD_DIESEL formulas...")
    sheet_bd = wb["BD_DIESEL"]
    for row in range(2, 1000):
        sheet_bd.cell(row=row, column=3).value = f'=IF(B{row}="","", "Semana "&WEEKNUM(B{row},2))'
        sheet_bd.cell(row=row, column=1).value = f'=IF(B{row}="","", IFERROR(VLOOKUP(E{row},CATALOGOS!$J$2:$K$50,2,0),"XX") & "-" & IFERROR(VLOOKUP(F{row},CATALOGOS!$A$2:$B$50,2,0),"XX") & "-" & WEEKNUM(B{row},2) & "-" & TEXT(ROW()-1,"000"))'

    if "BD_SOLICITUDES" in wb.sheetnames:
        print("Writing BD_SOLICITUDES formulas...")
        sheet_sol = wb["BD_SOLICITUDES"]
        for row in range(2, 1000):
            sheet_sol.cell(row=row, column=3).value = f'=IF(B{row}="","", "Semana "&WEEKNUM(B{row},2))'
            sheet_sol.cell(row=row, column=1).value = f'=IF(B{row}="","", "SO" & "-" & IFERROR(VLOOKUP(F{row},CATALOGOS!$A$2:$B$50,2,0),"XX") & "-" & WEEKNUM(B{row},2) & "-" & TEXT(ROW()-1,"000"))'

    if "BD_FACTURAS" in wb.sheetnames:
        print("Writing BD_FACTURAS formulas...")
        sheet_fac = wb["BD_FACTURAS"]
        for row in range(2, 1000):
            sheet_fac.cell(row=row, column=4).value = f'=IF(C{row}="","", "Semana "&WEEKNUM(C{row},2))'
            sheet_fac.cell(row=row, column=1).value = f'=IF(C{row}="","", "FA" & "-" & IFERROR(VLOOKUP(F{row},CATALOGOS!$A$2:$B$50,2,0),"XX") & "-" & WEEKNUM(C{row},2) & "-" & TEXT(ROW()-1,"000"))'

    print("Saving workbook...")
    wb.save(excel_path)
    print("Formulas inyectadas en las 3 hojas con exito usando openpyxl.")

except Exception as e:
    print(f"Error: {e}")
