# -*- coding: utf-8 -*-
"""
Generador Maestro de Reportes: Facturado por Obra y Día (Diesel y Gasolina)
- Lee facturas de 'diesel.facturas' y 'gasolina.facturas' en fenix_db
- Replica con máxima fidelidad los cuadros de Excel 'FACTURADO POR OBRA Y DÍA (FOLIOS, MONTOS CON IVA Y LITROS)'
- Genera PDF consolidado de alta calidad para impresión (1 página por semana, 0 encimados)
- Genera PDFs individuales para Diesel y Gasolina
- Genera libro Excel estructurado
"""

import os
import datetime
from collections import defaultdict
import psycopg2
from psycopg2.extras import DictCursor

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, legal, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, PageBreak
from reportlab.lib.styles import ParagraphStyle

BASE_DIR = r'c:\Users\JOSE\Desktop\Proyecto fenix'

def get_db_data(modulo):
    conn = psycopg2.connect(dbname='fenix_db', user='postgres', password='Bupito*268', host='localhost')
    cur = conn.cursor(cursor_factory=DictCursor)
    
    cur.execute(f"""
        SELECT DISTINCT semana 
        FROM {modulo}.facturas 
        WHERE (estatus_revision IS NULL OR estatus_revision = 'APROBADO')
          AND semana IS NOT NULL AND semana != ''
    """)
    raw_semanas = [r['semana'] for r in cur.fetchall()]
    semanas = sorted(raw_semanas, key=lambda s: int(s) if s.isdigit() else 999)
    
    semanas_data = {}
    
    DIAS_ES = {'Mon':'LUN','Tue':'MAR','Wed':'MIÉ','Thu':'JUE','Fri':'VIE','Sat':'SÁB','Sun':'DOM'}
    MESES_ES = {1:'ENE',2:'FEB',3:'MAR',4:'ABR',5:'MAY',6:'JUN',7:'JUL',8:'AGO',9:'SEP',10:'OCT',11:'NOV',12:'DIC'}
    
    for sem in semanas:
        cur.execute(f"""
            SELECT folio_factura, 
                   fecha_factura::text as fecha_str, 
                   COALESCE(NULLIF(obra_destino, ''), 'SIN OBRA') as obra,
                   COALESCE(importe_total, 0) as monto,
                   COALESCE(litros_facturados, 0) as litros
            FROM {modulo}.facturas
            WHERE semana = %s AND (estatus_revision IS NULL OR estatus_revision = 'APROBADO')
              AND (obra_destino IS NULL OR (obra_destino NOT ILIKE '%%Tanque Pegaso%%' AND obra_destino NOT ILIKE '%%COSUM%%' AND obra_destino NOT ILIKE '%%transporte%%'))
            ORDER BY fecha_str, folio_factura
        """, (sem,))
        rows = cur.fetchall()
        if not rows:
            continue
            
        fact_matrix = defaultdict(list)
        fact_obras_set = set()
        fact_fechas_dict = {}
        
        for fr in rows:
            f_date_raw = fr['fecha_str'] or 'Sin Fecha'
            f_obra = fr['obra'] or 'SIN OBRA'
            f_folio = fr['folio_factura'] or 'S/F'
            f_monto = float(fr['monto'] or 0)
            f_litros = float(fr['litros'] or 0)
            
            try:
                d_obj = datetime.datetime.strptime(f_date_raw, '%Y-%m-%d')
                nom_d = DIAS_ES.get(d_obj.strftime('%a'), d_obj.strftime('%a').upper())
                nom_m = MESES_ES.get(d_obj.month, str(d_obj.month))
                fecha_fmt = f"{nom_d} {d_obj.day:02d}-{nom_m}"
            except Exception:
                fecha_fmt = f_date_raw
                
            fact_fechas_dict[f_date_raw] = fecha_fmt
            fact_obras_set.add(f_obra)
            
            fact_matrix[(f_obra, f_date_raw)].append({
                'folio': f_folio,
                'monto': f_monto,
                'litros': f_litros
            })
            
        sorted_obras = sorted(list(fact_obras_set))
        sorted_fechas = sorted(list(fact_fechas_dict.keys()))
        
        semanas_data[sem] = {
            'semana': sem,
            'modulo': modulo.upper(),
            'fechas_raw': sorted_fechas,
            'fechas_dict': fact_fechas_dict,
            'obras': sorted_obras,
            'matrix': fact_matrix,
            'total_facturas': len(rows),
            'total_monto': sum(float(r['monto'] or 0) for r in rows),
            'total_litros': sum(float(r['litros'] or 0) for r in rows)
        }
        
    cur.close()
    conn.close()
    return semanas_data

def build_pdf_flowables(modules_semanas_list, avail_w):
    """
    Construye la lista de elementos (Table + PageBreak) para las semanas indicadas.
    """
    C_TITLE_BG = colors.HexColor('#1E3A5F')
    C_HEAD_BG = colors.HexColor('#1E3A5F')
    C_SUBHEAD_BG = colors.HexColor('#1F4E79')
    C_ACTIVE_CELL_BG = colors.HexColor('#E0F2FE')
    C_ALT_ROW_BG = colors.HexColor('#F8FAFC')
    C_TOTAL_BG = colors.HexColor('#1E3A5F')
    C_BORDER = colors.HexColor('#94A3B8')
    C_BLUE_TEXT = colors.HexColor('#0369A1')
    
    style_title = ParagraphStyle('PDFTitle', fontName='Helvetica-Bold', fontSize=8.0, leading=9.0, alignment=0, textColor=colors.white)
    style_head = ParagraphStyle('PDFHead', fontName='Helvetica-Bold', fontSize=5.8, leading=6.8, alignment=1, textColor=colors.white)
    style_subhead = ParagraphStyle('PDFSubHead', fontName='Helvetica-Bold', fontSize=4.6, leading=5.4, alignment=1, textColor=colors.white)
    style_obra = ParagraphStyle('PDFObra', fontName='Helvetica-Bold', fontSize=5.2, leading=6.0, alignment=0, textColor=colors.HexColor('#0F172A'))
    
    style_folio_active = ParagraphStyle('PDFFolioA', fontName='Helvetica-Bold', fontSize=4.8, leading=5.6, alignment=1, textColor=C_BLUE_TEXT)
    style_folio_empty = ParagraphStyle('PDFFolioE', fontName='Helvetica', fontSize=4.8, leading=5.6, alignment=1, textColor=colors.HexColor('#94A3B8'))
    
    style_num_active = ParagraphStyle('PDFNumA', fontName='Helvetica', fontSize=4.8, leading=5.6, alignment=2, textColor=colors.HexColor('#0F172A'))
    style_num_empty = ParagraphStyle('PDFNumE', fontName='Helvetica', fontSize=4.8, leading=5.6, alignment=1, textColor=colors.HexColor('#94A3B8'))
    
    style_tot_obra_cnt = ParagraphStyle('PDFTotObraCnt', fontName='Helvetica-Bold', fontSize=4.8, leading=5.6, alignment=1, textColor=colors.HexColor('#0F172A'))
    style_tot_obra_num = ParagraphStyle('PDFTotObraNum', fontName='Helvetica-Bold', fontSize=4.8, leading=5.6, alignment=2, textColor=colors.HexColor('#0F172A'))
    
    style_grand_lbl = ParagraphStyle('PDFGrandLbl', fontName='Helvetica-Bold', fontSize=5.2, leading=6.2, alignment=2, textColor=colors.white)
    style_grand_cnt = ParagraphStyle('PDFGrandCnt', fontName='Helvetica-Bold', fontSize=5.2, leading=6.2, alignment=1, textColor=colors.white)
    style_grand_num = ParagraphStyle('PDFGrandNum', fontName='Helvetica-Bold', fontSize=5.2, leading=6.2, alignment=2, textColor=colors.white)

    flowables = []
    
    total_items = len(modules_semanas_list)
    for item_idx, s_info in enumerate(modules_semanas_list):
        sem = s_info['semana']
        mod_name = s_info['modulo']
        fechas_raw = s_info['fechas_raw']
        fechas_dict = s_info['fechas_dict']
        obras = s_info['obras']
        matrix = s_info['matrix']
        num_dias = len(fechas_raw)
        
        table_matrix = []
        tot_cols = 1 + num_dias * 3 + 3
        
        # Row 0: Banner
        title_text = f"▶ FACTURADO POR OBRA Y DÍA — {mod_name} (FOLIOS, MONTOS CON IVA Y LITROS — SEMANA {sem})"
        table_matrix.append([Paragraph(title_text, style_title)] + [''] * (tot_cols - 1))
        
        # Row 1: Header Row 1
        h_row_1 = [Paragraph('OBRA / DESTINO', style_head)]
        for f_raw in fechas_raw:
            f_lbl = fechas_dict[f_raw]
            h_row_1.extend([Paragraph(f_lbl.upper(), style_head), '', ''])
        h_row_1.extend([Paragraph('TOTALES ACUMULADOS', style_head), '', ''])
        table_matrix.append(h_row_1)
        
        # Row 2: Header Row 2
        h_row_2 = ['']
        for _ in range(num_dias):
            h_row_2.extend([
                Paragraph('FACTURA', style_subhead),
                Paragraph('MONTO ($)', style_subhead),
                Paragraph('LITROS', style_subhead)
            ])
        h_row_2.extend([
            Paragraph('TOTAL FACTURAS', style_subhead),
            Paragraph('TOTAL MONTO', style_subhead),
            Paragraph('TOTAL LITROS', style_subhead)
        ])
        table_matrix.append(h_row_2)
        
        # Rows 3+: Obra data rows
        active_cell_styles = []
        data_row_heights = []
        
        for o_idx, obra_nom in enumerate(obras):
            row_idx = len(table_matrix)
            r_cells = [Paragraph(obra_nom[:32], style_obra)]
            
            tot_obra_cnt = 0
            tot_obra_monto = 0.0
            tot_obra_litros = 0.0
            max_lines_in_row = 1
            
            for f_idx, f_raw in enumerate(fechas_raw):
                inv_items = matrix.get((obra_nom, f_raw), [])
                col_start = 1 + f_idx * 3
                if inv_items:
                    folios_str = '<br/>'.join([x['folio'] for x in inv_items])
                    d_monto = sum(x['monto'] for x in inv_items)
                    d_litros = sum(x['litros'] for x in inv_items)
                    tot_obra_cnt += len(inv_items)
                    tot_obra_monto += d_monto
                    tot_obra_litros += d_litros
                    if len(inv_items) > max_lines_in_row:
                        max_lines_in_row = len(inv_items)
                        
                    r_cells.append(Paragraph(folios_str, style_folio_active))
                    r_cells.append(Paragraph(f"${d_monto:,.2f}", style_num_active))
                    r_cells.append(Paragraph(f"{d_litros:,.2f}", style_num_active))
                    active_cell_styles.append((col_start, row_idx, col_start + 2, row_idx))
                else:
                    r_cells.append(Paragraph('-', style_folio_empty))
                    r_cells.append(Paragraph('-', style_num_empty))
                    r_cells.append(Paragraph('-', style_num_empty))
                    
            # Totales Obra
            if tot_obra_cnt > 0:
                r_cells.append(Paragraph(f"{tot_obra_cnt} Factura(s)", style_tot_obra_cnt))
                r_cells.append(Paragraph(f"${tot_obra_monto:,.2f}", style_tot_obra_num))
                r_cells.append(Paragraph(f"{tot_obra_litros:,.2f}", style_tot_obra_num))
            else:
                r_cells.append(Paragraph('-', style_num_empty))
                r_cells.append(Paragraph('-', style_num_empty))
                r_cells.append(Paragraph('-', style_num_empty))
                
            table_matrix.append(r_cells)
            
            # Altura dinámica por fila
            base_h = 16.0 if len(obras) <= 8 else 13.0
            if max_lines_in_row > 1:
                base_h += (max_lines_in_row - 1) * 6.5
            data_row_heights.append(base_h)
            
        # Fila de Totales
        tot_row_idx = len(table_matrix)
        r_total = [Paragraph('TOTAL POR DÍA', style_grand_lbl)]
        
        tot_global_cnt = 0
        tot_global_monto = 0.0
        tot_global_litros = 0.0
        
        for f_idx, f_raw in enumerate(fechas_raw):
            day_tot_cnt = sum(len(matrix.get((o_nom, f_raw), [])) for o_nom in obras)
            day_tot_monto = sum(sum(x['monto'] for x in matrix.get((o_nom, f_raw), [])) for o_nom in obras)
            day_tot_litros = sum(sum(x['litros'] for x in matrix.get((o_nom, f_raw), [])) for o_nom in obras)
            
            tot_global_cnt += day_tot_cnt
            tot_global_monto += day_tot_monto
            tot_global_litros += day_tot_litros
            
            if day_tot_cnt > 0:
                r_total.append(Paragraph(f"{day_tot_cnt} Factura(s)", style_grand_cnt))
                r_total.append(Paragraph(f"${day_tot_monto:,.2f}", style_grand_num))
                r_total.append(Paragraph(f"{day_tot_litros:,.2f}", style_grand_num))
            else:
                r_total.append(Paragraph('-', style_grand_cnt))
                r_total.append(Paragraph('-', style_grand_num))
                r_total.append(Paragraph('-', style_grand_num))
                
        r_total.append(Paragraph(f"{tot_global_cnt} Factura(s)", style_grand_cnt))
        r_total.append(Paragraph(f"${tot_global_monto:,.2f}", style_grand_num))
        r_total.append(Paragraph(f"{tot_global_litros:,.2f}", style_grand_num))
        table_matrix.append(r_total)
        
        # Column Widths
        base_obra_w = 110.0
        base_fact_w = 34.0
        base_monto_w = 44.0
        base_lts_w = 36.0
        base_tot_cnt_w = 42.0
        base_tot_monto_w = 50.0
        base_tot_lts_w = 44.0
        
        raw_col_widths = [base_obra_w]
        for _ in range(num_dias):
            raw_col_widths.extend([base_fact_w, base_monto_w, base_lts_w])
        raw_col_widths.extend([base_tot_cnt_w, base_tot_monto_w, base_tot_lts_w])
        
        scale_w = avail_w / sum(raw_col_widths)
        scaled_widths = [w * scale_w for w in raw_col_widths]
        
        # Spans
        t_spans = [
            ('SPAN', (0, 0), (-1, 0)),
            ('SPAN', (0, 1), (0, 2)),
        ]
        for f_idx in range(num_dias):
            c_s = 1 + f_idx * 3
            t_spans.append(('SPAN', (c_s, 1), (c_s + 2, 1)))
            
        tot_c_s = 1 + num_dias * 3
        t_spans.append(('SPAN', (tot_c_s, 1), (tot_c_s + 2, 1)))
        
        # Styles
        t_style = [
            ('BACKGROUND', (0, 0), (-1, 0), C_TITLE_BG),
            ('BACKGROUND', (0, 1), (-1, 1), C_HEAD_BG),
            ('BACKGROUND', (0, 2), (-1, 2), C_SUBHEAD_BG),
            ('BACKGROUND', (0, tot_row_idx), (-1, tot_row_idx), C_TOTAL_BG),
            
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.35, C_BORDER),
            ('LINEBELOW', (0, tot_row_idx), (-1, tot_row_idx), 0.8, colors.HexColor('#0F172A')),
            
            ('TOPPADDING', (0, 0), (-1, -1), 0.2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0.2),
            ('LEFTPADDING', (0, 0), (-1, -1), 0.6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0.6),
        ] + t_spans
        
        for o_i in range(len(obras)):
            r_num = 3 + o_i
            bg_r = C_ALT_ROW_BG if (o_i % 2 == 1) else colors.white
            t_style.append(('BACKGROUND', (0, r_num), (-1, r_num), bg_r))
            
        for c1, r1, c2, r2 in active_cell_styles:
            t_style.append(('BACKGROUND', (c1, r1), (c2, r2), C_ACTIVE_CELL_BG))
            
        row_heights = [18.0, 13.0, 13.0] + data_row_heights + [16.0]
        
        table = Table(table_matrix, colWidths=scaled_widths, rowHeights=row_heights)
        table.setStyle(TableStyle(t_style))
        flowables.append(table)
        
        if item_idx < total_items - 1:
            flowables.append(PageBreak())
            
    return flowables

def generate_pdf_doc(modules_semanas_list, output_path, paper_size='legal'):
    psize = landscape(legal) if paper_size == 'legal' else landscape(letter)
    left_m = 14.0
    right_m = 14.0
    top_m = 14.0
    bot_m = 14.0
    avail_w = psize[0] - left_m - right_m
    
    doc = SimpleDocTemplate(
        output_path,
        pagesize=psize,
        leftMargin=left_m,
        rightMargin=right_m,
        topMargin=top_m,
        bottomMargin=bot_m
    )
    
    story = build_pdf_flowables(modules_semanas_list, avail_w)
    doc.build(story)
    print(f"PDF generado exitosamente ({len(modules_semanas_list)} páginas): {output_path}")

def generate_full_excel(diesel_data, gasolina_data, output_path):
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    
    f_title = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    f_header = Font(name='Calibri', size=9, bold=True, color='FFFFFF')
    f_subhead = Font(name='Calibri', size=8, bold=True, color='FFFFFF')
    f_data_bold = Font(name='Calibri', size=8.5, bold=True, color='0F172A')
    f_data = Font(name='Calibri', size=8.5, color='0F172A')
    f_folio_act = Font(name='Calibri', size=8.5, bold=True, color='0369A1')
    f_empty = Font(name='Calibri', size=8.5, color='94A3B8')
    f_total = Font(name='Calibri', size=9, bold=True, color='FFFFFF')
    
    fill_title = PatternFill(start_color='1E3A5F', end_color='1E3A5F', fill_type='solid')
    fill_head = PatternFill(start_color='1E3A5F', end_color='1E3A5F', fill_type='solid')
    fill_subhead = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    fill_act = PatternFill(start_color='E0F2FE', end_color='E0F2FE', fill_type='solid')
    fill_alt = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_tot = PatternFill(start_color='1E3A5F', end_color='1E3A5F', fill_type='solid')
    
    thin_s = Side(border_style='thin', color='CBD5E1')
    border_c = Border(left=thin_s, right=thin_s, top=thin_s, bottom=thin_s)
    
    align_c = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_l = Alignment(horizontal='left', vertical='center')
    align_r = Alignment(horizontal='right', vertical='center')

    all_modules = [('Diesel', diesel_data), ('Gasolina', gasolina_data)]
    
    for mod_name, s_data in all_modules:
        semanas_keys = sorted(s_data.keys(), key=lambda s: int(s) if s.isdigit() else 999)
        for sem in semanas_keys:
            s_info = s_data[sem]
            fechas_raw = s_info['fechas_raw']
            fechas_dict = s_info['fechas_dict']
            obras = s_info['obras']
            matrix = s_info['matrix']
            num_dias = len(fechas_raw)
            
            ws = wb.create_sheet(title=f"{mod_name} Sem {sem}")
            ws.views.sheetView[0].showGridLines = True
            ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
            ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
            ws.page_setup.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            
            tot_cols = 1 + num_dias * 3 + 3
            end_col_letter = get_column_letter(tot_cols)
            
            # Banner
            ws.merge_cells(f'A1:{end_col_letter}1')
            ws['A1'] = f"▶ FACTURADO POR OBRA Y DÍA — {mod_name.upper()} (FOLIOS, MONTOS CON IVA Y LITROS — SEMANA {sem})"
            ws['A1'].font = f_title
            ws['A1'].fill = fill_title
            ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
            ws.row_dimensions[1].height = 24
            
            # Header Row 1
            ws.merge_cells('A2:A3')
            ws['A2'] = 'OBRA / DESTINO'
            ws['A2'].font = f_header; ws['A2'].fill = fill_head; ws['A2'].alignment = align_c; ws['A2'].border = border_c
            ws['A3'].border = border_c; ws['A3'].fill = fill_head
            
            for f_idx, f_raw in enumerate(fechas_raw):
                c_start = 2 + f_idx * 3
                c_end = c_start + 2
                ws.merge_cells(start_row=2, start_column=c_start, end_row=2, end_column=c_end)
                c = ws.cell(2, c_start, fechas_dict[f_raw].upper())
                c.font = f_header; c.fill = fill_head; c.alignment = align_c
                for c_i in range(c_start, c_end + 1):
                    ws.cell(2, c_i).border = border_c
                    ws.cell(2, c_i).fill = fill_head
                    
            tot_c_start = 2 + num_dias * 3
            tot_c_end = tot_c_start + 2
            ws.merge_cells(start_row=2, start_column=tot_c_start, end_row=2, end_column=tot_c_end)
            c = ws.cell(2, tot_c_start, 'TOTALES ACUMULADOS')
            c.font = f_header; c.fill = fill_head; c.alignment = align_c
            for c_i in range(tot_c_start, tot_c_end + 1):
                ws.cell(2, c_i).border = border_c
                ws.cell(2, c_i).fill = fill_head
                
            ws.row_dimensions[2].height = 18
            
            # Sub-headers
            for f_idx in range(num_dias):
                c_start = 2 + f_idx * 3
                for offset, lbl in enumerate(['FACTURA', 'MONTO ($)', 'LITROS']):
                    cell = ws.cell(3, c_start + offset, lbl)
                    cell.font = f_subhead; cell.fill = fill_subhead; cell.alignment = align_c; cell.border = border_c
                    
            for offset, lbl in enumerate(['TOTAL FACTURAS', 'TOTAL MONTO', 'TOTAL LITROS']):
                cell = ws.cell(3, tot_c_start + offset, lbl)
                cell.font = f_subhead; cell.fill = fill_subhead; cell.alignment = align_c; cell.border = border_c
                
            ws.row_dimensions[3].height = 18
            
            # Data Rows
            r_curr = 4
            for o_idx, obra_nom in enumerate(obras):
                bg_row = fill_alt if (o_idx % 2 == 1) else PatternFill(fill_type=None)
                
                c_o = ws.cell(r_curr, 1, obra_nom)
                c_o.font = f_data_bold; c_o.alignment = align_l; c_o.border = border_c
                if bg_row.fill_type: c_o.fill = bg_row
                    
                tot_o_cnt = 0
                tot_o_monto = 0.0
                tot_o_lts = 0.0
                
                for f_idx, f_raw in enumerate(fechas_raw):
                    inv_items = matrix.get((obra_nom, f_raw), [])
                    c_start = 2 + f_idx * 3
                    if inv_items:
                        folios_str = ', '.join([x['folio'] for x in inv_items])
                        d_monto = sum(x['monto'] for x in inv_items)
                        d_litros = sum(x['litros'] for x in inv_items)
                        tot_o_cnt += len(inv_items)
                        tot_o_monto += d_monto
                        tot_o_lts += d_litros
                        
                        c0 = ws.cell(r_curr, c_start + 0, folios_str)
                        c0.font = f_folio_act; c0.fill = fill_act; c0.alignment = align_c; c0.border = border_c
                        
                        c1 = ws.cell(r_curr, c_start + 1, d_monto)
                        c1.font = f_data; c1.fill = fill_act; c1.alignment = align_r; c1.border = border_c
                        c1.number_format = '"$"#,##0.00'
                        
                        c2 = ws.cell(r_curr, c_start + 2, d_litros)
                        c2.font = f_data; c2.fill = fill_act; c2.alignment = align_r; c2.border = border_c
                        c2.number_format = '#,##0.00'
                    else:
                        for offset in range(3):
                            cell = ws.cell(r_curr, c_start + offset, '-')
                            cell.font = f_empty; cell.alignment = align_c; cell.border = border_c
                            if bg_row.fill_type: cell.fill = bg_row
                            
                c0 = ws.cell(r_curr, tot_c_start + 0, f"{tot_o_cnt} Factura(s)" if tot_o_cnt > 0 else '-')
                c0.font = f_data_bold; c0.alignment = align_c; c0.border = border_c
                if bg_row.fill_type: c0.fill = bg_row
                
                c1 = ws.cell(r_curr, tot_c_start + 1, tot_o_monto if tot_o_cnt > 0 else '-')
                c1.font = f_data_bold; c1.alignment = align_r if tot_o_cnt > 0 else align_c; c1.border = border_c
                if tot_o_cnt > 0: c1.number_format = '"$"#,##0.00'
                if bg_row.fill_type: c1.fill = bg_row
                
                c2 = ws.cell(r_curr, tot_c_start + 2, tot_o_lts if tot_o_cnt > 0 else '-')
                c2.font = f_data_bold; c2.alignment = align_r if tot_o_cnt > 0 else align_c; c2.border = border_c
                if tot_o_cnt > 0: c2.number_format = '#,##0.00'
                if bg_row.fill_type: c2.fill = bg_row
                
                ws.row_dimensions[r_curr].height = 18
                r_curr += 1
                
            # Totales
            c_tot_lbl = ws.cell(r_curr, 1, 'TOTAL POR DÍA')
            c_tot_lbl.font = f_total; c_tot_lbl.fill = fill_tot; c_tot_lbl.alignment = align_r; c_tot_lbl.border = border_c
            
            tot_g_cnt = 0
            tot_g_monto = 0.0
            tot_g_lts = 0.0
            
            for f_idx, f_raw in enumerate(fechas_raw):
                c_start = 2 + f_idx * 3
                d_cnt = sum(len(matrix.get((o, f_raw), [])) for o in obras)
                d_monto = sum(sum(x['monto'] for x in matrix.get((o, f_raw), [])) for o in obras)
                d_lts = sum(sum(x['litros'] for x in matrix.get((o, f_raw), [])) for o in obras)
                
                tot_g_cnt += d_cnt
                tot_g_monto += d_monto
                tot_g_lts += d_lts
                
                c0 = ws.cell(r_curr, c_start + 0, f"{d_cnt} Factura(s)" if d_cnt > 0 else '-')
                c0.font = f_total; c0.fill = fill_tot; c0.alignment = align_c; c0.border = border_c
                
                c1 = ws.cell(r_curr, c_start + 1, d_monto if d_cnt > 0 else '-')
                c1.font = f_total; c1.fill = fill_tot; c1.alignment = align_r if d_cnt > 0 else align_c; c1.border = border_c
                if d_cnt > 0: c1.number_format = '"$"#,##0.00'
                
                c2 = ws.cell(r_curr, c_start + 2, d_lts if d_cnt > 0 else '-')
                c2.font = f_total; c2.fill = fill_tot; c2.alignment = align_r if d_cnt > 0 else align_c; c2.border = border_c
                if d_cnt > 0: c2.number_format = '#,##0.00'
                
            c0 = ws.cell(r_curr, tot_c_start + 0, f"{tot_g_cnt} Factura(s)")
            c0.font = f_total; c0.fill = fill_tot; c0.alignment = align_c; c0.border = border_c
            
            c1 = ws.cell(r_curr, tot_c_start + 1, tot_g_monto)
            c1.font = f_total; c1.fill = fill_tot; c1.alignment = align_r; c1.border = border_c
            c1.number_format = '"$"#,##0.00'
            
            c2 = ws.cell(r_curr, tot_c_start + 2, tot_g_lts)
            c2.font = f_total; c2.fill = fill_tot; c2.alignment = align_r; c2.border = border_c
            c2.number_format = '#,##0.00'
            
            ws.row_dimensions[r_curr].height = 20
            
            # Widths
            ws.column_dimensions['A'].width = 28
            for f_idx in range(num_dias):
                c_start = 2 + f_idx * 3
                ws.column_dimensions[get_column_letter(c_start + 0)].width = 13
                ws.column_dimensions[get_column_letter(c_start + 1)].width = 14
                ws.column_dimensions[get_column_letter(c_start + 2)].width = 12
            ws.column_dimensions[get_column_letter(tot_c_start + 0)].width = 15
            ws.column_dimensions[get_column_letter(tot_c_start + 1)].width = 16
            ws.column_dimensions[get_column_letter(tot_c_start + 2)].width = 14

    if default_sheet.title in wb.sheetnames:
        wb.remove(default_sheet)
        
    wb.save(output_path)
    print(f"Excel maestro guardado exitosamente: {output_path}")

if __name__ == '__main__':
    print("=== GENERACIÓN DE REPORTES FACTURADO POR OBRA Y DÍA ===")
    
    # 1. Extracción de datos
    print("1. Extrayendo datos de la base de datos...")
    diesel_data = get_db_data('diesel')
    gasolina_data = get_db_data('gasolina')
    
    sorted_diesel_semanas = [diesel_data[k] for k in sorted(diesel_data.keys(), key=lambda s: int(s) if s.isdigit() else 999)]
    sorted_gasolina_semanas = [gasolina_data[k] for k in sorted(gasolina_data.keys(), key=lambda s: int(s) if s.isdigit() else 999)]
    
    # Lista consolidada completa (Diesel + Gasolina)
    consolidated_all = sorted_diesel_semanas + sorted_gasolina_semanas
    
    # 2. Generación de PDFs
    print("2. Generando archivos PDF...")
    
    # PDF Maestro Consolidado (Diesel + Gasolina)
    pdf_consolidado_oficio = os.path.join(BASE_DIR, 'Reporte_Facturado_Por_Obra_y_Dia_Consolidado_Oficio.pdf')
    pdf_consolidado_carta = os.path.join(BASE_DIR, 'Reporte_Facturado_Por_Obra_y_Dia_Consolidado_Carta.pdf')
    generate_pdf_doc(consolidated_all, pdf_consolidado_oficio, paper_size='legal')
    generate_pdf_doc(consolidated_all, pdf_consolidado_carta, paper_size='letter')
    
    # PDF Solo Diesel
    pdf_diesel_oficio = os.path.join(BASE_DIR, 'Reporte_Facturado_Por_Obra_y_Dia_Diesel_Oficio.pdf')
    pdf_diesel_carta = os.path.join(BASE_DIR, 'Reporte_Facturado_Por_Obra_y_Dia_Diesel_Carta.pdf')
    generate_pdf_doc(sorted_diesel_semanas, pdf_diesel_oficio, paper_size='legal')
    generate_pdf_doc(sorted_diesel_semanas, pdf_diesel_carta, paper_size='letter')
    
    # PDF Solo Gasolina
    pdf_gasolina_oficio = os.path.join(BASE_DIR, 'Reporte_Facturado_Por_Obra_y_Dia_Gasolina_Oficio.pdf')
    pdf_gasolina_carta = os.path.join(BASE_DIR, 'Reporte_Facturado_Por_Obra_y_Dia_Gasolina_Carta.pdf')
    generate_pdf_doc(sorted_gasolina_semanas, pdf_gasolina_oficio, paper_size='legal')
    generate_pdf_doc(sorted_gasolina_semanas, pdf_gasolina_carta, paper_size='letter')
    
    # 3. Generación de Excel Maestro
    print("3. Generando libro Excel estructurado...")
    excel_maestro = os.path.join(BASE_DIR, 'Reporte_Facturado_Por_Obra_y_Dia_Semanas_25_a_37.xlsx')
    generate_full_excel(diesel_data, gasolina_data, excel_maestro)
    
    print("\n¡PROCESO COMPLETADO EXITOSAMENTE!")
