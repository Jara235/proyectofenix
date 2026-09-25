import psycopg2
from psycopg2.extras import DictCursor
import openpyxl
import pandas as pd

conn = psycopg2.connect(dbname='fenix_db', user='postgres', password='Bupito*268', host='localhost')
cur = conn.cursor(cursor_factory=DictCursor)

# Load workbook daily values accurately
wb = openpyxl.load_workbook('google_sheet_download.xlsx', data_only=True)
week_sheets = sorted([s for s in wb.sheetnames if s.startswith("Semana")], key=lambda s: int(s.replace("Semana ", "")))

sheet_data = {}
for sname in week_sheets:
    ws = wb[sname]
    s_num = sname.replace("Semana ", "").strip()
    v_fres = float(ws['E5'].value or 0)
    v_carp = float(ws['E6'].value or 0)
    ton_carp = float(ws['G6'].value or 0)
    
    # Calculate sum of diesel from rows 27 to 38 daily columns F, J, N, R, V, Z
    d_tot = 0.0
    for r in range(27, 39):
        for c in ['F', 'J', 'N', 'R', 'V', 'Z']:
            v = ws[f'{c}{r}'].value
            if v is not None and v != '-' and str(v).replace('.','',1).isdigit():
                d_tot += float(v)
                
    sheet_data[s_num] = {
        'semana': sname,
        'fresado_m3': v_fres,
        'carpeta_m3': v_carp,
        'carpeta_ton': ton_carp,
        'vol_tot_m3': v_fres + v_carp,
        'diesel_sheet': d_tot
    }

# Query consumos from DB
cur.execute("""
    SELECT semana, count(*) as cant_cargas, sum(litros) as tot_litros, sum(importe_total) as tot_importe
    FROM diesel.consumos
    WHERE obra_destino ILIKE '%toluca%' OR obra_destino ILIKE '%mexico%'
    GROUP BY semana
    ORDER BY semana;
""")
db_consumos = {}
for r in cur.fetchall():
    s_clean = str(r['semana']).replace('Semana ', '').strip()
    db_consumos[s_clean] = {
        'cant': r['cant_cargas'],
        'litros': float(r['tot_litros'] or 0),
        'importe': float(r['tot_importe'] or 0)
    }

# Query facturas from DB
cur.execute("""
    SELECT semana, count(*) as cant_facturas, sum(litros_facturados) as tot_litros, sum(importe_total) as tot_importe
    FROM diesel.facturas
    WHERE obra_destino ILIKE '%toluca%' OR obra_destino ILIKE '%mexico%'
    GROUP BY semana
    ORDER BY semana;
""")
db_facturas = {}
for r in cur.fetchall():
    s_clean = str(r['semana']).replace('Semana ', '').strip()
    db_facturas[s_clean] = {
        'cant': r['cant_facturas'],
        'litros': float(r['tot_litros'] or 0),
        'importe': float(r['tot_importe'] or 0)
    }

# Query suppliers from facturas for Mexico Toluca
cur.execute("""
    SELECT proveedor, count(*) as cant, sum(litros_facturados) as tot_litros, sum(importe_total) as tot_importe
    FROM diesel.facturas
    WHERE obra_destino ILIKE '%toluca%' OR obra_destino ILIKE '%mexico%'
    GROUP BY proveedor
    ORDER BY tot_litros DESC;
""")
proveedores = cur.fetchall()

print("="*100)
print("PROVEEDORES QUE FACTURARON DIÉSEL PARA OBRA MÉXICO-TOLUCA:")
print("="*100)
for p in proveedores:
    print(f"  * {p['proveedor']:45s} | {p['cant']:2d} facturas | {p['tot_litros']:10,.2f} L | ${p['tot_importe']:12,.2f} MXN")

# Build comprehensive weekly reconciliation table for Semanas 24 to 37
rows = []
for s_num in [str(i) for i in range(24, 38)]:
    sh = sheet_data.get(s_num, {'vol_tot_m3': 0, 'fresado_m3': 0, 'carpeta_ton': 0, 'diesel_sheet': 0})
    dc = db_consumos.get(s_num, {'cant': 0, 'litros': 0, 'importe': 0})
    df = db_facturas.get(s_num, {'cant': 0, 'litros': 0, 'importe': 0})
    
    vol_tot = sh['vol_tot_m3']
    d_sheet = sh['diesel_sheet']
    d_db_cons = dc['litros']
    d_fact = df['litros']
    imp_fact = df['importe']
    
    # Differences
    diff_db_vs_sheet = d_db_cons - d_sheet
    diff_fact_vs_sheet = d_fact - d_sheet
    diff_fact_vs_db = d_fact - d_db_cons
    pct_desv_fact_vs_sheet = (diff_fact_vs_sheet / d_sheet * 100) if d_sheet > 0 else 0.0
    
    # Performance
    rend_real_sheet = d_sheet / vol_tot if vol_tot > 0 else 0.0
    rend_facturado = d_fact / vol_tot if vol_tot > 0 else 0.0
    
    rows.append({
        'Semana': f"Semana {s_num}",
        'Volumen_m3': vol_tot,
        'Diésel_Bitacora_L': d_sheet,
        'Diésel_DB_Cons_L': d_db_cons,
        'Diff_DB_vs_Bit_L': diff_db_vs_sheet,
        'Diésel_Facturado_L': d_fact,
        'Importe_Fact_MXN': imp_fact,
        'Diff_Fact_vs_Bit_L': diff_fact_vs_sheet,
        'Pct_Desv_%': pct_desv_fact_vs_sheet,
        'Rend_Bitacora_Lm3': rend_real_sheet,
        'Rend_Facturado_Lm3': rend_facturado
    })

df_all = pd.DataFrame(rows)
print("\n" + "="*120)
print("CONCILIACIÓN TRIANGULADA: BITÁCORA/GOOGLE SHEET VS. DB CONSUMOS VS. DB FACTURAS (SEMANAS 24 A 37)")
print("="*120)
pd.set_option('display.max_columns', 15)
pd.set_option('display.width', 1000)
print(df_all[['Semana', 'Volumen_m3', 'Diésel_Bitacora_L', 'Diésel_DB_Cons_L', 'Diff_DB_vs_Bit_L', 'Diésel_Facturado_L', 'Diff_Fact_vs_Bit_L', 'Pct_Desv_%', 'Rend_Bitacora_Lm3', 'Rend_Facturado_Lm3']].to_string(index=False))

# Totals for Semanas 24 to 37
tot_vol = df_all['Volumen_m3'].sum()
tot_bit = df_all['Diésel_Bitacora_L'].sum()
tot_db_c = df_all['Diésel_DB_Cons_L'].sum()
tot_fact = df_all['Diésel_Facturado_L'].sum()
tot_imp = df_all['Importe_Fact_MXN'].sum()

print("\n" + "="*80)
print("RESUMEN DE TOTALES CONSOLIDADOS (SEMANAS 24 A 37):")
print(f"Total Volumen Producido:                      {tot_vol:12,.2f} m³")
print(f"Total Diésel en Bitácora / Google Sheet:      {tot_bit:12,.2f} L")
print(f"Total Diésel en DB diesel.consumos:           {tot_db_c:12,.2f} L")
print(f"Total Diésel Facturado en DB diesel.facturas: {tot_fact:12,.2f} L")
print(f"Importe Total Facturado de Diésel:            ${tot_imp:12,.2f} MXN")
print(f"Precio Promedio Ponderado por Litro:          ${tot_imp / tot_fact:12.2f} MXN/L")
print("-" * 80)
print(f"Diferencia DB Consumos vs Bitácora:           {tot_db_c - tot_bit:+12,.2f} L ({(tot_db_c - tot_bit)/tot_bit*100:+6.2f}%)")
print(f"Diferencia Facturado vs Bitácora (Consumido): {tot_fact - tot_bit:+12,.2f} L ({(tot_fact - tot_bit)/tot_bit*100:+6.2f}%)")
print(f"Diferencia Facturado vs DB Consumos:          {tot_fact - tot_db_c:+12,.2f} L ({(tot_fact - tot_db_c)/tot_db_c*100:+6.2f}%)")
print("-" * 80)
print(f"Rendimiento Real Obra (Bitácora de Campo):    {tot_bit / tot_vol:12.3f} L/m³")
print(f"Rendimiento según DB Consumos:                {tot_db_c / tot_vol:12.3f} L/m³")
print(f"Rendimiento según Diésel Facturado:           {tot_fact / tot_vol:12.3f} L/m³")
print("="*80)
