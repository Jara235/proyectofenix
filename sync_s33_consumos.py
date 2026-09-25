from app_admin import get_db
import openpyxl

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)
ws = wb['SEMANA 33']

day_cols = [
    ('2026-08-10', 7, 'LEVET', None),
    ('2026-08-10', 8, 'MOBILE', 9),
    ('2026-08-10', 10, 'SI VALE', None),
    ('2026-08-11', 11, 'LEVET', None),
    ('2026-08-11', 12, 'MOBILE', 13),
    ('2026-08-11', 14, 'SI VALE', None),
    ('2026-08-12', 15, 'LEVET', None),
    ('2026-08-12', 16, 'MOBILE', 17),
    ('2026-08-12', 18, 'SI VALE', None),
    ('2026-08-13', 19, 'LEVET', None),
    ('2026-08-13', 20, 'MOBILE', 21),
    ('2026-08-13', 22, 'SI VALE', None),
    ('2026-08-14', 23, 'LEVET', None),
    ('2026-08-14', 24, 'MOBILE', 25),
    ('2026-08-14', 26, 'SI VALE', None),
    ('2026-08-15', 27, 'LEVET', None),
    ('2026-08-15', 28, 'MOBILE', 29),
    ('2026-08-15', 30, 'SI VALE', None),
]

sheet_cargas = []

# Table 1: JDJ
for r in range(4, 44):
    resp = ws.cell(r, 2).value
    obra = ws.cell(r, 3).value
    unidad = ws.cell(r, 4).value
    placa = ws.cell(r, 5).value
    
    resp_str = str(resp).strip() if resp is not None else ''
    obra_str = str(obra).strip() if obra is not None else ''
    unidad_str = str(unidad).strip() if unidad is not None else ''
    placa_str = str(placa).strip().upper() if placa is not None else ''
    if placa_str in ['PLACAS', 'S/P', 'NONE']: placa_str = ''
    
    conductor = resp_str or unidad_str
    if not conductor or 'TOTAL' in conductor.upper():
        continue
        
    for fecha, col_val, gas, col_folio in day_cols:
        val = ws.cell(r, col_val).value
        try:
            val_f = float(val or 0)
        except:
            val_f = 0.0
        if val_f > 0:
            folio = ''
            if col_folio:
                f_val = ws.cell(r, col_folio).value
                if f_val: folio = str(f_val).strip().upper()
            sheet_cargas.append({
                'row': r,
                'empresa': 'JDJ',
                'conductor': conductor,
                'obra': obra_str,
                'unidad': unidad_str,
                'placa': placa_str,
                'fecha': fecha,
                'gasolineria': gas,
                'importe': val_f,
                'folio_sheet': folio
            })

# Table 2: TRD
for r in range(50, 58):
    resp = ws.cell(r, 2).value
    obra = ws.cell(r, 3).value
    unidad = ws.cell(r, 4).value
    placa = ws.cell(r, 5).value
    
    resp_str = str(resp).strip() if resp is not None else ''
    obra_str = str(obra).strip() if obra is not None else ''
    unidad_str = str(unidad).strip() if unidad is not None else ''
    placa_str = str(placa).strip().upper() if placa is not None else ''
    if placa_str in ['PLACAS', 'S/P', 'NONE']: placa_str = ''
    
    conductor = resp_str or unidad_str
    if not conductor or 'TOTAL' in conductor.upper():
        continue
        
    for fecha, col_val, gas, col_folio in day_cols:
        val = ws.cell(r, col_val).value
        try:
            val_f = float(val or 0)
        except:
            val_f = 0.0
        if val_f > 0:
            folio = ''
            if col_folio:
                f_val = ws.cell(r, col_folio).value
                if f_val: folio = str(f_val).strip().upper()
            sheet_cargas.append({
                'row': r,
                'empresa': 'TRD',
                'conductor': conductor,
                'obra': obra_str,
                'unidad': unidad_str,
                'placa': placa_str,
                'fecha': fecha,
                'gasolineria': gas,
                'importe': val_f,
                'folio_sheet': folio
            })

print(f"Total cargas from Google Sheet: {len(sheet_cargas)} (Total: ${sum(c['importe'] for c in sheet_cargas):,.2f})")

# Let's inspect differences against DB
db = get_db()
db_cargas = db.execute("SELECT id, folio_conciliacion, fecha, conductor, gasolineria, importe_total, placa, foto_evidencia IS NOT NULL as has_foto FROM gasolina.consumos WHERE semana = '33' ORDER BY id").fetchall()
print(f"Total cargas in DB: {len(db_cargas)} (Total: ${sum(float(c['importe_total']) for c in db_cargas):,.2f})")

db.close()
