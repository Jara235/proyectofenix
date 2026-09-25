import openpyxl
from copy import copy

excel_path = r"c:\Users\JOSE\Desktop\Proyecto fenix\MAESTRO_CONTROL_DIESEL_NUEVO.xlsx"

try:
    wb = openpyxl.load_workbook(excel_path)
    
    # Check if BD_SOLICITUDES exists, if not create it
    if "BD_SOLICITUDES" not in wb.sheetnames:
        sheet_sol = wb.create_sheet("BD_SOLICITUDES")
    else:
        sheet_sol = wb["BD_SOLICITUDES"]
        sheet_sol.delete_rows(1, sheet_sol.max_row) # Clear if exists
        
    sheet_bd = wb["BD_DIESEL"]
    
    # 1. Copy headers
    for col in range(1, sheet_bd.max_column + 1):
        cell = sheet_bd.cell(row=1, column=col)
        new_cell = sheet_sol.cell(row=1, column=col)
        new_cell.value = cell.value
        if cell.has_style:
            new_cell.font = copy(cell.font)
            new_cell.border = copy(cell.border)
            new_cell.fill = copy(cell.fill)
            new_cell.number_format = copy(cell.number_format)
            new_cell.protection = copy(cell.protection)
            new_cell.alignment = copy(cell.alignment)
            
    # AutoFit approximate
    for col in sheet_sol.columns:
        sheet_sol.column_dimensions[col[0].column_letter].width = 15

    # 2. Extract rows where TIPO_MOVIMIENTO (Col E) is 'SOLICITUD'
    rows_to_move = []
    # Collect backwards so we can delete from BD_DIESEL safely
    for row in range(sheet_bd.max_row, 1, -1):
        mov_type = str(sheet_bd.cell(row=row, column=5).value or "").strip().upper()
        if "SOLICITUD" in mov_type:
            row_data = [sheet_bd.cell(row=row, column=c).value for c in range(1, sheet_bd.max_column + 1)]
            rows_to_move.append(row_data)
            sheet_bd.delete_rows(row, 1)
            
    # They were collected backwards, so reverse them to maintain chronological order
    rows_to_move.reverse()
    
    # Insert into BD_SOLICITUDES
    for idx, row_data in enumerate(rows_to_move, start=2):
        for col_idx, val in enumerate(row_data, start=1):
            sheet_sol.cell(row=idx, column=col_idx).value = val
            
    # Save the file before re-applying formulas using COM (COM handles formulas much better)
    wb.save(excel_path)
    print(f"Se movieron {len(rows_to_move)} solicitudes a la nueva pestana BD_SOLICITUDES.")

except Exception as e:
    print(f"Error: {e}")
