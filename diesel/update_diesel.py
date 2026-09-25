import openpyxl
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/MAESTRO_CONTROL_DIESEL_NUEVO.xlsx')
ws = wb['BD_DIESEL']

for i in range(2, max(ws.max_row + 1, 500)):
    # A = FOLIO_CONCILIACION
    # K = IMPORTE_TOTAL (BD_DIESEL)
    # BD_FACTURAS K = IMPORTE_TOTAL
    ws[f'N{i}'] = f'=IF(A{i}="", "", IF(ROUND(SUMIFS(BD_FACTURAS!K:K, BD_FACTURAS!A:A, A{i}), 2) = ROUND(K{i}, 2), "✅ Cuadrado", "⚠️ Diferencia"))'

wb.save('c:/Users/JOSE/Desktop/Proyecto fenix/MAESTRO_CONTROL_DIESEL_NUEVO.xlsx')
print('Formulas de conciliacion aplicadas en Diesel.')

