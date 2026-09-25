import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

maestro_path = r'c:\Users\JOSE\Desktop\Proyecto fenix\MAESTRO_CONTROL_GASOLINA_NUEVO.xlsx'

bd_gas = pd.read_excel(maestro_path, sheet_name='BD_GASOLINA')
bd_fac = pd.read_excel(maestro_path, sheet_name='BD_FACTURAS')
bd_gas['IMPORTE_TOTAL'] = pd.to_numeric(bd_gas['IMPORTE_TOTAL'], errors='coerce').fillna(0)
bd_gas['LITROS'] = pd.to_numeric(bd_gas['LITROS'], errors='coerce').fillna(0)
bd_fac['IMPORTE_TOTAL'] = pd.to_numeric(bd_fac['IMPORTE_TOTAL'], errors='coerce').fillna(0)
bd_fac['LITROS_FACTURADOS'] = pd.to_numeric(bd_fac['LITROS_FACTURADOS'], errors='coerce').fillna(0)

def normalize_gaso(s):
    s = str(s).strip().upper()
    if 'LEVET' in s: return 'LEVET'
    if 'MOBILE' in s or 'MOBIL' in s: return 'MOBILE HUIXQUILUCAN'
    if 'SI VALE' in s or 'SIVALE' in s: return 'SI VALE'
    if 'CASTILLA' in s or 'DERIVADOS' in s: return 'CASTILLA'
    return s

bd_gas['GASO_KEY'] = bd_gas['ORIGEN'].apply(normalize_gaso)
bd_fac['GASO_KEY'] = bd_fac['PROVEEDOR'].apply(normalize_gaso)

rows = []
for sem in sorted(bd_gas['SEMANA'].dropna().unique()):
    sem_int = int(sem)
    g = bd_gas[bd_gas['SEMANA'] == sem].groupby('GASO_KEY').agg(
        REP_REG=('IMPORTE_TOTAL','count'), REP_LTS=('LITROS','sum'), REP_IMP=('IMPORTE_TOTAL','sum')
    ).reset_index()
    f = bd_fac[bd_fac['SEMANA'] == sem_int].groupby('GASO_KEY').agg(
        FAC_REG=('IMPORTE_TOTAL','count'), FAC_LTS=('LITROS_FACTURADOS','sum'), FAC_IMP=('IMPORTE_TOTAL','sum')
    ).reset_index()
    m = pd.merge(g, f, on='GASO_KEY', how='outer').fillna(0)
    for _, r in m.iterrows():
        dif = round(float(r.REP_IMP) - float(r.FAC_IMP), 2)
        if abs(dif) <= 1.0: est = 'CUADRADO'
        elif dif > 0: est = 'FALTA FACTURA'
        else: est = 'FACTURA EXCEDE'
        rows.append({
            'SEMANA': sem_int, 'GASOLINERA': r.GASO_KEY,
            'REPORTES_REGISTROS': int(r.REP_REG), 'REPORTES_LITROS': round(float(r.REP_LTS),3),
            'REPORTES_IMPORTE': round(float(r.REP_IMP),2),
            'FACTURAS_REGISTROS': int(r.FAC_REG), 'FACTURAS_LITROS': round(float(r.FAC_LTS),3),
            'FACTURAS_IMPORTE': round(float(r.FAC_IMP),2),
            'DIFERENCIA_LITROS': round(float(r.REP_LTS) - float(r.FAC_LTS),3),
            'DIFERENCIA_IMPORTE': dif, 'ESTATUS': est
        })

df = pd.DataFrame(rows)

wb = load_workbook(maestro_path)
if 'CONCILIACION_GASOLINERAS' in wb.sheetnames:
    del wb['CONCILIACION_GASOLINERAS']
ws = wb.create_sheet('CONCILIACION_GASOLINERAS')

thin = Border(left=Side(style='thin',color='CCCCCC'),right=Side(style='thin',color='CCCCCC'),
              top=Side(style='thin',color='CCCCCC'),bottom=Side(style='thin',color='CCCCCC'))
HDR  = PatternFill('solid',fgColor='1A1A2E')
OK   = PatternFill('solid',fgColor='D4EDDA')
WARN = PatternFill('solid',fgColor='FFF3CD')
DANG = PatternFill('solid',fgColor='F8D7DA')

headers = ['SEMANA','GASOLINERA','REPORTES\n(Regs)','REPORTES\n(Lts)','REPORTES\n(Importe)',
           'FACTURAS\n(Regs)','FACTURAS\n(Lts)','FACTURAS\n(Importe)',
           'DIF. LITROS','DIF. IMPORTE','ESTATUS']
widths  = [10,24,14,14,18,14,14,18,14,16,18]
for c,(h,w) in enumerate(zip(headers,widths),1):
    cell = ws.cell(1,c,h)
    cell.fill = HDR
    cell.font = Font(bold=True,color='FFFFFF',size=9)
    cell.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
    cell.border = thin
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[1].height = 36

for r_idx,row in enumerate(df.itertuples(index=False),2):
    est = row.ESTATUS
    fill = OK if est=='CUADRADO' else (WARN if est=='FACTURA EXCEDE' else DANG)
    vals = [row.SEMANA,row.GASOLINERA,row.REPORTES_REGISTROS,row.REPORTES_LITROS,row.REPORTES_IMPORTE,
            row.FACTURAS_REGISTROS,row.FACTURAS_LITROS,row.FACTURAS_IMPORTE,
            row.DIFERENCIA_LITROS,row.DIFERENCIA_IMPORTE,est]
    for c,v in enumerate(vals,1):
        cell = ws.cell(r_idx,c,v)
        cell.border = thin
        cell.fill = fill
        cell.alignment = Alignment(vertical='center')
        if c in (5,8,10): cell.number_format = '#,##0.00'
        elif c in (4,7,9): cell.number_format = '#,##0.000'
        if c==11:
            if est=='CUADRADO': cell.font = Font(bold=True,color='1E8449')
            elif est=='FALTA FACTURA': cell.font = Font(bold=True,color='C0392B')
            else: cell.font = Font(bold=True,color='856404')

r_tot = len(df)+2
ws.cell(r_tot,1,'GRAN TOTAL').font = Font(bold=True,color='FFFFFF')
ws.cell(r_tot,1).fill = PatternFill('solid',fgColor='1A1A2E')
ws.merge_cells(start_row=r_tot,start_column=1,end_row=r_tot,end_column=2)
for c,col in [(5,'REPORTES_IMPORTE'),(8,'FACTURAS_IMPORTE'),(10,'DIFERENCIA_IMPORTE')]:
    tc = ws.cell(r_tot,c,round(df[col].sum(),2))
    tc.font = Font(bold=True,color='FFFF00')
    tc.fill = PatternFill('solid',fgColor='1A1A2E')
    tc.number_format = '#,##0.00'
    tc.border = thin

wb.save(maestro_path)
print('OK:', maestro_path)
print(df.to_string())
