import openpyxl, re

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)
ws = wb['Estado de Cuenta']

def clean_fol(f):
    m = re.findall(r'[0-9]+', str(f))
    return "".join(m) if m else str(f)

# Extraer listas de facturas por semana en el Estado de Cuenta
semanas_gs = {}
curr_sem = None

for r in range(31, ws.max_row + 1):
    val_a = str(ws.cell(r, 1).value or '').strip()
    val_b = str(ws.cell(r, 2).value or '').strip()
    val_c = str(ws.cell(r, 3).value or '').strip()
    val_d = str(ws.cell(r, 4).value or '').strip()
    
    m_sem = re.search(r'Semana\s+(\d+)', val_a, re.IGNORECASE)
    if m_sem:
        curr_sem = int(m_sem.group(1))
        semanas_gs[curr_sem] = {'gasolina': [], 'diesel': [], 'a_favor': [], 'monto_edo': 0.0}
        continue
    
    if curr_sem and val_b:
        tipo = val_b.lower()
        facturas_raw = val_c
        monto = float(ws.cell(r, 4).value or 0.0) if str(ws.cell(r, 4).value or '').replace('.','',1).isdigit() else 0.0
        
        # Extraer folios
        # split by commas, ' Y ', spaces
        tokens = re.split(r'[,yY\s]+', facturas_raw)
        folios = [clean_fol(t) for t in tokens if clean_fol(t) and len(clean_fol(t)) >= 3]
        
        if 'favor' in tipo:
            semanas_gs[curr_sem]['a_favor'].extend(folios)
        elif 'gasolina' in tipo:
            semanas_gs[curr_sem]['gasolina'].extend(folios)
        elif 'diesel' in tipo:
            semanas_gs[curr_sem]['diesel'].extend(folios)
        elif 'total' in tipo:
            semanas_gs[curr_sem]['monto_edo'] = monto

print("=== FACTURAS EXTRAIDAS DE GOOGLE SHEETS ESTADO DE CUENTA ===")
for s, d in semanas_gs.items():
    print(f"\nSemana {s} (Monto Edo: ${d['monto_edo']:,.2f}):")
    print(f"  Gasolina ({len(d['gasolina'])}): {d['gasolina']}")
    print(f"  Diesel ({len(d['diesel'])}): {d['diesel']}")
    print(f"  A favor ({len(d['a_favor'])}): {d['a_favor']}")
