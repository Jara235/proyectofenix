import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

filename = 'c:/Users/JOSE/Desktop/Proyecto fenix/Estimacion_Combustible_MexicoToluca_Lerma_Ene_Jun.xlsx'
wb = openpyxl.load_workbook(filename)

if "Bitácoras Mantenimiento Lerma" in wb.sheetnames:
    del wb["Bitácoras Mantenimiento Lerma"]

ws = wb.create_sheet(title="Bitácoras Mantenimiento Lerma")
ws.views.sheetView[0].showGridLines = True

def fill(hex_color): return PatternFill("solid", fgColor=hex_color)
def font(bold=False, italic=False, color='000000', size=10, name='Calibri'): return Font(bold=bold, italic=italic, color=color, size=size, name=name)
def align(h='center', v='center', wrap=False): return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def thin_border():
    s = Side(style='thin', color='CBD5E1')
    return Border(left=s, right=s, top=s, bottom=s)

WHITE = 'FFFFFF'
NAVY = '0F172A'
TEAL_HDR = '0F766E'
BLUE_HDR = '1E3A8A'

# Header Title
ws.merge_cells('A1:L1')
ws['A1'] = "🛠️ PROGRAMA Y BITÁCORAS DE MANTENIMIENTO PREVENTIVO DE MAQUINARIA PESADA"
ws['A1'].fill = fill(NAVY); ws['A1'].font = font(bold=True, color=WHITE, size=13); ws['A1'].alignment = align('left', 'center')
ws.row_dimensions[1].height = 28

ws.merge_cells('A2:L2')
ws['A2'] = "Obra: Lerma - Tres Marías | Periodo: Marzo a Agosto 2026 (6 Meses) | Intervalos: 200h, 1,000h, 2,000h e Inspección 8 Días"
ws['A2'].fill = fill('1E293B'); ws['A2'].font = font(italic=True, color='94A3B8', size=9); ws['A2'].alignment = align('left', 'center')
ws.row_dimensions[2].height = 20

headers = [
    "N°", "EQUIPO / MAQUINARIA", "MARCA / MODELO", 
    "MARZO (hrs)", "ABRIL (hrs)", "MAYO (hrs)", "JUNIO (hrs)", "JULIO (hrs)", "AGOSTO (hrs)", 
    "TOTAL HORAS (6 MESES)", "BITÁCORAS 200h", "MESES DE EJECUCIÓN (BITÁCORAS)"
]

ws.merge_cells('A4:L4')
ws['A4'] = "📋 1. HORAS TRABAJADAS POR MES Y BITÁCORAS DE 200 HORAS REQUERIDAS (MARZO A AGOSTO 2026)"
ws['A4'].fill = fill(TEAL_HDR); ws['A4'].font = font(bold=True, color=WHITE, size=10); ws['A4'].alignment = align('left', 'center')
ws.row_dimensions[4].height = 22

for c_i, h in enumerate(headers, 1):
    c = ws.cell(row=5, column=c_i, value=h)
    c.fill = fill('334155'); c.font = font(bold=True, color=WHITE, size=8); c.alignment = align('center', 'center', wrap=True); c.border = thin_border()
ws.row_dimensions[5].height = 24

data_mant = [
    (1, "Perfiladora Rodatec (Principal)", "RODATEC RX600E", 38.9, 32.1, 44.3, 21.8, 47.4, 31.8, 216.3, 1, "Agosto (alcanza 200h)"),
    (2, "Perfiladora Rodatec (Apoyo)", "RODATEC RX600-4", 19.7, 16.3, 22.5, 11.1, 24.1, 16.1, 109.8, 0, "Acumula 110h (Requiere horómetro previo)"),
    (3, "Pavimentadora Vögele (Principal)", "VÖGELE SUPER 1800-3", 78.9, 65.1, 90.0, 44.2, 96.2, 64.6, 439.0, 2, "Mayo (200h) y Agosto (400h)"),
    (4, "Pavimentadora Vögele (Apoyo)", "VÖGELE SUPER 1800-3i", 38.3, 31.6, 43.7, 21.4, 46.7, 31.3, 213.0, 1, "Agosto (alcanza 200h)"),
    (5, "Doble Rodillo Hamm (Principal)", "HAMM HD+120VV", 81.3, 67.1, 92.6, 45.5, 99.0, 66.5, 452.0, 2, "Mayo (200h) y Agosto (400h)"),
    (6, "Doble Rodillo Caterpillar (Apoyo)", "CAT CB66B", 44.1, 36.4, 50.3, 24.7, 53.8, 36.1, 245.4, 1, "Julio (alcanza 200h)"),
    (7, "Neumático Volvo (Principal)", "VOLVO PT240R", 60.4, 49.8, 68.8, 33.8, 73.6, 49.4, 335.8, 1, "Junio (alcanza 200h)"),
    (8, "Neumático Dynapac (Apoyo)", "DYNAPAC CP271", 39.5, 32.6, 45.0, 22.1, 48.1, 32.3, 219.6, 1, "Agosto (alcanza 200h)"),
    (9, "Retroexcavadora Case (Unidad 1)", "CASE 580N", 76.6, 63.2, 87.3, 42.9, 93.4, 62.7, 426.1, 2, "Mayo (200h) y Agosto (400h)"),
    (10, "Retroexcavadora Case (Unidad 2)", "CASE 580N", 39.5, 32.6, 45.0, 22.1, 48.1, 32.3, 219.6, 1, "Agosto (alcanza 200h)"),
    (11, "Barredora Laymor", "LAYMOR SM400", 48.8, 40.2, 55.6, 27.3, 59.4, 39.9, 271.2, 1, "Julio (alcanza 200h)"),
    (12, "Barredora Super Broom / Broce", "SUPER BROOM 2004", 31.3, 25.9, 35.7, 17.6, 38.2, 25.6, 174.3, 0, "Acumula 174h (Requiere horómetro previo)"),
]

r_curr = 6
for idx, d in enumerate(data_mant):
    bg = 'FFFFFF' if idx % 2 == 0 else 'F8FAFC'
    ws.cell(row=r_curr, column=1, value=d[0]).font = font(bold=True, size=9)
    ws.cell(row=r_curr, column=2, value=d[1]).font = font(bold=True, size=9)
    ws.cell(row=r_curr, column=3, value=d[2]).font = font(size=9)
    
    for c_i in range(4, 10):
        ws.cell(row=r_curr, column=c_i, value=d[c_i-1]).number_format = '#,##0.0'
        
    ws.cell(row=r_curr, column=10, value=d[9]).number_format = '#,##0.0'
    ws.cell(row=r_curr, column=10).font = font(bold=True, size=9)
    
    c_bit = ws.cell(row=r_curr, column=11, value=d[10])
    c_bit.number_format = '0'
    c_bit.font = font(bold=True, size=9)
    c_bit.fill = fill('FEF08A' if d[10] > 0 else 'E2E8F0')
    
    ws.cell(row=r_curr, column=12, value=d[11]).font = font(size=8, italic=True)
    
    for c in range(1, 13):
        cell = ws.cell(row=r_curr, column=c)
        if c != 11: cell.fill = fill(bg)
        cell.border = thin_border()
        if c in [1, 11]: cell.alignment = align('center', 'center')
        elif c in [2, 3, 12]: cell.alignment = align('left', 'center')
        else: cell.alignment = align('right', 'center')
    ws.row_dimensions[r_curr].height = 20
    r_curr += 1

# Total Row
ws.cell(row=r_curr, column=1, value="").fill = fill(TEAL_HDR)
ws.cell(row=r_curr, column=2, value="TOTAL FLOTA (12 EQUIPOS)").fill = fill(TEAL_HDR)
ws.cell(row=r_curr, column=2).font = font(bold=True, color=WHITE, size=9); ws.cell(row=r_curr, column=2).alignment = align('left', 'center')
ws.cell(row=r_curr, column=3, value="Lerma - Tres Marías").fill = fill(TEAL_HDR)
ws.cell(row=r_curr, column=3).font = font(bold=True, color=WHITE, size=8); ws.cell(row=r_curr, column=3).alignment = align('center', 'center')

tot_mar = sum(d[3] for d in data_mant)
tot_abr = sum(d[4] for d in data_mant)
tot_may = sum(d[5] for d in data_mant)
tot_jun = sum(d[6] for d in data_mant)
tot_jul = sum(d[7] for d in data_mant)
tot_ago = sum(d[8] for d in data_mant)
tot_6m = sum(d[9] for d in data_mant)
tot_bit = sum(d[10] for d in data_mant)

ws.cell(row=r_curr, column=4, value=tot_mar).number_format = '#,##0.0'
ws.cell(row=r_curr, column=5, value=tot_abr).number_format = '#,##0.0'
ws.cell(row=r_curr, column=6, value=tot_may).number_format = '#,##0.0'
ws.cell(row=r_curr, column=7, value=tot_jun).number_format = '#,##0.0'
ws.cell(row=r_curr, column=8, value=tot_jul).number_format = '#,##0.0'
ws.cell(row=r_curr, column=9, value=tot_ago).number_format = '#,##0.0'
ws.cell(row=r_curr, column=10, value=tot_6m).number_format = '#,##0.0'

c_tot_b = ws.cell(row=r_curr, column=11, value=tot_bit)
c_tot_b.number_format = '0'
c_tot_b.font = font(bold=True, color='000000', size=9)
c_tot_b.fill = fill('FDE047')

ws.cell(row=r_curr, column=12, value="13 Bitácoras Oficiales de 200 Horas a elaborar").fill = fill(TEAL_HDR)
ws.cell(row=r_curr, column=12).font = font(bold=True, color=WHITE, size=8); ws.cell(row=r_curr, column=12).alignment = align('left', 'center')

for c in range(1, 13):
    cell = ws.cell(row=r_curr, column=c)
    if c not in [11]:
        cell.fill = fill(TEAL_HDR); cell.font = font(bold=True, color=WHITE, size=9)
    cell.border = thin_border()
    if c in [1, 11]: cell.alignment = align('center', 'center')
    elif c in [2, 3, 12]: cell.alignment = align('left', 'center')
    else: cell.alignment = align('right', 'center')
ws.row_dimensions[r_curr].height = 22

# Widths
for col in ws.columns:
    col_letter = get_column_letter(col[0].column)
    if col_letter in ['B', 'L']:
        ws.column_dimensions[col_letter].width = 36
    elif col_letter == 'C':
        ws.column_dimensions[col_letter].width = 22
    else:
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str: val_str = max(val_str.split('\n'), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

wb.save(filename)
print("Saved Maintenance sheet in Excel successfully!")
