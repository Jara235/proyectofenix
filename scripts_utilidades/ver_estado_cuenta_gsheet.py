import openpyxl

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)

for sname in ['Estado de Cuenta', 'Hoja3', 'Hoja4', 'SEMANA 28', 'SEMANA 29', 'SEMANA 30', 'SEMANA 31']:
    if sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n=======================================================")
        print(f"SHEET: {sname} (rows={ws.max_row}, cols={ws.max_column})")
        print(f"=======================================================")
        for r in range(1, min(ws.max_row+1, 40)):
            row_vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, min(ws.max_column+1, 15))]
            if any(row_vals):
                print(f"R{r:02d}: " + " | ".join([v for v in row_vals if v]))
