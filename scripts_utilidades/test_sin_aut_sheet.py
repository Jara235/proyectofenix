import psycopg2
from psycopg2.extras import DictCursor
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import OrderedDict
import io, re, datetime as dt_module

conn = psycopg2.connect(dbname='fenix_db', user='postgres', password='Bupito*268', host='localhost')
cur = conn.cursor(cursor_factory=DictCursor)

def fill(hex_color): return PatternFill("solid", fgColor=hex_color)
def font(bold=False, italic=False, color='000000', size=10, name='Calibri'): return Font(bold=bold, italic=italic, color=color, size=size, name=name)
def align(h='center', v='center', wrap=False): return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def thin_border():
    s = Side(style='thin', color='CBD5E1')
    return Border(left=s, right=s, top=s, bottom=s)

WHITE = 'FFFFFF'
COL_HDR_FILL = '1E293B'
COL_ALT1 = 'F8FAFC'
COL_ALT2 = 'EFF6FF'
COL_TOT_FILL = '0F172A'

wb = openpyxl.Workbook()
ws_sin_aut = wb.active
ws_sin_aut.title = 'TAGs Sin Autorización'
ws_sin_aut.views.sheetView[0].showGridLines = True

# Title
ws_sin_aut.merge_cells("A1:K1")
t_cell = ws_sin_aut.cell(row=1, column=1, value="📋 RELACIÓN DE TAGS SIN AUTORIZACIÓN — PENDIENTES DE SOLICITUD DE TOPE")
t_cell.fill = fill('991B1B') # Red Dark
t_cell.font = font(bold=True, color=WHITE, size=12)
t_cell.alignment = align('center', 'center')
ws_sin_aut.row_dimensions[1].height = 28

ws_sin_aut.merge_cells("A2:K2")
sub_cell = ws_sin_aut.cell(row=2, column=1, value=f"Catálogo de dispositivos y unidades con saldo/consumo sin presupuesto asignado ($0.00) para trámite y asignación de presupuesto | Generado: {dt_module.datetime.now().strftime('%Y-%m-%d %H:%M')}")
sub_cell.fill = fill('1E293B'); sub_cell.font = font(color='94A3B8', size=9); sub_cell.alignment = align('center', 'center')
ws_sin_aut.row_dimensions[2].height = 20

headers = [
    'RESPONSABLE / CONDUCTOR',
    'NÚMERO DE TAG',
    'COMPAÑÍA',
    'NO. ECONÓMICO',
    'PLACAS',
    'TIPO DE UNIDAD',
    'ESTATUS DE AUTORIZACIÓN',
    'CONSUMO ACUMULADO ($)',
    'TOTAL PEAJES',
    'TOPE SEMANAL SOLICITADO ($)',
    'JUSTIFICACIÓN / OBSERVACIONES'
]

ws_sin_aut.row_dimensions[4].height = 24
for col_i, h_txt in enumerate(headers, start=1):
    c = ws_sin_aut.cell(row=4, column=col_i, value=h_txt)
    c.fill = fill('1E293B')
    c.font = font(bold=True, color=WHITE, size=9)
    c.alignment = align('center', 'center', wrap=True)
    c.border = thin_border()

# Query all tags sin autorizacion
cur.execute("""
    SELECT a.id, a.empresa, a.tag, a.responsable, a.no_economico, a.placas, a.tipo_unidad, a.monto_autorizado, a.estatus,
           COALESCE(SUM(ABS(m.importe)), 0) as consumo_acumulado,
           COUNT(m.id) as total_pasadas
    FROM tags.autorizaciones a
    LEFT JOIN tags.movimientos m ON (
        m.tag = a.tag 
        OR (m.tag IS NOT NULL AND a.tag IS NOT NULL AND LENGTH(a.tag) >= 6 AND (m.tag ILIKE '%%' || a.tag || '%%' OR a.tag ILIKE '%%' || m.tag || '%%'))
    )
    WHERE a.monto_autorizado IS NULL OR a.monto_autorizado = 0
    GROUP BY a.id, a.empresa, a.tag, a.responsable, a.no_economico, a.placas, a.tipo_unidad, a.monto_autorizado, a.estatus
    ORDER BY (CASE WHEN COALESCE(SUM(ABS(m.importe)), 0) > 0 THEN 0 ELSE 1 END), COALESCE(SUM(ABS(m.importe)), 0) DESC, a.empresa, a.responsable;
""")
rows = cur.fetchall()

curr_r = 5
tot_consumo_sin_aut = 0.0
tot_peajes_sin_aut = 0

for idx, r in enumerate(rows):
    cons = float(r['consumo_acumulado'] or 0)
    peajes = int(r['total_pasadas'] or 0)
    tot_consumo_sin_aut += cons
    tot_peajes_sin_aut += peajes

    bg = 'FEF2F2' if cons > 0 else (COL_ALT1 if idx % 2 == 0 else COL_ALT2)

    # 1. Responsable
    c = ws_sin_aut.cell(row=curr_r, column=1, value=r['responsable'] or '-')
    c.fill = fill(bg); c.font = font(bold=True, size=9); c.alignment = align('left', 'center'); c.border = thin_border()

    # 2. Tag
    c = ws_sin_aut.cell(row=curr_r, column=2, value=r['tag'] or '-')
    c.fill = fill(bg); c.font = font(bold=True, color='0284C7', size=9); c.alignment = align('center', 'center'); c.border = thin_border()

    # 3. Empresa / Compañía
    c = ws_sin_aut.cell(row=curr_r, column=3, value=r['empresa'] or '-')
    c.fill = fill('DBEAFE' if r['empresa'] == 'JDJ' else 'FEF3C7')
    c.font = font(bold=True, color='1E3A8A' if r['empresa'] == 'JDJ' else '92400E', size=9)
    c.alignment = align('center', 'center'); c.border = thin_border()

    # 4. No. Economico
    c = ws_sin_aut.cell(row=curr_r, column=4, value=r['no_economico'] or '-')
    c.fill = fill(bg); c.font = font(size=9); c.alignment = align('center', 'center'); c.border = thin_border()

    # 5. Placas
    c = ws_sin_aut.cell(row=curr_r, column=5, value=r['placas'] or '-')
    c.fill = fill(bg); c.font = font(size=9); c.alignment = align('center', 'center'); c.border = thin_border()

    # 6. Tipo de Unidad
    c = ws_sin_aut.cell(row=curr_r, column=6, value=r['tipo_unidad'] or '-')
    c.fill = fill(bg); c.font = font(size=9); c.alignment = align('left', 'center'); c.border = thin_border()

    # 7. Estatus Autorizacion
    c = ws_sin_aut.cell(row=curr_r, column=7, value='🔴 SIN AUTORIZACIÓN ($0.00)')
    c.fill = fill('FEE2E2'); c.font = font(bold=True, color='991B1B', size=8); c.alignment = align('center', 'center'); c.border = thin_border()

    # 8. Consumo Acumulado
    c = ws_sin_aut.cell(row=curr_r, column=8, value=cons)
    c.fill = fill(bg); c.font = font(bold=True, color='DC2626' if cons > 0 else '64748B', size=9)
    c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

    # 9. Total Peajes
    c = ws_sin_aut.cell(row=curr_r, column=9, value=peajes)
    c.fill = fill(bg); c.font = font(size=9); c.alignment = align('center', 'center'); c.border = thin_border(); c.number_format = '#,##0'

    # 10. Tope Solicitado (Empty for manual fill)
    c = ws_sin_aut.cell(row=curr_r, column=10, value="")
    c.fill = fill('FFFFFF'); c.font = font(bold=True, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

    # 11. Justificacion
    c = ws_sin_aut.cell(row=curr_r, column=11, value="Requiere asignación de tope" if cons > 0 else "Dispositivo sin tope asignado")
    c.fill = fill(bg); c.font = font(italic=True, color='64748B', size=8); c.alignment = align('left', 'center'); c.border = thin_border()

    ws_sin_aut.row_dimensions[curr_r].height = 20
    curr_r += 1

# Totals row
ws_sin_aut.cell(row=curr_r, column=1, value="TOTAL GENERAL").fill = fill('0F172A')
ws_sin_aut.cell(row=curr_r, column=1).font = font(bold=True, color=WHITE, size=9); ws_sin_aut.cell(row=curr_r, column=1).border = thin_border()

ws_sin_aut.cell(row=curr_r, column=2, value=f"{len(rows)} Registros").fill = fill('0F172A')
ws_sin_aut.cell(row=curr_r, column=2).font = font(bold=True, color=WHITE, size=9); ws_sin_aut.cell(row=curr_r, column=2).alignment = align('center', 'center'); ws_sin_aut.cell(row=curr_r, column=2).border = thin_border()

for c_i in [3, 4, 5, 6, 7]:
    c = ws_sin_aut.cell(row=curr_r, column=c_i, value="")
    c.fill = fill('0F172A'); c.border = thin_border()

c = ws_sin_aut.cell(row=curr_r, column=8, value=tot_consumo_sin_aut)
c.fill = fill('DC2626'); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('right', 'center'); c.border = thin_border(); c.number_format = '"$"#,##0.00'

c = ws_sin_aut.cell(row=curr_r, column=9, value=tot_peajes_sin_aut)
c.fill = fill('0F172A'); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center'); c.border = thin_border(); c.number_format = '#,##0'

ws_sin_aut.cell(row=curr_r, column=10, value="").fill = fill('0F172A'); ws_sin_aut.cell(row=curr_r, column=10).border = thin_border()
ws_sin_aut.cell(row=curr_r, column=11, value="").fill = fill('0F172A'); ws_sin_aut.cell(row=curr_r, column=11).border = thin_border()
ws_sin_aut.row_dimensions[curr_r].height = 22

# Column widths
ws_sin_aut.column_dimensions['A'].width = 32 # Responsable
ws_sin_aut.column_dimensions['B'].width = 18 # Tag
ws_sin_aut.column_dimensions['C'].width = 14 # Compañía
ws_sin_aut.column_dimensions['D'].width = 24 # No. Económico
ws_sin_aut.column_dimensions['E'].width = 14 # Placas
ws_sin_aut.column_dimensions['F'].width = 22 # Tipo Unidad
ws_sin_aut.column_dimensions['G'].width = 25 # Estatus
ws_sin_aut.column_dimensions['H'].width = 22 # Consumo Acumulado
ws_sin_aut.column_dimensions['I'].width = 14 # Total Peajes
ws_sin_aut.column_dimensions['J'].width = 24 # Tope Solicitado
ws_sin_aut.column_dimensions['K'].width = 30 # Observaciones

wb.save('test_sin_aut.xlsx')
print("Saved test_sin_aut.xlsx successfully.")
conn.close()
