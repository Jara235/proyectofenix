import openpyxl
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano ultimo(1).xlsx')
ws = wb['Captura_Facturas_Mobil']
factura_col = None
semana_col = None
for col in range(1, 20):
    val = ws.cell(row=3, column=col).value
    if val == 'No. Factura': factura_col = col
    elif val == 'Semana (Helper)': semana_col = col
if factura_col and semana_col:
    for r in range(4, ws.max_row + 1):
        factura = ws.cell(row=r, column=factura_col).value
        if factura in ['A-10270', 'A-10271']:
            ws.cell(row=r, column=semana_col).value = 26
            print(f'Updated {factura} to week 26')
wb.save('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_Actualizado.xlsx')
print('Excel updated successfully')

