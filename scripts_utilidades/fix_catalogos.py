import openpyxl

excel_path = r"c:\Users\JOSE\Desktop\Proyecto fenix\MAESTRO_CONTROL_DIESEL_NUEVO.xlsx"

try:
    wb = openpyxl.load_workbook(excel_path)
    sheet_cat = wb["CATALOGOS"]
    
    # Mover Codigos de Movimiento (lo que yo puse en E y F) hacia las columnas J y K
    sheet_cat['J1'] = "TIPO_OPERACION"
    sheet_cat['K1'] = "CODIGO_OPERACION"
    movimientos = [
        ("CARGA TANQUE", "CT"),
        ("CARGA MARIMBA", "CM"),
        ("CARGA MAQUINARIA", "CMQ"),
        ("SOLICITUD", "SO"),
        ("FACTURA", "FA")
    ]
    for i, (mov, code) in enumerate(movimientos, start=2):
        sheet_cat.cell(row=i, column=10).value = mov
        sheet_cat.cell(row=i, column=11).value = code
        
    # Restaurar los encabezados originales (aparentes) en E y F
    sheet_cat['E1'] = "NUMERO_ECONOMICO"
    sheet_cat['F1'] = "OPERADORES"
    
    # Limpiar las filas 2 a 6 en E y F que sobrescribí por error
    for row in range(2, 7):
        sheet_cat.cell(row=row, column=5).value = ""
        sheet_cat.cell(row=row, column=6).value = ""

    # Arreglar el código PAH que salió XX
    # Buscar "Huixquilucan" en cualquier parte de la celda de la columna A
    for row in range(2, 50):
        obra = str(sheet_cat.cell(row=row, column=1).value or "").strip().upper()
        if not obra:
            continue
            
        codigo_actual = sheet_cat.cell(row=row, column=2).value
        if "HUIXQUILUCAN" in obra or codigo_actual == "XX":
            # Reprocesar a prueba de balas
            if "PEGASO" in obra and "MAQUINARIA" in obra: code = "MP"
            elif "PEGASO" in obra: code = "PAP"
            elif "TOLUCA" in obra and "MEX" in obra: code = "MT"
            elif "TOLUCA" in obra and "MÉX" in obra: code = "MT"
            elif "LERMA" in obra or "TRES MAR" in obra or "L3M" in obra: code = "L3M"
            elif "BACHEO" in obra: code = "BT"
            elif "HUIXQUILUCAN" in obra: code = "PAH"
            elif "CHAMAPA" in obra: code = "CL"
            elif "DRAGONES" in obra: code = "DRA"
            elif "JALISCO" in obra: code = "JAL"
            elif "PROVIDENCIA" in obra: code = "PRO"
            elif "NO APLICA" in obra: code = "NA"
            else: code = "XX"
            
            sheet_cat.cell(row=row, column=2).value = code

    wb.save(excel_path)
    print("Correcciones aplicadas correctamente en CATALOGOS.")

except Exception as e:
    print(f"Error: {e}")
