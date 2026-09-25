import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import pandas as pd
import re

excel_path = r"c:\Users\JOSE\Desktop\Proyecto fenix\MAESTRO_CONTROL_DIESEL_NUEVO.xlsx"

try:
    print("Cargando archivo...")
    wb = openpyxl.load_workbook(excel_path)
    
    # Extraer Obras
    obras = []
    if "llave" in wb.sheetnames:
        llave_sheet = wb["llave"]
        for row in range(2, llave_sheet.max_row + 1):
            val = llave_sheet.cell(row=row, column=1).value
            if val and str(val).strip():
                obras.append(str(val).strip())
    else:
        print("No se encontro hoja 'llave'")
        
    # Extraer Semanas
    semanas = set()
    if "BD_DIESEL" in wb.sheetnames:
        sh = wb["BD_DIESEL"]
        for row in range(2, sh.max_row + 1):
            val = sh.cell(row=row, column=3).value
            if val and isinstance(val, str) and val.startswith("Semana"):
                semanas.add(val)
    if "BD_FACTURAS" in wb.sheetnames:
        sh = wb["BD_FACTURAS"]
        for row in range(2, sh.max_row + 1):
            val = sh.cell(row=row, column=4).value
            if val and isinstance(val, str) and val.startswith("Semana"):
                semanas.add(val)
                
    # Parse num
    def get_sem_num(s):
        m = re.search(r'\d+', s)
        return int(m.group()) if m else 0
        
    semanas = sorted(list(semanas), key=get_sem_num)
    if not semanas:
        semanas = ["Semana 24", "Semana 25", "Semana 26"]
        
    # Crear hoja
    if "TABLERO_CONCILIACION" in wb.sheetnames:
        del wb["TABLERO_CONCILIACION"]
    
    ws = wb.create_sheet("TABLERO_CONCILIACION", 0)
    
    # Estilos
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    # Fila 1: Titulo
    ws.cell(row=1, column=1).value = "OBRA / DESTINO"
    ws.cell(row=1, column=1).font = header_font
    ws.cell(row=1, column=1).fill = header_fill
    ws.cell(row=1, column=1).alignment = center_align
    ws.column_dimensions['A'].width = 30
    
    col_idx = 2
    for sem in semanas:
        ws.cell(row=1, column=col_idx).value = sem.upper()
        ws.cell(row=1, column=col_idx).font = header_font
        ws.cell(row=1, column=col_idx).fill = header_fill
        ws.cell(row=1, column=col_idx).alignment = center_align
        ws.merge_cells(start_row=1, start_column=col_idx, end_row=1, end_column=col_idx+4)
        
        # Subheaders
        headers = ["LTS FACTURADOS", "IMPORTE FACT.", "LTS CONSUMIDOS", "IMPORTE CONS.", "DIFERENCIA LTS"]
        for i, h in enumerate(headers):
            c = ws.cell(row=2, column=col_idx + i)
            c.value = h
            c.font = Font(bold=True, size=9)
            c.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            c.alignment = center_align
            c.border = thin_border
            ws.column_dimensions[c.column_letter].width = 15
            
        col_idx += 5
        
    # Llenar datos
    row_idx = 3
    for obra in obras:
        ws.cell(row=row_idx, column=1).value = obra
        ws.cell(row=row_idx, column=1).border = thin_border
        
        c_idx = 2
        for sem in semanas:
            col_ltr_fact_lts = openpyxl.utils.get_column_letter(c_idx)
            col_ltr_fact_imp = openpyxl.utils.get_column_letter(c_idx + 1)
            col_ltr_cons_lts = openpyxl.utils.get_column_letter(c_idx + 2)
            col_ltr_cons_imp = openpyxl.utils.get_column_letter(c_idx + 3)
            col_ltr_dif = openpyxl.utils.get_column_letter(c_idx + 4)
            
            # Facturados Lts (BD_FACTURAS G)
            ws.cell(row=row_idx, column=c_idx).value = f'=SUMIFS(BD_FACTURAS!$G:$G, BD_FACTURAS!$F:$F, $A{row_idx}, BD_FACTURAS!$D:$D, "{sem}")'
            ws.cell(row=row_idx, column=c_idx).number_format = '#,##0.00'
            
            # Facturados Imp (BD_FACTURAS K)
            ws.cell(row=row_idx, column=c_idx+1).value = f'=SUMIFS(BD_FACTURAS!$K:$K, BD_FACTURAS!$F:$F, $A{row_idx}, BD_FACTURAS!$D:$D, "{sem}")'
            ws.cell(row=row_idx, column=c_idx+1).number_format = '"$"#,##0.00'
            
            # Consumidos Lts (BD_DIESEL I)
            ws.cell(row=row_idx, column=c_idx+2).value = f'=SUMIFS(BD_DIESEL!$I:$I, BD_DIESEL!$F:$F, $A{row_idx}, BD_DIESEL!$C:$C, "{sem}")'
            ws.cell(row=row_idx, column=c_idx+2).number_format = '#,##0.00'
            
            # Consumidos Imp (BD_DIESEL K)
            ws.cell(row=row_idx, column=c_idx+3).value = f'=SUMIFS(BD_DIESEL!$K:$K, BD_DIESEL!$F:$F, $A{row_idx}, BD_DIESEL!$C:$C, "{sem}")'
            ws.cell(row=row_idx, column=c_idx+3).number_format = '"$"#,##0.00'
            
            # Diferencia (Fact - Cons)
            ws.cell(row=row_idx, column=c_idx+4).value = f'={col_ltr_fact_lts}{row_idx}-{col_ltr_cons_lts}{row_idx}'
            ws.cell(row=row_idx, column=c_idx+4).number_format = '#,##0.00'
            
            for i in range(5):
                ws.cell(row=row_idx, column=c_idx+i).border = thin_border
                
            c_idx += 5
        row_idx += 1
        
    print("Guardando...")
    wb.save(excel_path)
    print("Tablero generado con exito.")

except Exception as e:
    print(f"Error: {e}")
