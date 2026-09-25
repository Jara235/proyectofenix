import requests
import openpyxl

try:
    url = 'http://127.0.0.1:5002/api/admin/tags/reporte/excel'
    print(f"Requesting {url}...")
    resp = requests.get(url, timeout=15)
    print("Status code:", resp.status_code)
    print("Content length:", len(resp.content))
    print("Content-type:", resp.headers.get('content-type'))
    if resp.status_code == 200:
        filename = 'Reporte_TAGS_Test.xlsx'
        with open(filename, 'wb') as f:
            f.write(resp.content)
        print("Excel file successfully downloaded!")

        # Inspect sheets
        wb = openpyxl.load_workbook(filename)
        print("Sheets in generated Excel:", wb.sheetnames)

        # Inspect 'TAGs Sin Autorización'
        if 'TAGs Sin Autorización' in wb.sheetnames:
            ws = wb['TAGs Sin Autorización']
            print(f"\n--- Sheet 'TAGs Sin Autorización' (Rows: {ws.max_row}, Cols: {ws.max_column}) ---")
            for r in range(1, min(15, ws.max_row + 1)):
                vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
                print(f"Row {r:2d}: {vals}")

        # Inspect 'SEMANA 32'
        if 'SEMANA 32' in wb.sheetnames:
            ws = wb['SEMANA 32']
            print(f"\n--- Sheet 'SEMANA 32' (Rows: {ws.max_row}, Cols: {ws.max_column}) ---")
            for r in range(max(1, ws.max_row - 15), ws.max_row + 1):
                vals = [ws.cell(r, c).value for c in range(1, min(7, ws.max_column + 1))]
                print(f"Row {r:2d}: {vals}")
    else:
        print("Error response:", resp.text[:500])
except Exception as e:
    import traceback
    print("Request exception:", e)
    traceback.print_exc()
