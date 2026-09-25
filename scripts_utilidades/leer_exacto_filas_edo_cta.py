import openpyxl, re

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)
ws = wb['Estado de Cuenta']

for r in range(32, 62):
    row_data = [f"Col{c}: {ws.cell(r, c).value}" for c in [2, 4, 6, 8] if ws.cell(r, c).value is not None]
    if row_data:
        print(f"Row {r:02d} -> " + " | ".join(row_data))
