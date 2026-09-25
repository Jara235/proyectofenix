import openpyxl
print('--- MAESTRO ---')
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/MAESTRO_CONTROL_GASOLINA_NUEVO.xlsx', data_only=True)
print(wb.sheetnames)
print('--- SEMANA 26 ---')
wb2 = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/gasolina/SEMANA 26.xlsx', data_only=True)
print(wb2.sheetnames)

