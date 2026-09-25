import openpyxl
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/MAESTRO_CONTROL_DIESEL_NUEVO.xlsx', data_only=False)
ws = wb['BD_DIESEL']
print(ws['A3'].value)
print(ws['A4'].value)

