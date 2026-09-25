import openpyxl

wb = openpyxl.load_workbook('formatos/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_FINAL_V2.xlsx', data_only=True)

print("=== HOJA: Maquinaria_Inventario ===")
ws = wb['Maquinaria_Inventario']
for r in range(1, 35):
    row_vals = [ws.cell(r, c).value for c in range(1, 15)]
    if any(row_vals):
        print(f"R{r:02d}:", [str(v) if v is not None else "" for v in row_vals])

print("\n=== HOJA: Catalogos ===")
ws2 = wb['Catalogos']
for r in range(1, 35):
    row_vals = [ws2.cell(r, c).value for c in range(1, 15)]
    if any(row_vals):
        print(f"R{r:02d}:", [str(v) if v is not None else "" for v in row_vals])
