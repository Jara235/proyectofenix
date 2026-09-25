import openpyxl

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)

for sheetname in ['Estado de Cuenta', 'Hoja4', 'Hoja3']:
    if sheetname in wb.sheetnames:
        ws = wb[sheetname]
        print(f"\n=======================================================")
        print(f"SHEET: {sheetname} ({ws.max_row} rows, {ws.max_column} cols)")
        print(f"=======================================================")
        for r in range(1, ws.max_row + 1):
            row_vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)]
            if any(row_vals):
                print(f"R{r:02d}: " + " | ".join([v for v in row_vals if v]))
