import openpyxl
import datetime
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')
ws = wb['Captura_Facturas_Mobil']

for r in range(5, ws.max_row + 1):
    dt = ws.cell(row=r, column=2).value
    if dt is not None:
        try:
            if isinstance(dt, str): dt = datetime.datetime.strptime(dt, '%Y-%m-%d %H:%M:%S')
            week = dt.isocalendar().week
            ws.cell(row=r, column=14).value = week
            print(f'Row {r} Week {week}')
        except Exception as e:
            print(f'Skipped {r}: {e}')

wb.save('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')
print('Done!')

