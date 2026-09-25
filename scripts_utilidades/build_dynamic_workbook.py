import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

wb = openpyxl.Workbook()

# Styles
def fill(hex_color): return PatternFill("solid", fgColor=hex_color)
def font(bold=False, italic=False, color='000000', size=10, name='Calibri'): return Font(bold=bold, italic=italic, color=color, size=size, name=name)
def align(h='center', v='center', wrap=False): return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def thin_border():
    s = Side(style='thin', color='CBD5E1')
    return Border(left=s, right=s, top=s, bottom=s)

WHITE = 'FFFFFF'
NAVY = '0F172A'
BLUE_HDR = '1E3A8A'
TEAL_HDR = '0F766E'
AMBER_HDR = 'B45309'
GRAY_BG = 'F8FAFC'

# ─────────────────────────────────────────────────────────────
# HOJA 1: ESTIMACIÓN DINÁMICA MULTIFACTORIAL (ENERO - JUNIO)
# ─────────────────────────────────────────────────────────────
ws1 = wb.active
ws1.title = "Estimación Dinámica Ene-Jun"
ws1.views.sheetView[0].showGridLines = True

# Title banner
ws1.merge_cells('A1:M1')
t1 = ws1['A1']
t1.value = "📊 ESTIMACIÓN DINÁMICA DE DIESEL Y GASOLINA (ENERO - JUNIO)"
t1.fill = fill(NAVY); t1.font = font(bold=True, color=WHITE, size=13); t1.alignment = align('left', 'center')
ws1.row_dimensions[1].height = 28

ws1.merge_cells('A2:M2')
sub1 = ws1['A2']
sub1.value = "Modelo Multifactorial: Ajustado por Clima/Lluvias, Mantenimiento/Descomposturas de Maquinaria y Calendario Laboral | Fénix 2.0"
sub1.fill = fill('1E293B'); sub1.font = font(italic=True, color='94A3B8', size=9); sub1.alignment = align('left', 'center')
ws1.row_dimensions[2].height = 20

# KPI Summary Cards
kpis = [
    ("TOTAL DIESEL ESTIMADO", "107,832 L", "$2,911,475.23 MXN", '2563EB'),
    ("TOTAL GASOLINA ESTIMADA", "20,181 L", "$482,320.94 MXN", '059669'),
    ("PRESUPUESTO TOTAL DINÁMICO", "128,013 L", "$3,393,796.17 MXN", 'D97706')
]

col_start = 1
for title_kpi, lts_kpi, imp_kpi, color_kpi in kpis:
    col_end = col_start + 3 if col_start < 9 else 13
    ws1.merge_cells(start_row=4, start_column=col_start, end_row=4, end_column=col_end)
    c = ws1.cell(row=4, column=col_start, value=title_kpi)
    c.fill = fill('F1F5F9'); c.font = font(bold=True, color='475569', size=9); c.alignment = align('center', 'center')
    
    ws1.merge_cells(start_row=5, start_column=col_start, end_row=5, end_column=col_end)
    c2 = ws1.cell(row=5, column=col_start, value=f"{lts_kpi}  |  {imp_kpi}")
    c2.fill = fill(color_kpi); c2.font = font(bold=True, color=WHITE, size=11); c2.alignment = align('center', 'center')
    
    for r in range(4, 6):
        for c_idx in range(col_start, col_end + 1):
            ws1.cell(row=r, column=c_idx).border = thin_border()
    col_start = col_end + 1

ws1.row_dimensions[4].height = 18
ws1.row_dimensions[5].height = 24

# Header Table
ws1.merge_cells('A7:A8'); ws1['A7'] = "MES"; ws1['A7'].fill = fill(NAVY); ws1['A7'].font = font(bold=True, color=WHITE, size=9); ws1['A7'].alignment = align('center', 'center')
ws1.merge_cells('B7:B8'); ws1['B7'] = "FASE OPERATIVA"; ws1['B7'].fill = fill(NAVY); ws1['B7'].font = font(bold=True, color=WHITE, size=9); ws1['B7'].alignment = align('center', 'center')
ws1.merge_cells('C7:C8'); ws1['C7'] = "FACTOR DIESEL"; ws1['C7'].fill = fill('475569'); ws1['C7'].font = font(bold=True, color=WHITE, size=8); ws1['C7'].alignment = align('center', 'center')

ws1.merge_cells('D7:F7'); ws1['D7'] = "MÉXICO - TOLUCA"; ws1['D7'].fill = fill(BLUE_HDR); ws1['D7'].font = font(bold=True, color=WHITE, size=9); ws1['D7'].alignment = align('center', 'center')
ws1['D8'] = "Diesel (L)"; ws1['D8'].fill = fill('3B82F6'); ws1['D8'].font = font(bold=True, color=WHITE, size=8); ws1['D8'].alignment = align('center', 'center')
ws1['E8'] = "Gasolina (L)"; ws1['E8'].fill = fill('3B82F6'); ws1['E8'].font = font(bold=True, color=WHITE, size=8); ws1['E8'].alignment = align('center', 'center')
ws1['F8'] = "Importe ($)"; ws1['F8'].fill = fill('1D4ED8'); ws1['F8'].font = font(bold=True, color=WHITE, size=8); ws1['F8'].alignment = align('center', 'center')

ws1.merge_cells('G7:I7'); ws1['G7'] = "LERMA - TRES MARÍAS"; ws1['G7'].fill = fill(TEAL_HDR); ws1['G7'].font = font(bold=True, color=WHITE, size=9); ws1['G7'].alignment = align('center', 'center')
ws1['G8'] = "Diesel (L)"; ws1['G8'].fill = fill('14B8A6'); ws1['G8'].font = font(bold=True, color=WHITE, size=8); ws1['G8'].alignment = align('center', 'center')
ws1['H8'] = "Gasolina (L)"; ws1['H8'].fill = fill('14B8A6'); ws1['H8'].font = font(bold=True, color=WHITE, size=8); ws1['H8'].alignment = align('center', 'center')
ws1['I8'] = "Importe ($)"; ws1['I8'].fill = fill('0F766E'); ws1['I8'].font = font(bold=True, color=WHITE, size=8); ws1['I8'].alignment = align('center', 'center')

ws1.merge_cells('J7:M7'); ws1['J7'] = "TOTAL CONSOLIDADO AMBAS OBRAS"; ws1['J7'].fill = fill(NAVY); ws1['J7'].font = font(bold=True, color=WHITE, size=9); ws1['J7'].alignment = align('center', 'center')
ws1['J8'] = "Total Diesel (L)"; ws1['J8'].fill = fill('334155'); ws1['J8'].font = font(bold=True, color=WHITE, size=8); ws1['J8'].alignment = align('center', 'center')
ws1['K8'] = "Importe Diesel ($)"; ws1['K8'].fill = fill('334155'); ws1['K8'].font = font(bold=True, color=WHITE, size=8); ws1['K8'].alignment = align('center', 'center')
ws1['L8'] = "Total Gasolina (L)"; ws1['L8'].fill = fill('334155'); ws1['L8'].font = font(bold=True, color=WHITE, size=8); ws1['L8'].alignment = align('center', 'center')
ws1['M8'] = "GRAN TOTAL ($)"; ws1['M8'].fill = fill('0F172A'); ws1['M8'].font = font(bold=True, color='FDE047', size=9); ws1['M8'].alignment = align('center', 'center')

for r in range(7, 9):
    for c in range(1, 14):
        ws1.cell(row=r, column=c).border = thin_border()
ws1.row_dimensions[7].height = 18
ws1.row_dimensions[8].height = 18

# Data rows
dyn_data = [
    ("Enero", "Arranque de Año y Mantenimientos", 0.729, 7932.9, 1395.6, 247545.67, 5845.5, 1427.9, 191953.05, 13778.4, 372017.79, 2823.5, 439498.72),
    ("Febrero", "Operación Regular Creciente", 0.882, 9597.8, 1457.8, 293982.78, 7080.6, 1491.7, 226828.43, 16678.4, 450317.98, 2949.5, 520811.21),
    ("Marzo", "Temporada Alta (Estiaje Óptimo)", 1.161, 12633.9, 1979.0, 388395.73, 9320.4, 2025.9, 299989.14, 21950.7, 592668.49, 4004.9, 688384.87),
    ("Abril", "Avance Fuerte / Paro Semana Santa", 0.958, 10424.9, 1624.5, 330297.86, 7690.7, 1663.0, 237533.45, 18120.7, 489259.76, 3287.5, 567831.31),
    ("Mayo", "Mes Pico Máximo de Asfaltado", 1.323, 14396.8, 2169.8, 440570.61, 10620.9, 2221.2, 339852.03, 25017.7, 675476.97, 4391.0, 780422.64),
    ("Junio", "Inicio de Lluvias y Taller", 0.650, 7073.2, 1346.2, 223149.98, 5218.1, 1378.2, 173697.44, 12286.5, 331734.24, 2724.4, 396847.42),
]

row_curr = 9
for idx, d in enumerate(dyn_data):
    bg = 'FFFFFF' if idx % 2 == 0 else GRAY_BG
    ws1.cell(row=row_curr, column=1, value=d[0]).font = font(bold=True, size=9)
    ws1.cell(row=row_curr, column=2, value=d[1]).font = font(size=8, italic=True)
    ws1.cell(row=row_curr, column=3, value=d[2]).number_format = '0.00x'
    
    ws1.cell(row=row_curr, column=4, value=d[3]).number_format = '#,##0.0'
    ws1.cell(row=row_curr, column=5, value=d[4]).number_format = '#,##0.0'
    ws1.cell(row=row_curr, column=6, value=d[5]).number_format = '"$"#,##0.00'
    
    ws1.cell(row=row_curr, column=7, value=d[6]).number_format = '#,##0.0'
    ws1.cell(row=row_curr, column=8, value=d[7]).number_format = '#,##0.0'
    ws1.cell(row=row_curr, column=9, value=d[8]).number_format = '"$"#,##0.00'
    
    ws1.cell(row=row_curr, column=10, value=d[9]).number_format = '#,##0.0'
    ws1.cell(row=row_curr, column=11, value=d[10]).number_format = '"$"#,##0.00'
    ws1.cell(row=row_curr, column=12, value=d[11]).number_format = '#,##0.0'
    
    c_gt = ws1.cell(row=row_curr, column=13, value=d[12])
    c_gt.number_format = '"$"#,##0.00'
    c_gt.font = font(bold=True, color='0F172A', size=9)
    c_gt.fill = fill('FEF08A')
    
    for c in range(1, 13):
        cell = ws1.cell(row=row_curr, column=c)
        cell.fill = fill(bg)
        cell.border = thin_border()
        if c in [1, 3]: cell.alignment = align('center', 'center')
        elif c == 2: cell.alignment = align('left', 'center')
        else: cell.alignment = align('right', 'center')
    c_gt.border = thin_border(); c_gt.alignment = align('right', 'center')
    ws1.row_dimensions[row_curr].height = 19
    row_curr += 1

# Totals Row
ws1.cell(row=row_curr, column=1, value="TOTAL SEMESTRE").fill = fill(NAVY)
ws1.cell(row=row_curr, column=1).font = font(bold=True, color=WHITE, size=9); ws1.cell(row=row_curr, column=1).alignment = align('left', 'center')
ws1.cell(row=row_curr, column=2, value="6 Meses Operativos").fill = fill(NAVY)
ws1.cell(row=row_curr, column=2).font = font(bold=True, color=WHITE, size=8); ws1.cell(row=row_curr, column=2).alignment = align('center', 'center')
ws1.cell(row=row_curr, column=3, value="0.95x Prom").fill = fill(NAVY)
ws1.cell(row=row_curr, column=3).font = font(bold=True, color=WHITE, size=8); ws1.cell(row=row_curr, column=3).alignment = align('center', 'center')

ws1.cell(row=row_curr, column=4, value=62059.5).number_format = '#,##0.0'
ws1.cell(row=row_curr, column=5, value=9973.1).number_format = '#,##0.0'
ws1.cell(row=row_curr, column=6, value=1923942.91).number_format = '"$"#,##0.00'

ws1.cell(row=row_curr, column=7, value=45772.9).number_format = '#,##0.0'
ws1.cell(row=row_curr, column=8, value=10207.9).number_format = '#,##0.0'
ws1.cell(row=row_curr, column=9, value=1469853.26).number_format = '"$"#,##0.00'

ws1.cell(row=row_curr, column=10, value=107832.4).number_format = '#,##0.0'
ws1.cell(row=row_curr, column=11, value=2911475.23).number_format = '"$"#,##0.00'
ws1.cell(row=row_curr, column=12, value=20180.8).number_format = '#,##0.0'

c_tot_gt = ws1.cell(row=row_curr, column=13, value=3393796.17)
c_tot_gt.number_format = '"$"#,##0.00'
c_tot_gt.font = font(bold=True, color='000000', size=10)
c_tot_gt.fill = fill('FACC15')

for c in range(1, 13):
    cell = ws1.cell(row=row_curr, column=c)
    if c > 3:
        cell.fill = fill(NAVY)
        cell.font = font(bold=True, color=WHITE, size=9)
        cell.alignment = align('right', 'center')
    cell.border = thin_border()
c_tot_gt.border = thin_border(); c_tot_gt.alignment = align('right', 'center')
ws1.row_dimensions[row_curr].height = 22

# ─────────────────────────────────────────────────────────────
# HOJA 2: MATRIZ DE FACTORES OPERATIVOS
# ─────────────────────────────────────────────────────────────
ws_fac = wb.create_sheet(title="Factores Operativos")
ws_fac.views.sheetView[0].showGridLines = True

ws_fac.merge_cells('A1:G1')
ws_fac['A1'] = "⚙️ MATRIZ DE VARIABLES Y FACTORES MENSUALES (CLIMA, MAQUINARIA Y CALENDARIO)"
ws_fac['A1'].fill = fill(NAVY); ws_fac['A1'].font = font(bold=True, color=WHITE, size=12); ws_fac['A1'].alignment = align('left', 'center')
ws_fac.row_dimensions[1].height = 26

headers_mf = ["MES", "CLIMA / LLUVIA (F1)", "MAQUINARIA / TALLER (F2)", "DÍAS HÁBILES (F3)", "FACTOR COMPUESTO DIESEL", "FACTOR GASOLINA", "JUSTIFICACIÓN TÉCNICA Y OPERATIVA"]
for c_i, h in enumerate(headers_mf, 1):
    c = ws_fac.cell(row=3, column=c_i, value=h)
    c.fill = fill('334155'); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center'); c.border = thin_border()
ws_fac.row_dimensions[3].height = 20

factores_expl = [
    ("Enero", "0.90 (Frío matutino extremo en Toluca/La Marquesa)", "0.88 (Reactivación post-vacaciones y afinación)", "0.92 (21 d. hábiles, arranque paulatino)", "0.729 (-27.1%)", "0.830", "El frío intenso matutino retrasa el tendido de asfalto caliente hasta que sube la temperatura. Mantenimiento preventivo anual."),
    ("Febrero", "0.98 (Seco, buenas horas de sol)", "1.00 (Equipos 100% disponibles)", "0.90 (Mes corto de 28 días / 20 d. hábiles)", "0.882 (-11.8%)", "0.874", "Condiciones favorables de trabajo. El consumo total baja respecto a meses de 31 días únicamente por ser mes de 28 días."),
    ("Marzo", "1.15 (Temporada seca óptima / Estiaje)", "0.98 (Operación continua sin paros mayores)", "1.03 (22 d. hábiles completos)", "1.161 (+16.1%)", "1.100", "Inicio del periodo óptimo de asfalto. Sin lluvias, frentes de fresado y pavimentación trabajan horas completas y turnos extra."),
    ("Abril", "1.12 (Calor y clima seco)", "0.92 (Cambio de picas y mantenimiento plancha)", "0.93 (Paro de Semana Santa: 19 d. hábiles)", "0.958 (-4.2%)", "0.950", "Mes de alta productividad interrumpido 3-4 días por Semana Santa, aprovechados para mantenimiento de fresadora y pavimentadora."),
    ("Mayo", "1.20 (Máxima ventana seca previa a lluvias)", "1.05 (Dobles turnos, frentes acelerados)", "1.05 (23 d. hábiles intensos)", "1.323 (+32.3%)", "1.240", "MES PICO DEL SEMESTRE. Se acelera el avance físico de obra para cerrar tramos antes de la temporada de lluvias. Máximo consumo de Diesel."),
    ("Junio", "0.78 (Inicio de lluvias fuertes en montaña)", "0.85 (Equipos entran a taller por paros)", "0.98 (21 d. hábiles afectados por tormentas)", "0.650 (-35.0%)", "0.800", "Las lluvias torrenciales impiden el riego de liga y tendido de mezcla asfáltica. Se generan paros forzados y traslados a taller.")
]

r_f = 4
for idx, f in enumerate(factores_expl):
    bg = 'FFFFFF' if idx % 2 == 0 else GRAY_BG
    for c_i, val in enumerate(f, 1):
        cell = ws_fac.cell(row=r_f, column=c_i, value=val)
        cell.fill = fill(bg); cell.border = thin_border()
        if c_i == 1: cell.font = font(bold=True, size=9); cell.alignment = align('center', 'center')
        elif c_i in [5, 6]: cell.font = font(bold=True, color='1E40AF', size=9); cell.alignment = align('center', 'center')
        elif c_i == 7: cell.font = font(size=8); cell.alignment = align('left', 'center')
        else: cell.font = font(size=8); cell.alignment = align('center', 'center')
    ws_fac.row_dimensions[r_f].height = 24
    r_f += 1

# ─────────────────────────────────────────────────────────────
# HOJA 3: PROMEDIOS DE REFERENCIA
# ─────────────────────────────────────────────────────────────
ws_ref = wb.create_sheet(title="Promedios de Referencia")
ws_ref.views.sheetView[0].showGridLines = True

ws_ref.merge_cells('A1:G1')
ws_ref['A1'] = "📌 PROMEDIOS REALES DE REFERENCIA MENSUAL Y SEMANAL"
ws_ref['A1'].fill = fill(NAVY); ws_ref['A1'].font = font(bold=True, color=WHITE, size=12); ws_ref['A1'].alignment = align('left', 'center')
ws_ref.row_dimensions[1].height = 26

headers_t1 = ["OBRA", "COMBUSTIBLE", "PROMEDIO SEMANAL (L)", "PROMEDIO MENSUAL BASE (L)", "PRECIO PROM. ($/L)", "IMPORTE MENSUAL ($)", "PARTICIPACIÓN %"]
for c_i, h in enumerate(headers_t1, 1):
    c = ws_ref.cell(row=3, column=c_i, value=h)
    c.fill = fill('334155'); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center'); c.border = thin_border()
ws_ref.row_dimensions[3].height = 20

data_t1 = [
    ("MÉXICO - TOLUCA", "DIESEL", 2511.2, 10881.9, 27.00, 293810.40, "49.6%"),
    ("MÉXICO - TOLUCA", "GASOLINA", 392.8, 1701.9, 23.90, 40676.00, "6.9%"),
    ("MÉXICO - TOLUCA (SUBTOTAL)", "AMBOS", 2904.0, 12583.8, 26.58, 334486.40, "56.4%"),
    
    ("LERMA - TRES MARÍAS", "DIESEL", 1852.6, 8027.9, 27.00, 216753.00, "36.6%"),
    ("LERMA - TRES MARÍAS", "GASOLINA", 401.1, 1738.0, 23.90, 41538.50, "7.0%"),
    ("LERMA - TRES MARÍAS (SUBTOTAL)", "AMBOS", 2253.7, 9765.9, 26.45, 258291.50, "43.6%"),
]

r_curr = 4
for idx, d in enumerate(data_t1):
    is_sub = "SUBTOTAL" in d[0]
    bg = 'E2E8F0' if is_sub else ('FFFFFF' if idx % 2 == 0 else 'F8FAFC')
    
    ws_ref.cell(row=r_curr, column=1, value=d[0]).font = font(bold=is_sub, size=9)
    ws_ref.cell(row=r_curr, column=2, value=d[1]).font = font(bold=is_sub, color='1E40AF' if d[1]=='DIESEL' else ('047857' if d[1]=='GASOLINA' else '000000'), size=9)
    ws_ref.cell(row=r_curr, column=3, value=d[2]).number_format = '#,##0.0'
    ws_ref.cell(row=r_curr, column=4, value=d[3]).number_format = '#,##0.0'
    ws_ref.cell(row=r_curr, column=5, value=d[4]).number_format = '"$"#,##0.00'
    ws_ref.cell(row=r_curr, column=6, value=d[5]).number_format = '"$"#,##0.00'
    ws_ref.cell(row=r_curr, column=7, value=d[6])
    
    for c in range(1, 8):
        cell = ws_ref.cell(row=r_curr, column=c)
        cell.fill = fill(bg)
        cell.border = thin_border()
        if c in [1, 2]: cell.alignment = align('left' if c==1 else 'center', 'center')
        elif c == 7: cell.alignment = align('center', 'center')
        else: cell.alignment = align('right', 'center')
    ws_ref.row_dimensions[r_curr].height = 18
    r_curr += 1

# Total General
ws_ref.cell(row=r_curr, column=1, value="GRAN TOTAL CONSOLIDADO MENSUAL").fill = fill(NAVY)
ws_ref.cell(row=r_curr, column=1).font = font(bold=True, color=WHITE, size=9); ws_ref.cell(row=r_curr, column=1).alignment = align('left', 'center')
ws_ref.cell(row=r_curr, column=2, value="AMBAS OBRAS").fill = fill(NAVY)
ws_ref.cell(row=r_curr, column=2).font = font(bold=True, color=WHITE, size=9); ws_ref.cell(row=r_curr, column=2).alignment = align('center', 'center')

ws_ref.cell(row=r_curr, column=3, value=5157.7).number_format = '#,##0.0'
ws_ref.cell(row=r_curr, column=4, value=22349.7).number_format = '#,##0.0'
ws_ref.cell(row=r_curr, column=5, value=26.52).number_format = '"$"#,##0.00'

c_tot = ws_ref.cell(row=r_curr, column=6, value=592777.90)
c_tot.number_format = '"$"#,##0.00'
c_tot.font = font(bold=True, color='000000', size=10)
c_tot.fill = fill('FDE047')

ws_ref.cell(row=r_curr, column=7, value="100.0%").fill = fill(NAVY)
ws_ref.cell(row=r_curr, column=7).font = font(bold=True, color=WHITE, size=9); ws_ref.cell(row=r_curr, column=7).alignment = align('center', 'center')

for c in range(1, 8):
    cell = ws_ref.cell(row=r_curr, column=c)
    if c != 6:
        cell.fill = fill(NAVY); cell.font = font(bold=True, color=WHITE, size=9)
    if c in [3, 4, 5]: cell.alignment = align('right', 'center')
    cell.border = thin_border()
c_tot.border = thin_border(); c_tot.alignment = align('right', 'center')
ws_ref.row_dimensions[r_curr].height = 22

# Auto adjust columns
for ws in [ws1, ws_fac, ws_ref]:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 13)

filename = 'c:/Users/JOSE/Desktop/Proyecto fenix/Estimacion_Combustible_MexicoToluca_Lerma_Ene_Jun.xlsx'
wb.save(filename)
print("Saved complete dynamic workbook to:", filename)
