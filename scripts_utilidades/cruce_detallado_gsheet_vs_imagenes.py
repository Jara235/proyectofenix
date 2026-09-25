import openpyxl, re, psycopg2
from psycopg2.extras import DictCursor

# Facturas pagadas en imágenes
tabla_1_raw = [
    ("20 de julio de 2026", 1857.15, "BACHEO TOLUCA", "A-10484"),
    ("20 de julio de 2026", 12150.08, "P. DE ASFALTO HUIXQUILUCAN", "A-10485"),
    ("20 de julio de 2026", 21600.09, "P. DE ASFALTO HUIXQUILUCAN", "A-10491"),
    ("20 de julio de 2026", 10800.97, "ALFREDO DEL MAZO", "A-10492"),
    ("14 de julio de 2026", 27000.11, "COLEGIO MILITAR", "A-10496"),
    ("21 de julio de 2026", 10800.13, "P. DE ASFALTO HUIXQUILUCAN", "A-10520"),
    ("21 de julio de 2026", 21600.15, "DEZAZOLVE CANALES", "A-10521"),
    ("21 de julio de 2026", 25925.00, "MEXICO TOLUCA", "A-10522"),
    ("21 de julio de 2026", 5400.05, "ALFREDO DEL MAZO", "A-10532"),
    ("22 de julio de 2026", 21600.15, "P. DE ASFALTO HUIXQUILUCAN", "A-10535"),
    ("22 de julio de 2026", 21600.15, "DEZAZOLVE CANALES", "A-10543"),
    ("23 de julio de 2026", 8100.00, "ALFREDO DEL MAZO", "A-10544"),
    ("23 de julio de 2026", 10800.10, "P. DE ASFALTO HUIXQUILUCAN", "A-10556"),
    ("23 de julio de 2026", 25650.13, "MEXICO TOLUCA", "A-10563"),
    ("23 de julio de 2026", 4860.05, "DEZAZOLVE CANALES", "A-10564"),
    ("23 de julio de 2026", 7290.05, "CONSTITUCION", "A-10565"),
    ("23 de julio de 2026", 14850.08, "ALFREDO DEL MAZO", "A-10572"),
    ("23 de julio de 2026", 24840.10, "ALFREDO DEL MAZO", "A-10577"),
    ("24 de julio de 2026", 10800.10, "P. DE ASFALTO HUIXQUILUCAN", "A-10579"),
    ("27 de julio de 2026", 10800.10, "P. DE ASFALTO HUIXQUILUCAN", "A-10581"),
    ("24 de julio de 2026", 18900.06, "ALFREDO DEL MAZO", "A-10583"),
    ("24 de julio de 2026", 14364.05, "DEZAZOLVE CANALES", "A-10586"),
    ("24 de julio de 2026", 1080.03, "MEXICO TOLUCA", "A-10587"),
    ("24 de julio de 2026", 12960.06, "ALFREDO DEL MAZO", "A-10588"),
    ("25 de julio de 2026", 24840.10, "MEXICO TOLUCA", "A-10589"),
    ("27 de julio de 2026", 9720.03, "LERMA", "A-10603"),
    ("27 de julio de 2026", 3240.00, "PROVIDENCIA", "A-10604"),
    ("28 de julio de 2026", 10800.03, "P. DE ASFALTO HUIXQUILUCAN", "A-10605"),
    ("28 de julio de 2026", 30502.11, "MEXICO TOLUCA", "A-10623"),
    ("28 de julio de 2026", 23490.11, "DEZAZOLVE CANALES", "A-10624"),
    ("28 de julio de 2026", 10800.10, "P. DE ASFALTO HUIXQUILUCAN", "A-10625"),
    ("28 de julio de 2026", 26959.90, "P. DE ASFALTO HUIXQUILUCAN", "A-10630"),
    ("28 de julio de 2026", 1620.03, "P. DE ASFALTO HUIXQUILUCAN", "A-10634"),
    ("29 de julio de 2026", 1558.03, "P. DE ASFALTO HUIXQUILUCAN", "A-10635"),
    ("28 de julio de 2026", 10260.10, "ALFREDO DEL MAZO", "A-10636"),
    ("29 de julio de 2026", 24564.14, "LERMA", "A-10637"),
    ("31 de julio de 2026", 3510.00, "PROVIDENCIA", "A-10644"),
    ("31 de julio de 2026", 17820.03, "ALFREDO DEL MAZO", "A-10645"),
    ("31 de julio de 2026", 25650.09, "DEZAZOLVE CANALES", "A-10646"),
    ("30 de julio de 2026", 21600.46, "P. DE ASFALTO HUIXQUILUCAN", "A-10652"),
    ("27 de julio de 2026", 3105.00, "BACHEO TOLUCA", "A-10653"),
    ("27 de julio de 2026", 24852.31, "DEZAZOLVE CANALES", "A-10654"),
    ("1 de agosto de 2026", 8100.05, "P. DE ASFALTO HUIXQUILUCAN", "A-10669"),
    ("1 de agosto de 2026", 25920.06, "MEXICO TOLUCA", "A-10673"),
    ("1 de agosto de 2026", 8100.05, "P. DE ASFALTO HUIXQUILUCAN", "A-10674"),
    ("29 de julio de 2026", 20520.09, "DEZAZOLVE CANALES", "A-10702"),
    ("29 de julio de 2026", 1620.03, "PROVIDENCIA", "A-10712"),
    ("29 de julio de 2026", 9736.66, "ALFREDO DEL MAZO", "A-10713"),
    ("30 de julio de 2026", 21600.14, "P. DE ASFALTO HUIXQUILUCAN", "A-10722"),
    ("30 de julio de 2026", 24840.10, "MEXICO TOLUCA", "A-10733"),
    ("30 de julio de 2026", 1485.03, "P. DE ASFALTO HUIXQUILUCAN", "A-10748"),
    ("17 de julio de 2026", 5400.05, "P. DE ASFALTO HUIXQUILUCAN", "A-23783100"),
]

tabla_2_raw = [
    ("6 de julio de 2026", 66605.40, "TANQUE PEGASO", "A-10360"),
    ("7 de julio de 2026", 8100.05, "TANQUE PEGASO", "A-10380"),
    ("7 de julio de 2026", 3000.00, "TANQUE PEGASO", "A-10386"),
    ("7 de julio de 2026", 18900.06, "LERMA", "A-10381"),
    ("7 de julio de 2026", 2700.08, "LERMA", "A-10388"),
    ("6 de julio de 2026", 8100.05, "P. DE ASFALTO HUIXQUILUCAN", "A-10361"),
    ("6 de julio de 2026", 10800.10, "P. DE ASFALTO HUIXQUILUCAN", "A-10365"),
    ("6 de julio de 2026", 16200.10, "P. DE ASFALTO HUIXQUILUCAN", "A-10368"),
    ("7 de julio de 2026", 16200.10, "P. DE ASFALTO HUIXQUILUCAN", "A-10884"),
    ("8 de julio de 2026", 10800.10, "P. DE ASFALTO HUIXQUILUCAN", "A-10405"),
    ("8 de julio de 2026", 8100.05, "P. DE ASFALTO HUIXQUILUCAN", "A-10406"),
    ("8 de julio de 2026", 8100.05, "P. DE ASFALTO HUIXQUILUCAN", "A-10408"),
    ("6 de julio de 2026", 6291.00, "PROVIDENCIA", "A-10362"),
    ("6 de julio de 2026", 7695.00, "BACHEO TOLUCA", "A-10366"),
    ("7 de julio de 2026", 21600.15, "BACHEO TOLUCA", "A-10387"),
    ("6 de julio de 2026", 19305.10, "MEXICO TOLUCA", "A-10367"),
    ("8 de julio de 2026", 16200.10, "MEXICO TOLUCA", "A-10398"),
    ("8 de julio de 2026", 23206.21, "DEZAZOLVE CANALES", "A-10409"),
    ("13 de julio de 2026", 601.43, "TANQUE PEGASO", "A-10420"),
    ("15 de julio de 2026", 108000.36, "TANQUE PEGASO", "A-10510"),
    ("13 de julio de 2026", 4197.68, "TANQUE PEGASO", "A-10511"),
    ("10 de julio de 2026", 18043.21, "LERMA", "A-10894"),
    ("13 de julio de 2026", 17804.45, "LERMA", "A-10426"),
    ("13 de julio de 2026", 3000.00, "LERMA", "A-10427"),
    ("9 de julio de 2026", 2700.08, "MAQUINARIA", "A-10886"),
    ("9 de julio de 2026", 16201.80, "P. DE ASFALTO HUIXQUILUCAN", "A-10888"),
    ("10 de julio de 2026", 27000.11, "P. DE ASFALTO HUIXQUILUCAN", "A-10899"),
    ("11 de julio de 2026", 13500.06, "P. DE ASFALTO HUIXQUILUCAN", "A-10904"),
    ("12 de julio de 2026", 8100.05, "P. DE ASFALTO HUIXQUILUCAN", "A-10912"),
    ("14 de julio de 2026", 16200.10, "P. DE ASFALTO HUIXQUILUCAN", "A-10495"),
    ("15 de julio de 2026", 8100.05, "P. DE ASFALTO HUIXQUILUCAN", "A-10438"),
    ("15 de julio de 2026", 27000.11, "P. DE ASFALTO HUIXQUILUCAN", "A-10514"),
    ("13 de julio de 2026", 14445.05, "VICENTE LOMBARDO", "A-10422"),
    ("14 de julio de 2026", 11880.00, "VICENTE LOMBARDO", "A-10499"),
    ("9 de julio de 2026", 10800.10, "PROVIDENCIA", "A-10887"),
    ("15 de julio de 2026", 5400.05, "PROVIDENCIA", "A-10513"),
    ("10 de julio de 2026", 5400.05, "ALFREDO DEL MAZO", "A-10893"),
    ("13 de julio de 2026", 22652.97, "ALFREDO DEL MAZO", "A-10425"),
    ("15 de julio de 2026", 10800.13, "ALFREDO DEL MAZO", "A-10512"),
    ("15 de julio de 2026", 29693.92, "MEXICO TOLUCA", "A-10509"),
    ("10 de julio de 2026", 24705.06, "DEZAZOLVE CANALES", "A-10900"),
    ("13 de julio de 2026", 14445.05, "DEZAZOLVE CANALES", "A-10423"),
    ("14 de julio de 2026", 11880.00, "DEZAZOLVE CANALES", "A-10500"),
    ("16 de julio de 2026", 16200.10, "ALFREDO DEL MAZO", "A-10446"),
    ("16 de julio de 2026", 10800.10, "PROVIDENCIA", "A-10447"),
    ("16 de julio de 2026", 13500.55, "P. DE ASFALTO HUIXQUILUCAN", "A-10449"),
    ("16 de julio de 2026", 22950.13, "DEZAZOLVE CANALES", "A-10453"),
    ("16 de julio de 2026", 1800.02, "TANQUE PEGASO", "A-10454"),
    ("17 de julio de 2026", 27000.11, "LERMA", "A-10457"),
    ("17 de julio de 2026", 10800.06, "P. DE ASFALTO HUIXQUILUCAN", "A-10458"),
    ("17 de julio de 2026", 32400.12, "MEXICO TOLUCA", "A-10459"),
    ("18 de julio de 2026", 20250.13, "ALFREDO DEL MAZO", "A-10474"),
    ("20 de julio de 2026", 12285.08, "PROVIDENCIA", "A-10479"),
    ("20 de julio de 2026", 14715.08, "LERMA", "A-10480"),
    ("20 de julio de 2026", 16200.10, "DEZAZOLVE CANALES", "A-10483"),
    ("20 de julio de 2026", 11642.85, "BACHEO TOLUCA", "A-10484"),
]

def clean_fol(f):
    m = re.findall(r'[0-9]+', str(f))
    return "".join(m) if m else str(f)

# Folios pagados en las imágenes
pagados_por_folio = {}
for f in tabla_1_raw + tabla_2_raw:
    cf = clean_fol(f[3])
    if cf not in pagados_por_folio:
        pagados_por_folio[cf] = {'orig': f[3], 'importe_pagado': 0.0, 'items': []}
    pagados_por_folio[cf]['importe_pagado'] += f[1]
    pagados_por_folio[cf]['items'].append(f)

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)
ws = wb['Estado de Cuenta']

semanas_info = {}
curr_sem = None

for r in range(31, 65):
    c2 = str(ws.cell(r, 2).value or '').strip()
    c4 = str(ws.cell(r, 4).value or '').strip()
    c6 = str(ws.cell(r, 6).value or '').strip()
    c8 = ws.cell(r, 8).value
    
    m_sem = re.search(r'Semana\s+(\d+)', c2, re.IGNORECASE)
    if m_sem:
        curr_sem = int(m_sem.group(1))
        semanas_info[curr_sem] = {'rows': []}
        continue
    
    if curr_sem and c4:
        semanas_info[curr_sem]['rows'].append({
            'tipo': c4,
            'facturas_str': c6,
            'monto': float(c8 or 0.0) if str(c8).replace('.','',1).isdigit() else 0.0
        })

print("================================================================================")
print("CRUCE DETALLADO POR SEMANA: ESTADO DE CUENTA (GSHEET) vs PAGOS REALIZADOS (IMÁGENES)")
print("================================================================================")

for sem, data in semanas_info.items():
    print(f"\n==================== SEMANA {sem} ====================")
    folios_sem_gs = set()
    for row in data['rows']:
        # extraer folios de la celda
        raw = row['facturas_str']
        # Buscar todos los folios tipo A10372, A-10372, 10372, 23783097
        fols = re.findall(r'A\s*([0-9]+)|([0-9]{4,8})', raw, re.IGNORECASE)
        found = []
        for a, b in fols:
            val = a if a else b
            if val:
                found.append(val)
                folios_sem_gs.add(val)
        print(f"  [{row['tipo']:<18}] Monto: ${row['monto']:>10,.2f} | Facturas ({len(found)}): {found}")
    
    # Ver qué facturas de este set se pagaron y cuáles no
    pagadas_en_sem = []
    pendientes_en_sem = []
    
    for f in sorted(folios_sem_gs):
        if f in pagados_por_folio:
            pagadas_en_sem.append((f, pagados_por_folio[f]['importe_pagado']))
        else:
            pendientes_en_sem.append(f)
            
    print(f"\n  >> Resumen Conciliación Semana {sem}:")
    print(f"     * Facturas del Edo Cta YA PAGADAS en las imágenes: {len(pagadas_en_sem)} facturas (Total: ${sum(x[1] for x in pagadas_en_sem):,.2f})")
    print(f"     * Facturas del Edo Cta PENDIENTES DE PAGO / NO EN LAS IMÁGENES: {len(pendientes_en_sem)} facturas -> {pendientes_en_sem}")

    # Ver si se pagaron facturas en esta semana que NO estaban en el Edo Cta de esta semana
    # (por ejemplo Tanque Pegaso u otras)
    pagadas_extras = []
    for cf, pdata in pagados_por_folio.items():
        # check if it belongs to this week by date
        # (check if first item date falls in this week)
        for item in pdata['items']:
            # we can check
            pass
