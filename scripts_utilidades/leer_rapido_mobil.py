import openpyxl

wb = openpyxl.load_workbook('formatos/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_FINAL_V2.xlsx', read_only=True, data_only=True)

for sheet in ['Estados_Cuenta_Mobil', 'Captura_Facturas_Mobil']:
    if sheet in wb.sheetnames:
        ws = wb[sheet]
        print(f"=== Sheet: {sheet} ===")
        count = 0
        for row in ws.iter_rows(values_only=True):
            vals = [str(v).encode('ascii', 'ignore').decode('ascii') for v in row if v is not None]
            if vals:
                count += 1
                if count <= 25:
                    print(f"  R{count:02d}: " + " | ".join(vals[:8]))
        print(f"Total rows in {sheet}: {count}")
