import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

maestro_path = r'c:\Users\JOSE\Desktop\Proyecto fenix\MAESTRO_CONTROL_GASOLINA_NUEVO.xlsx'

# 1. Load Data
df_auth = pd.read_excel(maestro_path, sheet_name='BD_AUTORIZACIONES')
df_gas = pd.read_excel(maestro_path, sheet_name='BD_GASOLINA')

# Normalize columns for grouping
df_auth['VEHICULO'] = df_auth['VEHICULO'].fillna('NO REPORTADO').astype(str).str.strip().str.upper()
df_auth['PLACA'] = df_auth['PLACA'].fillna('S/P').astype(str).str.strip().str.upper()
df_auth['IMPORTE_AUTORIZADO'] = pd.to_numeric(df_auth['IMPORTE_AUTORIZADO'], errors='coerce').fillna(0)

df_gas['VEHICULO'] = df_gas['VEHICULO'].fillna('NO REPORTADO').astype(str).str.strip().str.upper()
df_gas['PLACA'] = df_gas['PLACA'].fillna('S/P').astype(str).str.strip().str.upper()
df_gas['IMPORTE_TOTAL'] = pd.to_numeric(df_gas['IMPORTE_TOTAL'], errors='coerce').fillna(0)

# Group Auth
grp_auth = df_auth.groupby(['SEMANA', 'VEHICULO', 'PLACA']).agg({
    'IMPORTE_AUTORIZADO': 'sum',
    'RESPONSABLE': 'first'
}).reset_index()

# Group Gas
grp_gas = df_gas.groupby(['SEMANA', 'VEHICULO', 'PLACA']).agg({
    'IMPORTE_TOTAL': 'sum',
    'ORIGEN': 'first' # Just to keep some context
}).reset_index()
grp_gas.rename(columns={'IMPORTE_TOTAL': 'IMPORTE_REPORTADO'}, inplace=True)

# Merge
df_merge = pd.merge(grp_auth, grp_gas, on=['SEMANA', 'VEHICULO', 'PLACA'], how='outer')

# Fill NaNs
df_merge['IMPORTE_AUTORIZADO'] = df_merge['IMPORTE_AUTORIZADO'].fillna(0)
df_merge['IMPORTE_REPORTADO'] = df_merge['IMPORTE_REPORTADO'].fillna(0)
df_merge['RESPONSABLE'] = df_merge['RESPONSABLE'].fillna('Sin Registro en Autorizados')

# Calculate Difference
df_merge['DIFERENCIA (Saldo a favor)'] = df_merge['IMPORTE_AUTORIZADO'] - df_merge['IMPORTE_REPORTADO']

# Status
def get_status(row):
    diff = row['DIFERENCIA (Saldo a favor)']
    auth = row['IMPORTE_AUTORIZADO']
    if auth == 0 and row['IMPORTE_REPORTADO'] > 0:
        return 'ALERTA: Sin Presupuesto'
    elif diff < 0:
        return 'ALERTA: Sobrepasó Presupuesto'
    elif diff == 0:
        return 'Exacto al Presupuesto'
    else:
        return 'Ahorro / Dentro de Presupuesto'

df_merge['ESTATUS'] = df_merge.apply(get_status, axis=1)

# Sort
df_merge = df_merge.sort_values(by=['SEMANA', 'DIFERENCIA (Saldo a favor)'], ascending=[False, True])

# Formatting and Writing
wb = load_workbook(maestro_path)
sheet_name = 'Control_Presupuestal'
if sheet_name in wb.sheetnames:
    del wb[sheet_name]
ws = wb.create_sheet(sheet_name)

# Headers
headers = list(df_merge.columns)
ws.append(headers)

# Header style
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1A3C5E", end_color="1A3C5E", fill_type="solid")
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center')

# Data
for r in dataframe_to_rows(df_merge, index=False, header=False):
    ws.append(r)

# Format Currency & Alerts
green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
green_font = Font(color="006100")
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
red_font = Font(color="9C0006")
yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
yellow_font = Font(color="9C6500")

for row in ws.iter_rows(min_row=2, max_col=len(headers)):
    # Currency formats for Importes and Diferencia
    for col_idx in [headers.index('IMPORTE_AUTORIZADO')+1, headers.index('IMPORTE_REPORTADO')+1, headers.index('DIFERENCIA (Saldo a favor)')+1]:
        row[col_idx-1].number_format = '"$"#,##0.00'
        
    # Status Alert formatting
    status_cell = row[headers.index('ESTATUS')]
    if 'Sin Presupuesto' in status_cell.value:
        status_cell.fill = yellow_fill
        status_cell.font = yellow_font
    elif 'Sobrepasó' in status_cell.value:
        status_cell.fill = red_fill
        status_cell.font = red_font
    elif 'Ahorro' in status_cell.value or 'Exacto' in status_cell.value:
        status_cell.fill = green_fill
        status_cell.font = green_font

# Adjust widths
for col in ws.columns:
    max_length = 0
    col_letter = col[0].column_letter
    for cell in col:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(cell.value)
        except: pass
    ws.column_dimensions[col_letter].width = max_length + 2

# Totals row
max_row = ws.max_row
ws.cell(row=max_row+1, column=1, value="TOTALES").font = Font(bold=True)
for col_name in ['IMPORTE_AUTORIZADO', 'IMPORTE_REPORTADO', 'DIFERENCIA (Saldo a favor)']:
    col_idx = headers.index(col_name) + 1
    col_letter = ws.cell(row=1, column=col_idx).column_letter
    cell = ws.cell(row=max_row+1, column=col_idx)
    cell.value = f"=SUM({col_letter}2:{col_letter}{max_row})"
    cell.number_format = '"$"#,##0.00'
    cell.font = Font(bold=True)

wb.save(maestro_path)
print(f"Sheet '{sheet_name}' created successfully with {len(df_merge)} summary rows.")
