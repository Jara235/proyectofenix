import openpyxl

wb = openpyxl.load_workbook('formatos/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_FINAL_V2.xlsx', data_only=True)

for sheet in ['Estados_Cuenta_Mobil', 'Captura_Facturas_Mobil', 'Dashboard_Ejecutivo', 'Resumen_Ejecutivo']:
    if sheet in wb.sheetnames:
        ws = wb[sheet]
        print(f"=== Sheet: {sheet} (rows={ws.max_row}, cols={ws.max_column}) ===")
        for r in range(1, min(ws.max_row+1, 50)):
            vals = [str(ws.cell(r, c).value) for c in range(1, min(ws.max_column+1, 15)) if ws.cell(r, c).value is not None]
            if vals:
                print(f"  R{r:02d}: " + " | ".join(vals[:8]))
