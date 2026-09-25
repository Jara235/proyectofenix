import openpyxl

wb = openpyxl.load_workbook('formatos/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_FINAL_V2.xlsx', data_only=True)
ws = wb['Maquinaria_Inventario']

current_year = 2026
degradation_rate = 0.012 # 1.2% anual

equipos = []
for r in range(2, ws.max_row + 1):
    num = ws.cell(r, 1).value
    unidad = ws.cell(r, 2).value
    empresa = ws.cell(r, 3).value
    marca = ws.cell(r, 4).value
    modelo = ws.cell(r, 5).value
    anio = ws.cell(r, 6).value
    serie = ws.cell(r, 7).value
    rend_nuevo = ws.cell(r, 8).value
    um = ws.cell(r, 9).value
    eco = ws.cell(r, 12).value
    
    if unidad and rend_nuevo:
        try:
            r_base = float(rend_nuevo)
            a = int(anio) if anio and str(anio).isdigit() else 2015 # default if missing
            edad = max(0, current_year - a)
            
            if um == 'km/L':
                # For km/L, older vehicle gives less km per liter
                r_ajustado = r_base / (1.0 + (edad * degradation_rate))
                rango_min = r_ajustado * 0.85
                rango_max = r_ajustado * 1.15
            else:
                # For L/h, older machine consumes more liters per hour
                r_ajustado = r_base * (1.0 + (edad * degradation_rate))
                rango_min = r_ajustado * 0.85
                rango_max = r_ajustado * 1.20
                
            equipos.append({
                'eco': str(eco or f'EQ-{r-1}').strip(),
                'unidad': str(unidad).strip(),
                'marca': str(marca or '').strip(),
                'modelo': str(modelo or '').strip(),
                'anio': a if anio and str(anio).isdigit() else 'S/A',
                'edad': edad,
                'rend_nuevo': r_base,
                'um': str(um or 'L/h').strip(),
                'desgaste_pct': edad * degradation_rate * 100,
                'rend_ajustado': r_ajustado,
                'rango_min': rango_min,
                'rango_max': rango_max
            })
        except Exception as e:
            pass

print(f"Total equipos procesados: {len(equipos)}")
print("-" * 135)
print(f"{'Eco':<7} | {'Tipo de Máquina':<28} | {'Marca / Modelo':<26} | {'Año':<5} | {'Edad':<5} | {'Rend. Base':<10} | {'Desgaste':<8} | {'Rend. Esperado 2026':<20} | {'Rango Normal':<18}")
print("-" * 135)

for eq in equipos:
    marca_mod = f"{eq['marca']} {eq['modelo']}"[:25]
    print(f"{eq['eco']:<7} | {eq['unidad']:<28} | {marca_mod:<26} | {str(eq['anio']):<5} | {eq['edad']:<3} a | {eq['rend_nuevo']:>6.1f} {eq['um']:<3} | {eq['desgaste_pct']:>5.1f}%  | {eq['rend_ajustado']:>8.2f} {eq['um']:<4} | {eq['rango_min']:>5.1f} - {eq['rango_max']:>5.1f} {eq['um']}")

print("-" * 135)
