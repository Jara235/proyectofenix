import openpyxl

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx', data_only=True)
print("Sheet names in downloaded workbook:")
for idx, name in enumerate(wb.sheetnames):
    ws = wb[name]
    print(f"[{idx}] '{name}' -> rows: {ws.max_row}, cols: {ws.max_column}")
    # print first 5 rows
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [str(ws.cell(r, c).value or '') for c in range(1, min(12, ws.max_column + 1))]
        if any(vals):
            print(f"   R{r}:", [v for v in vals if v][:8])
