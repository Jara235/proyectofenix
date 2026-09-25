import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

maestro_path = r'c:\Users\JOSE\Desktop\Proyecto fenix\MAESTRO_CONTROL_GASOLINA_NUEVO.xlsx'

# ---- Build the normalized dataset ----
def load_gt(path, semana):
    df = pd.read_excel(path, sheet_name='Consumos Semanales GT', header=2)
    df.columns = ['_idx','RESPONSABLE','VEHICULO','PLACA','OBRA_DESTINO','IMPORTE_AUTORIZADO','GASOLINERA']
    df = df[df['RESPONSABLE'].notna() & (df['RESPONSABLE'] != 'RESPONSABLE') & (df['RESPONSABLE'] != 'Total General')].copy()
    df['SEMANA'] = semana
    return df.drop(columns=['_idx'])

gt26 = load_gt(r'c:\Users\JOSE\Desktop\Proyecto fenix\CONSUMOS_DE_GASOLINA_SEMANALES_GT.xlsx', 26)
gt27 = load_gt(r'c:\Users\JOSE\Desktop\Proyecto fenix\CONSUMOS_DE_GASOLINA_SEMANALES_GT_LEVET_HUIX.xlsx', 27)
df_estado = pd.concat([gt26, gt27], ignore_index=True)

# Normalize
for col in ['RESPONSABLE','VEHICULO','OBRA_DESTINO','GASOLINERA']:
    df_estado[col] = df_estado[col].astype(str).str.strip().str.upper()
df_estado['PLACA'] = df_estado['PLACA'].astype(str).str.strip().str.upper().str.replace(r'[\s\-]','',regex=True)
df_estado.loc[df_estado['PLACA'] == 'NAN', 'PLACA'] = 'S/P'
df_estado.loc[df_estado['VEHICULO'] == 'NAN', 'VEHICULO'] = 'EQUIPO MENOR'
df_estado['IMPORTE_AUTORIZADO'] = pd.to_numeric(df_estado['IMPORTE_AUTORIZADO'], errors='coerce').fillna(0)

# Read BD_GASOLINA consumos
bd_gas = pd.read_excel(maestro_path, sheet_name='BD_GASOLINA')
bd_gas['PLACA'] = bd_gas['PLACA'].astype(str).str.strip().str.upper().str.replace(r'[\s\-]','',regex=True)
bd_gas['SEMANA'] = pd.to_numeric(bd_gas['SEMANA'], errors='coerce').fillna(0).astype(int)
bd_gas['IMPORTE_TOTAL'] = pd.to_numeric(bd_gas['IMPORTE_TOTAL'], errors='coerce').fillna(0)
bd_gas['LITROS'] = pd.to_numeric(bd_gas['LITROS'], errors='coerce').fillna(0)

agg = bd_gas.groupby(['PLACA','SEMANA']).agg(
    CONSUMO_REAL=('IMPORTE_TOTAL','sum'),
    LITROS_REALES=('LITROS','sum')
).reset_index()

merged = pd.merge(df_estado, agg, on=['PLACA','SEMANA'], how='left')
merged['CONSUMO_REAL'] = merged['CONSUMO_REAL'].fillna(0)
merged['LITROS_REALES'] = merged['LITROS_REALES'].fillna(0)
merged['DIFERENCIA'] = merged['IMPORTE_AUTORIZADO'] - merged['CONSUMO_REAL']
merged['ESTATUS'] = merged['DIFERENCIA'].apply(lambda x: 'OK - Saldo a favor' if x >= 0 else 'ALERTA - Excedido')

final_cols = ['SEMANA','RESPONSABLE','VEHICULO','PLACA','OBRA_DESTINO','GASOLINERA',
              'IMPORTE_AUTORIZADO','CONSUMO_REAL','LITROS_REALES','DIFERENCIA','ESTATUS']
final = merged[final_cols].sort_values(['SEMANA','OBRA_DESTINO','RESPONSABLE'])

# ---- Write to Excel ----
wb = load_workbook(maestro_path)

# Remove existing sheet if present
if 'BD_ESTADOS_CUENTA' in wb.sheetnames:
    del wb['BD_ESTADOS_CUENTA']

ws = wb.create_sheet('BD_ESTADOS_CUENTA', 3)  # after BD_FACTURAS

# Colors
HDR_FILL = PatternFill('solid', fgColor='1A1A2E')
HDR_FONT = Font(bold=True, color='FFFFFF', size=10)
OK_FILL  = PatternFill('solid', fgColor='D4EDDA')
ALT_FILL = PatternFill('solid', fgColor='F8D7DA')
SEM_FILLS = {26: PatternFill('solid', fgColor='EBF5FB'), 27: PatternFill('solid', fgColor='FDF2E9')}
CUR_NUM  = '#,##0.00'
thin = Border(left=Side(style='thin',color='CCCCCC'), right=Side(style='thin',color='CCCCCC'),
              top=Side(style='thin',color='CCCCCC'), bottom=Side(style='thin',color='CCCCCC'))

# Header row
headers = ['SEMANA','RESPONSABLE','VEHICULO','PLACA','FRENTE DE TRABAJO','GASOLINERA',
           'IMPORTE AUTORIZADO','CONSUMO REAL','LITROS REALES','DIFERENCIA (REMANENTE)','ESTATUS']
col_widths = [10,35,22,14,40,22,18,14,13,20,20]

for c, (h, w) in enumerate(zip(headers, col_widths), 1):
    cell = ws.cell(1, c, h)
    cell.fill = HDR_FILL
    cell.font = HDR_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin
    ws.column_dimensions[get_column_letter(c)].width = w

ws.row_dimensions[1].height = 30

# Data rows
for r, row in enumerate(final.itertuples(index=False), 2):
    sem = int(row.SEMANA)
    is_alt = row.ESTATUS == 'ALERTA - Excedido'
    row_fill = ALT_FILL if is_alt else SEM_FILLS.get(sem, PatternFill())
    
    vals = [sem, row.RESPONSABLE, row.VEHICULO, row.PLACA, row.OBRA_DESTINO, row.GASOLINERA,
            row.IMPORTE_AUTORIZADO, row.CONSUMO_REAL, row.LITROS_REALES, row.DIFERENCIA, row.ESTATUS]
    
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v)
        cell.border = thin
        cell.fill = row_fill
        cell.alignment = Alignment(vertical='center')
        if c in (7,8,10):
            cell.number_format = CUR_NUM
        elif c == 9:
            cell.number_format = '#,##0.00'
        if is_alt and c == 11:
            cell.font = Font(bold=True, color='C0392B')
        elif c == 11:
            cell.font = Font(color='1E8449')

# Totals
def add_total_row(ws, semana, rows_data, start_row):
    sub = rows_data[rows_data['SEMANA'] == semana]
    if len(sub) == 0:
        return
    cell = ws.cell(start_row, 1, f'TOTAL SEMANA {semana}')
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='2C3E50')
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=6)
    for c, val in [(7, sub.IMPORTE_AUTORIZADO.sum()), (8, sub.CONSUMO_REAL.sum()),
                   (9, sub.LITROS_REALES.sum()), (10, sub.DIFERENCIA.sum())]:
        tc = ws.cell(start_row, c, round(val,2))
        tc.font = Font(bold=True, color='FFFFFF')
        tc.fill = PatternFill('solid', fgColor='2C3E50')
        tc.number_format = CUR_NUM if c != 9 else '#,##0.00'
        tc.border = thin

r_now = len(final) + 2
sub26 = final[final.SEMANA == 26]
sub27 = final[final.SEMANA == 27]

# Insert a blank row + total between weeks
# Actually just put totals at bottom
add_total_row(ws, 26, final, r_now)
r_now += 1
add_total_row(ws, 27, final, r_now)
r_now += 1
# Grand total
ws.cell(r_now, 1, 'GRAN TOTAL').font = Font(bold=True, color='FFFFFF')
ws.cell(r_now, 1).fill = PatternFill('solid', fgColor='1A1A2E')
ws.merge_cells(start_row=r_now, start_column=1, end_row=r_now, end_column=6)
for c, val in [(7, final.IMPORTE_AUTORIZADO.sum()), (8, final.CONSUMO_REAL.sum()),
               (9, final.LITROS_REALES.sum()), (10, final.DIFERENCIA.sum())]:
    tc = ws.cell(r_now, c, round(val,2))
    tc.font = Font(bold=True, color='FFFF00')
    tc.fill = PatternFill('solid', fgColor='1A1A2E')
    tc.number_format = CUR_NUM if c != 9 else '#,##0.00'
    tc.border = thin

wb.save(maestro_path)
print('=== DONE ===')
print(f'Total rows written: {len(final)}')
print(f'Semana 26 - Auth: {sub26.IMPORTE_AUTORIZADO.sum():,.2f} | Consumo: {sub26.CONSUMO_REAL.sum():,.2f}')
print(f'Semana 27 - Auth: {sub27.IMPORTE_AUTORIZADO.sum():,.2f} | Consumo: {sub27.CONSUMO_REAL.sum():,.2f}')
print(f'Gran Total Autorizado: {final.IMPORTE_AUTORIZADO.sum():,.2f}')
