import openpyxl

excel_path = r"c:\Users\JOSE\Desktop\Proyecto fenix\FACTURAS_EXTRAIDAS_REVISION.xlsx"

try:
    wb = openpyxl.load_workbook(excel_path, data_only=False)
    sheet = wb.active
    
    headers = [cell.value for cell in sheet[1]]
    print("Columnas actuales en FACTURAS_EXTRAIDAS_REVISION.xlsx:")
    print(headers)
    
    print("\nRevisando las primeras 3 filas para ver formulas:")
    for row in sheet.iter_rows(min_row=2, max_row=4):
        row_data = []
        for cell in row:
            val = str(cell.value)
            if val.startswith('='):
                row_data.append(f"FORMULA[{val}]")
            else:
                # Just show a preview if not formula
                row_data.append(val[:15])
        print(row_data)
        
except Exception as e:
    print(f"Error revisando FACTURAS_EXTRAIDAS_REVISION.xlsx: {e}")

# Check Maestro file as well
excel_path_maestro = r"c:\Users\JOSE\Desktop\Proyecto fenix\MAESTRO_CONTROL_DIESEL_NUEVO.xlsx"
try:
    wb2 = openpyxl.load_workbook(excel_path_maestro, data_only=False)
    if "BD_FACTURAS" in wb2.sheetnames:
        sheet2 = wb2["BD_FACTURAS"]
        headers2 = [cell.value for cell in sheet2[1]]
        print("\nColumnas actuales en BD_FACTURAS (Maestro):")
        print(headers2)
except Exception as e:
    print(f"Error revisando Maestro: {e}")
