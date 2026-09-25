import openpyxl
import datetime
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')
ws = wb['Captura_Facturas_Mobil']

fecha_col = None
semana_col = None
for col in range(1, 20):
    val = ws.cell(row=4, column=col).value
    if val == 'Fecha': fecha_col = col
    elif val == 'Semana (Helper)': semana_col = col

if fecha_col and semana_col:
    for r in range(5, ws.max_row + 1):
        dt = ws.cell(row=r, column=fecha_col).value
        if isinstance(dt, datetime.datetime):
            week = dt.isocalendar().week
            ws.cell(row=r, column=semana_col).value = week
            print(f'Row {r}: {dt.date()} -> Week {week}')

wb.save('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')
print('Hardcoded week integers!')

