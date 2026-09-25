import sys
sys.path.insert(0, r'c:\Users\JOSE\Desktop\Proyecto fenix')

from app_admin import app
import openpyxl

with app.test_client() as client:
    # Set session
    with client.session_transaction() as sess:
        sess['user_id'] = 1
        sess['username'] = 'admin'
        sess['user_role'] = 'ADMIN'

    resp = client.get('/api/admin/tags/reporte/excel')
    print("Status code:", resp.status_code)
    print("Content length:", len(resp.data))
    print("Content-type:", resp.content_type)
    
    if resp.status_code == 200:
        filename = 'Reporte_TAGS_Test_Client.xlsx'
        with open(filename, 'wb') as f:
            f.write(resp.data)
        print("Saved Excel successfully!")

        wb = openpyxl.load_workbook(filename)
        print("\nSheets in generated Excel:", wb.sheetnames)

        # Inspect 'TAGs Sin Autorización'
        if 'TAGs Sin Autorización' in wb.sheetnames:
            ws = wb['TAGs Sin Autorización']
            print(f"\n--- Sheet 'TAGs Sin Autorización' (Rows: {ws.max_row}, Cols: {ws.max_column}) ---")
            for r in range(1, min(25, ws.max_row + 1)):
                vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
                print(f"Row {r:2d}: {vals}")

        # Inspect 'SEMANA 32'
        if 'SEMANA 32' in wb.sheetnames:
            ws = wb['SEMANA 32']
            print(f"\n--- Sheet 'SEMANA 32' (Rows: {ws.max_row}, Cols: {ws.max_column}) ---")
            for r in range(max(1, ws.max_row - 18), ws.max_row + 1):
                vals = [ws.cell(r, c).value for c in range(1, min(7, ws.max_column + 1))]
                print(f"Row {r:2d}: {vals}")
    else:
        print("Error:", resp.data[:500])
