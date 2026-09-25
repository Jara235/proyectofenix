import openpyxl

wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/MAQUINARÍA 2026.xlsx', data_only=True)
print("Sheet names in MAQUINARÍA 2026.xlsx:", wb.sheetnames)
ws = wb.active
for r in range(1, 25):
    row_vals = [cell.value for cell in ws[r]]
    if any(row_vals):
        print(f"Row {r}: {row_vals}")
