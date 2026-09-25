import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

filename = 'c:/Users/JOSE/Desktop/Proyecto fenix/Estimacion_Combustible_MexicoToluca_Lerma_Ene_Jun.xlsx'
wb = openpyxl.load_workbook(filename)

# Add a dedicated Reference sheet or insert Reference Block in ws1
ws_ref = wb.create_sheet(title="Promedios de Referencia", index=0)
ws_ref.views.sheetView[0].showGridLines = True

def fill(hex_color): return PatternFill("solid", fgColor=hex_color)
def font(bold=False, italic=False, color='000000', size=10, name='Calibri'): return Font(bold=bold, italic=italic, color=color, size=size, name=name)
def align(h='center', v='center', wrap=False): return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def thin_border():
    s = Side(style='thin', color='CBD5E1')
    return Border(left=s, right=s, top=s, bottom=s)

WHITE = 'FFFFFF'
NAVY = '0F172A'

# Header
ws_ref.merge_cells('A1:G1')
ws_ref['A1'] = "📌 PROMEDIOS REALES DE REFERENCIA MENSUAL Y SEMANAL"
ws_ref['A1'].fill = fill(NAVY); ws_ref['A1'].font = font(bold=True, color=WHITE, size=13); ws_ref['A1'].alignment = align('left', 'center')
ws_ref.row_dimensions[1].height = 26

ws_ref.merge_cells('A2:G2')
ws_ref['A2'] = "Datos normalizados calculados a partir de los registros históricos reales de campo y facturación | Fénix 2.0"
ws_ref['A2'].fill = fill('1E293B'); ws_ref['A2'].font = font(italic=True, color='94A3B8', size=9); ws_ref['A2'].alignment = align('left', 'center')
ws_ref.row_dimensions[2].height = 18

# Table 1: Promedios Mensuales de Referencia (Mes Típico 30.4 días / 4.333 semanas)
ws_ref.merge_cells('A4:G4')
ws_ref['A4'] = "1. PROMEDIO REAL MENSUAL DE REFERENCIA (MES TÍPICO DE OPERACIÓN)"
ws_ref['A4'].fill = fill('1E3A8A'); ws_ref['A4'].font = font(bold=True, color=WHITE, size=10); ws_ref['A4'].alignment = align('left', 'center')
ws_ref.row_dimensions[4].height = 20

headers_t1 = ["OBRA", "COMBUSTIBLE", "PROMEDIO SEMANAL (L)", "PROMEDIO MENSUAL (L)", "PRECIO PROM. ($/L)", "IMPORTE MENSUAL ($)", "PARTICIPACIÓN %"]
for c_i, h in enumerate(headers_t1, 1):
    c = ws_ref.cell(row=5, column=c_i, value=h)
    c.fill = fill('334155'); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center'); c.border = thin_border()
ws_ref.row_dimensions[5].height = 20

data_t1 = [
    ("MÉXICO - TOLUCA", "DIESEL", 2511.2, 10881.9, 27.00, 293810.40, "49.6%"),
    ("MÉXICO - TOLUCA", "GASOLINA", 392.8, 1701.9, 23.90, 40676.00, "6.9%"),
    ("MÉXICO - TOLUCA (SUBTOTAL)", "AMBOS", 2904.0, 12583.8, 26.58, 334486.40, "56.4%"),
    
    ("LERMA - TRES MARÍAS", "DIESEL", 1852.6, 8027.9, 27.00, 216753.00, "36.6%"),
    ("LERMA - TRES MARÍAS", "GASOLINA", 401.1, 1738.0, 23.90, 41538.50, "7.0%"),
    ("LERMA - TRES MARÍAS (SUBTOTAL)", "AMBOS", 2253.7, 9765.9, 26.45, 258291.50, "43.6%"),
]

r_curr = 6
for idx, d in enumerate(data_t1):
    is_sub = "SUBTOTAL" in d[0]
    bg = 'E2E8F0' if is_sub else ('FFFFFF' if idx % 2 == 0 else 'F8FAFC')
    
    ws_ref.cell(row=r_curr, column=1, value=d[0]).font = font(bold=is_sub, size=9)
    ws_ref.cell(row=r_curr, column=2, value=d[1]).font = font(bold=is_sub, color='1E40AF' if d[1]=='DIESEL' else ('047857' if d[1]=='GASOLINA' else '000000'), size=9)
    ws_ref.cell(row=r_curr, column=3, value=d[2]).number_format = '#,##0.0'
    ws_ref.cell(row=r_curr, column=4, value=d[3]).number_format = '#,##0.0'
    ws_ref.cell(row=r_curr, column=5, value=d[4]).number_format = '"$"#,##0.00'
    ws_ref.cell(row=r_curr, column=6, value=d[5]).number_format = '"$"#,##0.00'
    ws_ref.cell(row=r_curr, column=7, value=d[6])
    
    for c in range(1, 8):
        cell = ws_ref.cell(row=r_curr, column=c)
        cell.fill = fill(bg)
        cell.border = thin_border()
        if c in [1, 2]: cell.alignment = align('left' if c==1 else 'center', 'center')
        elif c == 7: cell.alignment = align('center', 'center')
        else: cell.alignment = align('right', 'center')
    ws_ref.row_dimensions[r_curr].height = 18
    r_curr += 1

# Total General
ws_ref.cell(row=r_curr, column=1, value="GRAN TOTAL CONSOLIDADO MENSUAL").fill = fill(NAVY)
ws_ref.cell(row=r_curr, column=1).font = font(bold=True, color=WHITE, size=9); ws_ref.cell(row=r_curr, column=1).alignment = align('left', 'center')
ws_ref.cell(row=r_curr, column=2, value="AMBAS OBRAS").fill = fill(NAVY)
ws_ref.cell(row=r_curr, column=2).font = font(bold=True, color=WHITE, size=9); ws_ref.cell(row=r_curr, column=2).alignment = align('center', 'center')

ws_ref.cell(row=r_curr, column=3, value=5157.7).number_format = '#,##0.0'
ws_ref.cell(row=r_curr, column=4, value=22349.7).number_format = '#,##0.0'
ws_ref.cell(row=r_curr, column=5, value=26.52).number_format = '"$"#,##0.00'

c_tot = ws_ref.cell(row=r_curr, column=6, value=592777.90)
c_tot.number_format = '"$"#,##0.00'
c_tot.font = font(bold=True, color='000000', size=10)
c_tot.fill = fill('FDE047')

ws_ref.cell(row=r_curr, column=7, value="100.0%").fill = fill(NAVY)
ws_ref.cell(row=r_curr, column=7).font = font(bold=True, color=WHITE, size=9); ws_ref.cell(row=r_curr, column=7).alignment = align('center', 'center')

for c in range(1, 8):
    cell = ws_ref.cell(row=r_curr, column=c)
    if c not in [6]:
        cell.fill = fill(NAVY)
        cell.font = font(bold=True, color=WHITE, size=9)
    if c in [3, 4, 5]: cell.alignment = align('right', 'center')
    cell.border = thin_border()
c_tot.border = thin_border(); c_tot.alignment = align('right', 'center')
ws_ref.row_dimensions[r_curr].height = 22

# Auto adjust columns for ws_ref
for col in ws_ref.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        if '\n' in val_str:
            val_str = max(val_str.split('\n'), key=len)
        max_len = max(max_len, len(val_str))
    ws_ref.column_dimensions[col_letter].width = max(max_len + 3, 14)

wb.save(filename)
print("Updated Excel file with Reference sheet!")
