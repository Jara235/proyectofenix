import openpyxl

wb = openpyxl.load_workbook('gasolina/conciliacion_gasolina_temp.xlsx', data_only=True)
ws = wb['SEMANA 33']

excedidos = []
for r in range(128, 140):
    resp = ws.cell(r, 2).value # Col B
    ct = ws.cell(r, 3).value   # Col C
    uni = ws.cell(r, 4).value  # Col D
    pla = ws.cell(r, 5).value  # Col E
    aut = ws.cell(r, 6).value  # Col F
    con = ws.cell(r, 7).value  # Col G
    exc = ws.cell(r, 8).value  # Col H
    
    if resp:
        excedidos.append({
            'row': r,
            'responsable': str(resp).strip(),
            'centro_trabajo': str(ct or '').strip(),
            'unidad': str(uni or '').strip(),
            'placa': str(pla or 'S/P').strip(),
            'autorizado': float(aut or 0),
            'consumo': float(con or 0),
            'excedente': float(exc or 0)
        })

print("=" * 125)
print("TABLA OFICIAL DE CONSUMOS EXCEDIDOS DE GASOLINA - SEMANA 33 (10 AL 15 DE AGOSTO 2026)")
print("=" * 125)
print(f"{'#':<2} | {'Responsable':<28} | {'Centro de Trabajo':<27} | {'Unidad':<20} | {'Placas':<8} | {'Autorizado':<10} | {'Consumo':<10} | {'Excedido':<10}")
print("-" * 125)

tot_aut = 0
tot_con = 0
tot_exc = 0

for idx, e in enumerate(excedidos, 1):
    tot_aut += e['autorizado']
    tot_con += e['consumo']
    tot_exc += abs(e['excedente'])
    print(f"{idx:<2} | {e['responsable']:<28} | {e['centro_trabajo']:<27} | {e['unidad']:<20} | {e['placa']:<8} | ${e['autorizado']:>8.2f} | ${e['consumo']:>8.2f} | -${abs(e['excedente']):>8.2f}")

print("-" * 125)
print(f"TOTALES: {len(excedidos)} unidades excedidas | Autorizado: ${tot_aut:,.2f} | Consumo: ${tot_con:,.2f} | Total Excedido: -${tot_exc:,.2f}")
print("=" * 125)
