import sqlite3
db = sqlite3.connect('c:/Users/JOSE/Desktop/Proyecto fenix/fenix.db')
db.execute("UPDATE fenix_facturas_documentos SET semana=26 WHERE folio='10320'")
db.commit()

import openpyxl
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')
ws = wb['Captura_Facturas_Mobil']
for i in range(5, ws.max_row+1):
    if ws.cell(row=i, column=1).value == 'A-10320':
        ws.cell(row=i, column=14).value = 26
        print('Updated A-10320 to week 26 in Excel')
wb.save('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_LIMPIO.xlsx')

