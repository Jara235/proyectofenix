import openpyxl
from copy import copy
from openpyxl.utils import get_column_letter

# Cargar el archivo de Google Sheets local
wb_gs = openpyxl.load_workbook('gasolina/conciliacion_gasolina_google_sheets.xlsx')

sheet_title = "CONCILIACION_PAGOS_MOBIL"
if sheet_title in wb_gs.sheetnames:
    del wb_gs[sheet_title]

ws = wb_gs.create_sheet(title=sheet_title)
ws.views.sheetView[0].showGridLines = True

wb_master = openpyxl.load_workbook('formatos/CONCILIACION_PAGOS_MOBIL_SEMANAS_28_A_32.xlsx')
ws_master = wb_master['Conciliacion_Facturas']
ws_res = wb_master['Resumen_Ejecutivo']

# Copiar resumen y detalle
def copy_cells(src_sheet, dest_sheet, start_src_row, end_src_row, start_dest_row, max_col):
    for r in range(start_src_row, end_src_row + 1):
        dest_r = start_dest_row + (r - start_src_row)
        for c in range(1, max_col + 1):
            cell_orig = src_sheet.cell(r, c)
            cell_dest = dest_sheet.cell(dest_r, c, value=cell_orig.value)
            if cell_orig.has_style:
                if cell_orig.font: cell_dest.font = copy(cell_orig.font)
                if cell_orig.fill: cell_dest.fill = copy(cell_orig.fill)
                if cell_orig.border: cell_dest.border = copy(cell_orig.border)
                if cell_orig.alignment: cell_dest.alignment = copy(cell_orig.alignment)
                if cell_orig.number_format: cell_dest.number_format = copy(cell_orig.number_format)
        dest_sheet.row_dimensions[dest_r].height = src_sheet.row_dimensions[r].height

# Copiar encabezado y tabla de resumen
copy_cells(ws_res, ws, 1, 16, 1, 8)

# Copiar detalle
start_det_r = 19
ws.cell(start_det_r - 1, 1, value="").fill = copy(ws_res.cell(1, 1).fill)
copy_cells(ws_master, ws, 1, ws_master.max_row, start_det_r, 9)

# Ajustar anchos
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        if cell.coordinate in ws.merged_cells:
            continue
        if len(val_str) > max_len:
            max_len = len(val_str)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

wb_gs.save('gasolina/conciliacion_gasolina_google_sheets.xlsx')
print("Pestaña CONCILIACION_PAGOS_MOBIL agregada exitosamente a gasolina/conciliacion_gasolina_google_sheets.xlsx")
