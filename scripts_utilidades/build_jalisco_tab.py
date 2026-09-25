import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import shutil

# Respaldar el archivo por seguridad
shutil.copy('MAESTRO_CONTROL_DIESEL_NUEVO.xlsx', 'MAESTRO_CONTROL_DIESEL_NUEVO_backup.xlsx')

print("Leyendo datos de BD_DIESEL...")
df_diesel = pd.read_excel('MAESTRO_CONTROL_DIESEL_NUEVO.xlsx', sheet_name='BD_DIESEL')

# Filtrar solo Jalisco
df_jalisco = df_diesel[df_diesel['OBRA_DESTINO'].str.contains('Jalisco', na=False, case=False)]

# Agrupar por Semana
resumen = df_jalisco.groupby('SEMANA').agg({
    'LITROS': 'sum',
    'IMPORTE_TOTAL': 'sum'
}).reset_index()

# Ordenar por semana
resumen = resumen.sort_values('SEMANA')

print("Abriendo el archivo de Excel para escribir...")
wb = openpyxl.load_workbook('MAESTRO_CONTROL_DIESEL_NUEVO.xlsx')

sheet_name = 'BD_Jalisco'
if sheet_name in wb.sheetnames:
    del wb[sheet_name]

ws = wb.create_sheet(title=sheet_name)

# Estilos
header_fill = PatternFill(start_color="6F8E45", end_color="6F8E45", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
border_side = Side(border_style="thin", color="000000")
border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# Títulos superiores
ws.merge_cells('A1:I1')
ws['A1'] = 'RESUMEN DE CONSUMO SEMANAL POR OBRA (DETALLE HISTÓRICO POR SEMANA)'
ws['A1'].font = Font(bold=True, size=12)

ws['B3'] = 'Obra (Tabla Semanal):'
ws['B3'].font = Font(bold=True)
ws['C3'] = 'Jalisco'
ws['C3'].fill = PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid")
ws['C3'].border = border

ws['B4'] = 'Precio Diésel ($/L):'
ws['B4'].font = Font(bold=True)
ws['C4'] = 27.00
ws['C4'].number_format = '"$"#,##0.00'
ws['C4'].fill = PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid")
ws['C4'].border = border

# Encabezados de tabla
headers = [
    "Semana (Lunes)", "Saldo Inicial (Vales)", "Litros Despachados", 
    "Costo/L ($)", "Costo Combustible ($)", "Gasolina", 
    "Saldo Final (Vales)", "Ajuste semanal", "Litros de Diesel Acumulados"
]

for col_num, header_text in enumerate(headers, 1):
    cell = ws.cell(row=6, column=col_num)
    cell.value = header_text
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = align_center
    cell.border = border
    ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 20

# Escribir datos
start_row = 7
acumulado_litros = 0

# Set previous week manual data based on the screenshot
# We need to map Semana -> Start Date or just use the Semana text.
# The screenshot uses dates like 2026-05-25 (Week 22), 2026-06-01 (Week 23), etc.
# But Fenix uses "Semana 26", "Semana 27". We'll just use the Semana string from DB.

# Let's collect the DB data into a dictionary for easy access
db_data = {row['SEMANA']: {'litros': row['LITROS'], 'costo': row['IMPORTE_TOTAL']} for _, row in resumen.iterrows()}

# We know the historical manual data from the screenshot for earlier weeks:
# Row 1 (Semana 22? 2026-05-25): SaldoIni 21600, Lts 556, Ajuste 132, Gasolina 0
# Row 2 (Semana 23? 2026-06-01): SaldoIni 6720, Lts 824, Ajuste 276, Gasolina 3000
# Row 3 (Semana 24? 2026-06-08): SaldoIni -18252, Lts 180, Ajuste 72, Gasolina 1556
# Row 4 (Semana 25? 2026-06-15): SaldoIni -24596, Lts 690, Ajuste 200, Gasolina 3700

historico = [
    {"semana": "Semana 22", "saldo_ini": 21600.0, "lts": 556, "ajuste": 132, "gasolina": 0},
    {"semana": "Semana 23", "saldo_ini": 6720.0, "lts": 824, "ajuste": 276, "gasolina": 3000},
    {"semana": "Semana 24", "saldo_ini": -18252.0, "lts": 180, "ajuste": 72, "gasolina": 1556},
    {"semana": "Semana 25", "saldo_ini": -24596.0, "lts": 690, "ajuste": 200, "gasolina": 3700},
]

current_row = start_row
litros_acumulados = 0

# Escribir el historico primero
for h in historico:
    ws.cell(row=current_row, column=1, value=h["semana"]).alignment = align_center
    ws.cell(row=current_row, column=2, value=h["saldo_ini"]).number_format = '"$"#,##0.00'
    ws.cell(row=current_row, column=3, value=h["lts"]).alignment = align_center
    ws.cell(row=current_row, column=4, value=27.00).number_format = '"$"#,##0.00'
    
    # Formula Costo Combustible: Litros * Costo/L
    ws.cell(row=current_row, column=5, value=f"=C{current_row}*D{current_row}").number_format = '"$"#,##0.00'
    
    ws.cell(row=current_row, column=6, value=h["gasolina"]).number_format = '"$"#,##0.00'
    
    # Formula Saldo Final: Saldo Inicial - Costo Combustible - Gasolina + Ajuste
    ws.cell(row=current_row, column=7, value=f"=B{current_row}-E{current_row}-F{current_row}+H{current_row}").number_format = '"$"#,##0.00'
    
    ws.cell(row=current_row, column=8, value=h["ajuste"]).number_format = '"$"#,##0.00'
    
    # Formula Acumulado
    if current_row == start_row:
        ws.cell(row=current_row, column=9, value=f"=C{current_row}")
    else:
        ws.cell(row=current_row, column=9, value=f"=C{current_row}+I{current_row-1}")
        
    for col in range(1, 10): ws.cell(row=current_row, column=col).border = border
    current_row += 1

# Ahora agregar las semanas que están en la base de datos (Ej: Semana 26, Semana 27, etc)
# Excluyendo las que ya metimos en el historico
semanas_hist = [h['semana'] for h in historico]

for semana in sorted(db_data.keys()):
    if semana in semanas_hist: continue # Skip if already in historical manual
    
    data = db_data[semana]
    lts = data['litros']
    
    ws.cell(row=current_row, column=1, value=semana).alignment = align_center
    # Formula Saldo Inicial = Saldo Final de la semana anterior + Ajuste (We just take Saldo Final of previous row)
    ws.cell(row=current_row, column=2, value=f"=G{current_row-1}").number_format = '"$"#,##0.00'
    
    ws.cell(row=current_row, column=3, value=lts).alignment = align_center
    ws.cell(row=current_row, column=4, value=27.00).number_format = '"$"#,##0.00'
    ws.cell(row=current_row, column=5, value=f"=C{current_row}*D{current_row}").number_format = '"$"#,##0.00'
    
    ws.cell(row=current_row, column=6, value=0.0).number_format = '"$"#,##0.00' # Gasolina default 0
    ws.cell(row=current_row, column=7, value=f"=B{current_row}-E{current_row}-F{current_row}+H{current_row}").number_format = '"$"#,##0.00'
    ws.cell(row=current_row, column=8, value=0.0).number_format = '"$"#,##0.00' # Ajuste default 0
    ws.cell(row=current_row, column=9, value=f"=C{current_row}+I{current_row-1}")
    
    for col in range(1, 10): ws.cell(row=current_row, column=col).border = border
    current_row += 1

# Fila de Totales
ws.cell(row=current_row, column=1, value="TOTAL ACUMULADO").font = Font(bold=True)
ws.cell(row=current_row, column=1).fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")

# Sumatorias
ws.cell(row=current_row, column=2, value=f"=B{start_row}").number_format = '"$"#,##0.00' # El saldo inicial total es el del dia 1
ws.cell(row=current_row, column=3, value=f"=SUM(C{start_row}:C{current_row-1})").alignment = align_center
ws.cell(row=current_row, column=5, value=f"=SUM(E{start_row}:E{current_row-1})").number_format = '"$"#,##0.00'
ws.cell(row=current_row, column=6, value=f"=SUM(F{start_row}:F{current_row-1})").number_format = '"$"#,##0.00'
ws.cell(row=current_row, column=7, value=f"=G{current_row-1}").number_format = '"$"#,##0.00' # Saldo final es el ultimo
ws.cell(row=current_row, column=8, value=f"=SUM(H{start_row}:H{current_row-1})").number_format = '"$"#,##0.00'
ws.cell(row=current_row, column=9, value=f"=I{current_row-1}")

for col in range(1, 10): 
    cell = ws.cell(row=current_row, column=col)
    cell.border = border
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")

print("Guardando archivo final...")
wb.save('MAESTRO_CONTROL_DIESEL_NUEVO.xlsx')
print("Exito!")
