import openpyxl

wb = openpyxl.load_workbook('formatos/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_FINAL_V2.xlsx', data_only=False)
ws = wb['Maquinaria_Inventario']

print("=== COLUMNAS Y FORMULAS DE Maquinaria_Inventario ===")
for r in range(1, 20):
    row_vals = [ws.cell(r, c).value for c in range(1, 14)]
    print(f"Row {r:02d}:", row_vals)

ws2 = wb['Catalogos']
print("\n=== HOJA: Catalogos (Formulas / Valores) ===")
for r in range(1, 30):
    row_vals = [ws2.cell(r, c).value for c in range(1, 15) if ws2.cell(r, c).value is not None]
    if row_vals:
        print(f"Row {r:02d}:", row_vals)
