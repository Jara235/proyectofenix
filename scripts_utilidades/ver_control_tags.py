import openpyxl, os

for f in ['CONTROL TAG.xlsx', 'CONTROL_MAESTRO_TAGS.xlsx']:
    path = os.path.join('TAGS', f)
    if os.path.exists(path):
        wb = openpyxl.load_workbook(path, data_only=True)
        print(f"\n============================\nARCHIVO: {path}\nHOJAS: {wb.sheetnames}")
        for sname in wb.sheetnames[:3]:
            ws = wb[sname]
            print(f"\n--- Hoja: {sname} (filas: {ws.max_row}, cols: {ws.max_column}) ---")
            for r in range(1, min(ws.max_row+1, 15)):
                row_vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, min(ws.max_column+1, 12))]
                if any(row_vals):
                    print(" | ".join(row_vals))
