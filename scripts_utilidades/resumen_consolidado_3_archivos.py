import openpyxl, os

files = {
    'MÉXICO-TOLUCA': 'acarreos/OBRA MEXICO TOLUCA 2026 (6).xlsx',
    'LERMA-3 MARÍAS': 'acarreos/OBRA LERMA 3 MARIAS 2026 (13).xlsx',
    'ALFREDO DEL MAZO': 'acarreos/OBRA ALFREDO DEL MAZO (9).xlsx'
}

def get_summary():
    total_general_viajes = 0
    total_general_m3 = 0.0
    total_general_ton = 0.0
    total_general_importe = 0.0
    
    resumen_obras = {}
    
    for obra, fpath in files.items():
        wb = openpyxl.load_workbook(fpath, data_only=True)
        resumen_obras[obra] = {'hojas': {}, 'total_viajes': 0, 'total_m3': 0.0, 'total_ton': 0.0, 'total_importe': 0.0}
        
        for sname in wb.sheetnames:
            if sname.lower().startswith('tabla') or sname.lower().startswith('t-') or sname.lower() in ['edo cuenta ', 'resum']:
                continue
            ws = wb[sname]
            
            # Encabezado
            hrow = 1
            headers = [str(ws.cell(hrow, c).value or '').upper().strip() for c in range(1, 30)]
            
            idx_sem = None
            idx_folio = None
            idx_fecha = None
            idx_mat = None
            idx_placas = None
            idx_m3 = None
            idx_kg = None
            idx_imp = None
            
            for i, h in enumerate(headers, 1):
                if 'SEM' in h and idx_sem is None: idx_sem = i
                elif 'FOLIO' in h and idx_folio is None: idx_folio = i
                elif 'FECHA' in h and idx_fecha is None: idx_fecha = i
                elif 'MATERIAL' in h and idx_mat is None: idx_mat = i
                elif 'PLACA' in h and idx_placas is None: idx_placas = i
                elif ('M3' in h or 'CAPACIDAD' in h) and idx_m3 is None: idx_m3 = i
                elif ('PESO' in h or 'KG' in h) and idx_kg is None: idx_kg = i
                elif 'IMPORTE' in h and idx_imp is None: idx_imp = i
                
            viajes = 0
            m3_sum = 0.0
            ton_sum = 0.0
            imp_sum = 0.0
            semanas_dict = {}
            
            for r in range(hrow + 1, ws.max_row + 1):
                val_sem = ws.cell(r, idx_sem).value if idx_sem else None
                val_placa = ws.cell(r, idx_placas).value if idx_placas else None
                val_fecha = ws.cell(r, idx_fecha).value if idx_fecha else None
                if val_sem is None and val_placa is None and val_fecha is None:
                    continue
                
                # Semana
                sem_str = 'S/D'
                if val_sem:
                    try:
                        sem_str = str(int(float(str(val_sem).replace('SEM', '').replace('#', '').strip())))
                    except:
                        sem_str = str(val_sem).strip()
                        
                val_m3 = ws.cell(r, idx_m3).value if idx_m3 else 0.0
                val_kg = ws.cell(r, idx_kg).value if idx_kg else 0.0
                val_imp = ws.cell(r, idx_imp).value if idx_imp else 0.0
                
                try: m3_f = float(val_m3 or 0.0)
                except: m3_f = 0.0
                try: kg_f = float(val_kg or 0.0)
                except: kg_f = 0.0
                try: imp_f = float(val_imp or 0.0)
                except: imp_f = 0.0
                
                # Toneladas
                if kg_f > 0:
                    ton_f = kg_f / 1000.0
                elif 'CARPETA' in sname.upper() or 'MEZCLA' in sname.upper():
                    ton_f = m3_f * 2.35 # estimación si no viene kg
                else:
                    ton_f = 0.0
                    
                viajes += 1
                m3_sum += m3_f
                ton_sum += ton_f
                imp_sum += imp_f
                
                if sem_str not in semanas_dict:
                    semanas_dict[sem_str] = {'viajes': 0, 'm3': 0.0, 'ton': 0.0, 'importe': 0.0}
                semanas_dict[sem_str]['viajes'] += 1
                semanas_dict[sem_str]['m3'] += m3_f
                semanas_dict[sem_str]['ton'] += ton_f
                semanas_dict[sem_str]['importe'] += imp_f
                
            resumen_obras[obra]['hojas'][sname] = {
                'viajes': viajes, 'm3': m3_sum, 'ton': ton_sum, 'importe': imp_sum, 'semanas': semanas_dict
            }
            resumen_obras[obra]['total_viajes'] += viajes
            resumen_obras[obra]['total_m3'] += m3_sum
            resumen_obras[obra]['total_ton'] += ton_sum
            resumen_obras[obra]['total_importe'] += imp_sum
            
            total_general_viajes += viajes
            total_general_m3 += m3_sum
            total_general_ton += ton_sum
            total_general_importe += imp_sum

    print("\n" + "=" * 95)
    print("RESUMEN GLOBAL CONSOLIDADO DE LOS 3 ARCHIVOS")
    print("=" * 95)
    print(f"* Total Viajes Registrados:  {total_general_viajes:,} viajes")
    print(f"* Volumen Total Acarreado:   {total_general_m3:,.2f} m3")
    print(f"* Toneladas Totales Mezcla:  {total_general_ton:,.2f} Toneladas")
    print(f"* Importe Total Fletes:      ${total_general_importe:,.2f} MXN\n")
    
    for obra, info in resumen_obras.items():
        print("-" * 95)
        print(f"OBRA: {obra} (Viajes: {info['total_viajes']} | Vol: {info['total_m3']:,.2f} m3 | Mezcla: {info['total_ton']:,.2f} Ton | Fletes: ${info['total_importe']:,.2f})")
        print("-" * 95)
        for hname, hdata in info['hojas'].items():
            print(f"  > Concepto: {hname:<18} | Viajes: {hdata['viajes']:<5} | Vol: {hdata['m3']:>10,.2f} m3 | Ton: {hdata['ton']:>9,.2f} | Importe: ${hdata['importe']:>11,.2f}")
            sems = sorted([int(k) for k in hdata['semanas'].keys() if k.isdigit()])
            print(f"    Semanas: {sems}")

get_summary()
