import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

filename = 'c:/Users/JOSE/Desktop/Proyecto fenix/Estimacion_Combustible_MexicoToluca_Lerma_Ene_Jun.xlsx'
wb = openpyxl.load_workbook(filename)

# Check if sheet exists and remove/replace
if "Comparativo vs Referencia" in wb.sheetnames:
    del wb["Comparativo vs Referencia"]

ws_comp = wb.create_sheet(title="Comparativo vs Referencia", index=0)
ws_comp.views.sheetView[0].showGridLines = True

def fill(hex_color): return PatternFill("solid", fgColor=hex_color)
def font(bold=False, italic=False, color='000000', size=10, name='Calibri'): return Font(bold=bold, italic=italic, color=color, size=size, name=name)
def align(h='center', v='center', wrap=False): return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def thin_border():
    s = Side(style='thin', color='CBD5E1')
    return Border(left=s, right=s, top=s, bottom=s)

WHITE = 'FFFFFF'
NAVY = '0F172A'
BLUE_HDR = '1E3A8A'
GREEN_HDR = '065F46'
RED_TEXT = '991B1B'
GREEN_TEXT = '166534'

# Title
ws_comp.merge_cells('A1:J1')
ws_comp['A1'] = "📊 COMPARATIVO: ESTIMACIÓN MENSUAL VS VALOR REAL DE REFERENCIA"
ws_comp['A1'].fill = fill(NAVY); ws_comp['A1'].font = font(bold=True, color=WHITE, size=13); ws_comp['A1'].alignment = align('left', 'center')
ws_comp.row_dimensions[1].height = 28

ws_comp.merge_cells('A2:J2')
ws_comp['A2'] = "Valor de Referencia Promedio Real Histórico: $592,777.90 MXN / mes (22,349.7 L) | Obras México-Toluca y Lerma-Tres Marías"
ws_comp['A2'].fill = fill('1E293B'); ws_comp['A2'].font = font(italic=True, color='94A3B8', size=9); ws_comp['A2'].alignment = align('left', 'center')
ws_comp.row_dimensions[2].height = 20

# Header
headers_comp = [
    "MES", 
    "DIESEL EST. (L)", 
    "GASOLINA EST. (L)", 
    "TOTAL VOLUMEN (L)", 
    "PRESUPUESTO ESTIMADO ($)", 
    "VALOR REFERENCIA ($)", 
    "DIFERENCIA ($)", 
    "VARIACIÓN (%)", 
    "¿SUPERA REFERENCIA?", 
    "JUSTIFICACIÓN TÉCNICA Y OPERATIVA DE CAMPO"
]

for c_i, h in enumerate(headers_comp, 1):
    c = ws_comp.cell(row=4, column=c_i, value=h)
    c.fill = fill('334155'); c.font = font(bold=True, color=WHITE, size=9); c.alignment = align('center', 'center'); c.border = thin_border()
ws_comp.row_dimensions[4].height = 24

data_comp = [
    ("Enero", 13778.4, 2823.5, 16601.9, 439498.72, 592777.90, -153279.18, -25.86, "NO", "Frío extremo matutino y heladas en Toluca/La Marquesa retrasan el colado de asfalto caliente hasta media mañana. Arranque paulatino y afinación anual de maquinaria."),
    ("Febrero", 16678.4, 2949.5, 19627.9, 520811.22, 592777.90, -71966.68, -12.14, "NO", "Clima seco favorable, pero el mes cuenta con sólo 28 días naturales (20 días hábiles efectivos), lo que reduce el volumen total mensual respecto al promedio."),
    ("Marzo", 21950.7, 4004.9, 25955.6, 688384.86, 592777.90, 95606.96, 16.13, "SÍ", "Temporada óptima de estiaje: cero lluvias, altas temperaturas para emulsión y asfalto, frentes de fresado nocturno continuos y máximos tiros de pavimentadora."),
    ("Abril", 18120.7, 3287.5, 21408.2, 567831.31, 592777.90, -24946.59, -4.21, "NO", "Clima excelente, pero compensado por el paro de Semana Santa (3-4 días hábiles), aprovechado para mantenimiento preventivo de fresadora y pavimentadora."),
    ("Mayo", 25017.7, 4391.0, 29408.7, 780422.64, 592777.90, 187644.74, 31.66, "SÍ", "MES PICO ANUAL: Se aceleran los frentes constructivos a doble turno para cerrar metas antes del inicio de lluvias. Máxima demanda simultánea de maquinaria pesada y pipas."),
    ("Junio", 12286.5, 2724.4, 15010.9, 396847.42, 592777.90, -195930.48, -33.05, "NO", "Inicio de lluvias torrenciales en zona de montaña. Imposibilidad técnica de regar liga o tirar asfalto sobre base mojada. Equipos pesados entran a taller para mantenimiento.")
]

r_curr = 5
for idx, d in enumerate(data_comp):
    supera = d[8] == "SÍ"
    bg = 'FEF2F2' if supera else 'F0FDF4'
    tag_bg = 'DC2626' if supera else '16A34A'
    
    ws_comp.cell(row=r_curr, column=1, value=d[0]).font = font(bold=True, size=9)
    ws_comp.cell(row=r_curr, column=2, value=d[1]).number_format = '#,##0.0'
    ws_comp.cell(row=r_curr, column=3, value=d[2]).number_format = '#,##0.0'
    ws_comp.cell(row=r_curr, column=4, value=d[3]).number_format = '#,##0.0'
    ws_comp.cell(row=r_curr, column=5, value=d[4]).number_format = '"$"#,##0.00'
    ws_comp.cell(row=r_curr, column=6, value=d[5]).number_format = '"$"#,##0.00'
    
    c_dif = ws_comp.cell(row=r_curr, column=7, value=d[6])
    c_dif.number_format = '"+"#,##0.00;"-"#,##0.00;"$0.00"'
    c_dif.font = font(bold=True, color=RED_TEXT if supera else GREEN_TEXT, size=9)
    
    c_pct = ws_comp.cell(row=r_curr, column=8, value=d[7]/100.0)
    c_pct.number_format = '+0.0%;-0.0%;0.0%'
    c_pct.font = font(bold=True, color=RED_TEXT if supera else GREEN_TEXT, size=9)
    
    c_sup = ws_comp.cell(row=r_curr, column=9, value=d[8])
    c_sup.fill = fill(tag_bg); c_sup.font = font(bold=True, color=WHITE, size=9); c_sup.alignment = align('center', 'center')
    
    c_jus = ws_comp.cell(row=r_curr, column=10, value=d[9])
    c_jus.font = font(size=8); c_jus.alignment = align('left', 'center', wrap=True)
    
    for c in range(1, 11):
        cell = ws_comp.cell(row=r_curr, column=c)
        if c != 9:
            cell.fill = fill(bg)
        cell.border = thin_border()
        if c in [1]: cell.alignment = align('center', 'center')
        elif c in [2, 3, 4, 5, 6, 7, 8]: cell.alignment = align('right', 'center')
    
    ws_comp.row_dimensions[r_curr].height = 28
    r_curr += 1

# Total Semestre
ws_comp.cell(row=r_curr, column=1, value="TOTAL SEMESTRE").fill = fill(NAVY)
ws_comp.cell(row=r_curr, column=1).font = font(bold=True, color=WHITE, size=9); ws_comp.cell(row=r_curr, column=1).alignment = align('center', 'center')

ws_comp.cell(row=r_curr, column=2, value=107832.4).number_format = '#,##0.0'
ws_comp.cell(row=r_curr, column=3, value=20180.8).number_format = '#,##0.0'
ws_comp.cell(row=r_curr, column=4, value=128013.2).number_format = '#,##0.0'
ws_comp.cell(row=r_curr, column=5, value=3393796.17).number_format = '"$"#,##0.00'
ws_comp.cell(row=r_curr, column=6, value=592777.90 * 6).number_format = '"$"#,##0.00'

c_tdif = ws_comp.cell(row=r_curr, column=7, value=3393796.17 - (592777.90 * 6))
c_tdif.number_format = '"+"#,##0.00;"-"#,##0.00;"$0.00"'
c_tdif.font = font(bold=True, color='FDE047', size=9)

c_tpct = ws_comp.cell(row=r_curr, column=8, value=(3393796.17 - (592777.90 * 6)) / (592777.90 * 6))
c_tpct.number_format = '+0.0%;-0.0%;0.0%'
c_tpct.font = font(bold=True, color='FDE047', size=9)

ws_comp.cell(row=r_curr, column=9, value="2 de 6 Meses").fill = fill(NAVY)
ws_comp.cell(row=r_curr, column=9).font = font(bold=True, color=WHITE, size=8); ws_comp.cell(row=r_curr, column=9).alignment = align('center', 'center')

ws_comp.cell(row=r_curr, column=10, value="Presupuesto Semestral Dinámico es -4.6% inferior al promedio lineal por el impacto de lluvias de junio y frío de enero.").fill = fill(NAVY)
ws_comp.cell(row=r_curr, column=10).font = font(italic=True, color='94A3B8', size=8); ws_comp.cell(row=r_curr, column=10).alignment = align('left', 'center')

for c in range(1, 11):
    cell = ws_comp.cell(row=r_curr, column=c)
    if c not in [1, 9, 10]:
        cell.fill = fill(NAVY); cell.font = font(bold=True, color=WHITE, size=9)
        cell.alignment = align('right', 'center')
    cell.border = thin_border()
ws_comp.row_dimensions[r_curr].height = 24

# Widths
for col in ws_comp.columns:
    col_letter = get_column_letter(col[0].column)
    if col_letter == 'J':
        ws_comp.column_dimensions[col_letter].width = 45
    else:
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str: val_str = max(val_str.split('\n'), key=len)
            max_len = max(max_len, len(val_str))
        ws_comp.column_dimensions[col_letter].width = max(max_len + 3, 13)

wb.save(filename)
print("Saved comparison table in Excel successfully!")
