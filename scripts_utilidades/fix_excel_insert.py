# -*- coding: utf-8 -*-
import openpyxl
from datetime import datetime
wb = openpyxl.load_workbook('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano ultimo(1).xlsx')
ws = wb['Captura_Facturas_Mobil']
empty_row = ws.max_row + 1
for r in range(5, ws.max_row + 1):
    if not ws.cell(row=r, column=1).value:
        empty_row = r
        break
# Insert A-10270
ws.cell(row=empty_row, column=1).value = 'A-10270'
ws.cell(row=empty_row, column=2).value = datetime(2026, 6, 27)
ws.cell(row=empty_row, column=3).value = 39
ws.cell(row=empty_row, column=4).value = 'México-Toluca'
ws.cell(row=empty_row, column=5).value = 290
ws.cell(row=empty_row, column=6).value = 27.00
ws.cell(row=empty_row, column=7).value = 7830.05
ws.cell(row=empty_row, column=14).value = 26
# Insert A-10271
empty_row += 1
ws.cell(row=empty_row, column=1).value = 'A-10271'
ws.cell(row=empty_row, column=2).value = datetime(2026, 6, 27)
ws.cell(row=empty_row, column=3).value = 39
ws.cell(row=empty_row, column=4).value = 'Planta Asflato Huixquilucan'
ws.cell(row=empty_row, column=5).value = 600
ws.cell(row=empty_row, column=6).value = 27.00
ws.cell(row=empty_row, column=7).value = 16200.1
ws.cell(row=empty_row, column=14).value = 26
wb.save('c:/Users/JOSE/Desktop/Proyecto fenix/GC-COMB-3.1_Excel_Maestro_Control_Combustible_Grupo_Trujano_Actualizado.xlsx')
print('Inserted missing invoices successfully!')

