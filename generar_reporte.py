#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generar_Reporte.py  —  Reporte Semanal de Combustible y Acarreos
Uso: python generar_reporte.py --semana 26
"""
import argparse, sqlite3, os, sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (PatternFill, Font, Alignment, Border, Side,
                               numbers as xlnumbers)
from openpyxl.utils import get_column_letter

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fenix.db")
LOGO    = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logo_empresa.png")
PRECIO_LEVET = 22.89  # MXN/L precio promedio Levet

# ─── ESTILOS ──────────────────────────────────────────────────────────────────
HDR_FILL   = PatternFill("solid", fgColor="1a3c5e")
HDR_FONT   = Font(bold=True, color="FFFFFF", size=10)
HDR_ALN    = Alignment(horizontal="center", vertical="center", wrap_text=True)
SUB_FILL   = PatternFill("solid", fgColor="2d5a8e")
SUB_FONT   = Font(bold=True, color="E2E8F0", size=9)
ROW_FILL1  = PatternFill("solid", fgColor="F8FAFC")
ROW_FILL2  = PatternFill("solid", fgColor="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1a3c5e")
RED_FONT   = Font(bold=True, color="DC2626")
GRN_FONT   = Font(bold=True, color="059669")
AMB_FONT   = Font(bold=True, color="D97706")
THIN       = Side(style="thin", color="CBD5E1")
BORDER     = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM_MXN    = '#,##0.00'
NUM_L      = '#,##0.0'
NUM_INT    = '#,##0'

def get_db():
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    return db

def style_header(ws, row, cols, fill=HDR_FILL, font=HDR_FONT):
    for c in range(1, cols+1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill; cell.font = font; cell.alignment = HDR_ALN
        cell.border = BORDER

def style_row(ws, row, cols, idx=0):
    fill = ROW_FILL1 if idx%2==0 else ROW_FILL2
    for c in range(1, cols+1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill; cell.border = BORDER
        cell.alignment = Alignment(vertical="center")

def title_row(ws, row, title, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = TITLE_FONT; cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 28

def add_filter(ws, row, ncols):
    ws.auto_filter.ref = f"A{row}:{get_column_letter(ncols)}{ws.max_row}"

# ─── PESTÑA 1: RESUMEN ────────────────────────────────────────────────────────
def build_resumen(wb, semana):
    ws = wb.create_sheet("RESUMEN")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 20

    ws.row_dimensions[1].height = 40
    ws.merge_cells("A1:B1")
    ws['A1'] = f"REPORTE SEMANAL — SEMANA {semana}"
    ws['A1'].font = Font(bold=True, size=16, color="1a3c5e")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A2:B2")
    ws['A2'] = f"J.D.J. Equipo y Construcciones S.A. de C.V."
    ws['A2'].font = Font(italic=True, size=11, color="64748B")
    ws['A2'].alignment = Alignment(horizontal="center")
    ws.merge_cells("A3:B3")
    ws['A3'] = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws['A3'].font = Font(size=9, color="94A3B8"); ws['A3'].alignment = Alignment(horizontal="center")

    db = get_db(); cur = db.cursor()

    # Gasolina
    cur.execute("SELECT ROUND(SUM(monto_autorizado),2) FROM fenix_gas_presupuestos WHERE semana=?", (semana,))
    aut = cur.fetchone()[0] or 0
    cur.execute("SELECT ROUND(SUM(litros),2), ROUND(SUM(importe),2) FROM fenix_gas_tickets_reales WHERE semana=?", (semana,))
    gr = cur.fetchone(); lts_real = gr[0] or 0; imp_real = gr[1] or 0
    cur.execute("SELECT ROUND(SUM(importe),2) FROM fenix_gas_tickets_reales WHERE semana=? AND placa='NO REPORTADA'", (semana,))
    sin_id = cur.fetchone()[0] or 0
    lts_aut = round(aut / PRECIO_LEVET, 1) if aut else 0
    dif = round(imp_real - aut, 2)

    # Diesel
    cur.execute("SELECT ROUND(SUM(litros),2) FROM fenix_movimientos_combustible WHERE semana=? AND tipo_combustible='Diesel' AND tipo_movimiento='CONSUMO'", (semana,))
    lts_diesel = cur.fetchone()[0] or 0

    # Acarreos
    cur.execute("SELECT COUNT(*), ROUND(SUM(total),2) FROM fenix_viajes_acarreo WHERE semana=?", (semana,))
    ar = cur.fetchone(); viajes = ar[0] or 0; total_ac = ar[1] or 0

    db.close()

    data = [
        ("", ""),
        ("━━━  GASOLINA  ━━━", ""),
        ("Litros Autorizados (semana)", f"{lts_aut:,.1f} L"),
        ("Importe Autorizado", f"${aut:,.2f}"),
        ("Litros Reales (gasolinera)", f"{lts_real:,.1f} L"),
        ("Importe Real (gasolinera)", f"${imp_real:,.2f}"),
        ("Diferencia (Real − Autorizado)", f"${dif:,.2f}"),
        ("Tickets sin identificar", f"${sin_id:,.2f}"),
        ("", ""),
        ("━━━  DIÉSEL  ━━━", ""),
        ("Litros Consumidos", f"{lts_diesel:,.1f} L"),
        ("", ""),
        ("━━━  ACARREOS  ━━━", ""),
        ("Total Viajes", f"{viajes:,}"),
        ("Total a Pagar", f"${total_ac:,.2f}"),
    ]

    for i, (label, value) in enumerate(data, start=5):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
        if "━━" in label:
            ws.cell(row=i, column=1).font = Font(bold=True, color="1a3c5e", size=11)
        elif value.startswith("$"):
            val_num = float(value.replace("$","").replace(",",""))
            ws.cell(row=i, column=2).value = val_num
            ws.cell(row=i, column=2).number_format = NUM_MXN
            if "Diferencia" in label and val_num > 0:
                ws.cell(row=i, column=2).font = RED_FONT
            elif "Diferencia" in label:
                ws.cell(row=i, column=2).font = GRN_FONT
        elif "L" in value and "L " not in value:
            ws.cell(row=i, column=2).font = Font(bold=True, color="059669")

# ─── PESTAÑA 2: GASOLINA CONCILIACIÓN ────────────────────────────────────────
def build_gasolina_conciliacion(wb, semana):
    ws = wb.create_sheet("GAS - CONCILIACIÓN")
    ws.sheet_view.showGridLines = False
    widths = [14,28,22,30,6,14,14,14,14,14,16]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width = w

    title_row(ws, 1, f"CONCILIACIÓN DE GASOLINA — SEMANA {semana} (Lunes a Jueves)", len(widths))
    ws.row_dimensions[2].height = 5

    # ── A: Por Placa ──
    ws.cell(row=3, column=1, value="A. AUTORIZADO vs CONSUMO REAL POR UNIDAD")
    ws.cell(row=3, column=1).font = Font(bold=True, size=11, color="1a3c5e")

    headers_a = ["Semana","Responsable","Unidad","Obra","Placa",
                  "Lts Autorizados","Lts Gasolinera","$ Autorizado","$ Real Gasolinera","Diferencia $","Diferencia %"]
    for c,h in enumerate(headers_a,1):
        ws.cell(row=4, column=c, value=h)
    style_header(ws, 4, len(headers_a))

    db = get_db(); cur = db.cursor()
    cur.execute("""
        SELECT p.semana, p.responsable, p.unidad, p.obra, p.placa,
               ROUND(p.monto_autorizado,2) as aut
        FROM fenix_gas_presupuestos p
        WHERE p.semana=? AND p.placa IS NOT NULL AND p.placa!=''
        ORDER BY p.obra, p.responsable
    """, (semana,))
    presupuestos = cur.fetchall()

    cur.execute("""
        SELECT placa, ROUND(SUM(litros),3) as lts, ROUND(SUM(importe),2) as imp
        FROM fenix_gas_tickets_reales WHERE semana=? AND placa!='NO REPORTADA' AND placa IS NOT NULL
        GROUP BY placa
    """, (semana,))
    tick_map = {r['placa']: r for r in cur.fetchall()}

    row = 5
    total_aut = total_real = 0
    for idx,p in enumerate(presupuestos):
        t = tick_map.get(p['placa'], None)
        lts_aut_u = round(p['aut']/PRECIO_LEVET, 1)
        lts_real_u = t['lts'] if t else 0
        imp_real_u = t['imp'] if t else 0
        dif = round(imp_real_u - p['aut'], 2)
        pct = round((dif/p['aut'])*100, 1) if p['aut'] else 0
        style_row(ws, row, len(headers_a), idx)
        vals = [p['semana'], p['responsable'], p['unidad'], p['obra'], p['placa'],
                lts_aut_u, lts_real_u, p['aut'], imp_real_u, dif, pct/100]
        for c,v in enumerate(vals,1):
            cell = ws.cell(row=row, column=c, value=v)
            if c in [6,7]: cell.number_format = NUM_L
            if c in [8,9]: cell.number_format = NUM_MXN
            if c==10:
                cell.number_format = NUM_MXN
                cell.font = RED_FONT if dif>50 else (GRN_FONT if dif<-50 else Font())
            if c==11: cell.number_format = '0.0%'
        total_aut += p['aut']; total_real += imp_real_u
        row += 1

    # Totales sección A
    ws.cell(row=row, column=7, value="TOTAL:")
    ws.cell(row=row, column=7).font = Font(bold=True)
    ws.cell(row=row, column=8, value=total_aut); ws.cell(row=row, column=8).number_format = NUM_MXN; ws.cell(row=row, column=8).font = Font(bold=True)
    ws.cell(row=row, column=9, value=total_real); ws.cell(row=row, column=9).number_format = NUM_MXN; ws.cell(row=row, column=9).font = Font(bold=True)
    dif_tot = total_real - total_aut
    ws.cell(row=row, column=10, value=dif_tot); ws.cell(row=row, column=10).number_format = NUM_MXN
    ws.cell(row=row, column=10).font = RED_FONT if dif_tot>0 else GRN_FONT
    row += 2

    # ── B: Por Obra ──
    ws.cell(row=row, column=1, value="B. RESUMEN POR OBRA/CENTRO DE TRABAJO")
    ws.cell(row=row, column=1).font = Font(bold=True, size=11, color="1a3c5e")
    row += 1

    cur.execute("""
        SELECT COALESCE(obra,'Sin Asignar') as obra,
               COUNT(*) as unidades, ROUND(SUM(monto_autorizado),2) as aut
        FROM fenix_gas_presupuestos WHERE semana=? GROUP BY obra ORDER BY aut DESC
    """, (semana,))
    obras = cur.fetchall()

    cur.execute("""
        SELECT p.obra, ROUND(SUM(t.litros),2) as lts, ROUND(SUM(t.importe),2) as imp
        FROM fenix_gas_tickets_reales t
        JOIN fenix_gas_presupuestos p ON p.placa=t.placa AND p.semana=t.semana
        WHERE t.semana=? AND t.placa!='NO REPORTADA'
        GROUP BY p.obra
    """, (semana,))
    obra_tick = {r['obra']: r for r in cur.fetchall()}

    headers_b = ["Obra / Centro de Trabajo","Unidades","$ Autorizado","Lts Reales","$ Real","Diferencia","Desviación %"]
    for c,h in enumerate(headers_b,1):
        ws.cell(row=row, column=c, value=h)
    style_header(ws, row, len(headers_b)); row+=1

    for idx,o in enumerate(obras):
        t = obra_tick.get(o['obra'],None)
        imp_r = t['imp'] if t else 0; lts_r = t['lts'] if t else 0
        dif = round(imp_r - o['aut'],2)
        style_row(ws, row, len(headers_b), idx)
        vals = [o['obra'], o['unidades'], o['aut'], lts_r, imp_r, dif, (dif/o['aut']) if o['aut'] else 0]
        for c,v in enumerate(vals,1):
            cell = ws.cell(row=row, column=c, value=v)
            if c in [3,5]: cell.number_format = NUM_MXN
            if c==4: cell.number_format = NUM_L
            if c==6:
                cell.number_format = NUM_MXN
                cell.font = RED_FONT if dif>0 else GRN_FONT
            if c==7: cell.number_format = '0.0%'
        row += 1
    row += 1

    # ── C: Tickets sin Identificar ──
    ws.cell(row=row, column=1, value="C. TICKETS SIN IDENTIFICAR (Placa NO REPORTADA)")
    ws.cell(row=row, column=1).font = Font(bold=True, size=11, color="B91C1C")
    row += 1

    cur.execute("""
        SELECT ticket, gasolinera, fecha, ROUND(litros,3) as litros,
               ROUND(precio_litro,2) as precio, ROUND(importe,2) as importe
        FROM fenix_gas_tickets_reales WHERE semana=? AND placa='NO REPORTADA'
        ORDER BY ticket
    """, (semana,))
    sin_id_rows = cur.fetchall()
    db.close()

    headers_c = ["# Ticket","Gasolinera","Fecha","Litros","$/Litro","Importe","Nota"]
    for c,h in enumerate(headers_c,1):
        ws.cell(row=row, column=c, value=h)
    style_header(ws, row, len(headers_c), fill=PatternFill("solid", fgColor="991B1B")); row+=1

    if sin_id_rows:
        total_sin_id = 0
        for idx,r_ in enumerate(sin_id_rows):
            style_row(ws, row, len(headers_c), idx)
            vals = [r_['ticket'], r_['gasolinera'], r_['fecha'], r_['litros'],
                    r_['precio'], r_['importe'], "⚠️ PENDIENTE DE IDENTIFICAR"]
            for c,v in enumerate(vals,1):
                cell = ws.cell(row=row, column=c, value=v)
                if c in [4]: cell.number_format = NUM_L
                if c in [5,6]: cell.number_format = NUM_MXN
                if c==6: cell.font = RED_FONT
            total_sin_id += r_['importe']
            row += 1
        ws.cell(row=row, column=5, value="TOTAL SIN ID:"); ws.cell(row=row, column=5).font = Font(bold=True, color="B91C1C")
        ws.cell(row=row, column=6, value=total_sin_id); ws.cell(row=row, column=6).number_format = NUM_MXN; ws.cell(row=row, column=6).font = RED_FONT
    else:
        ws.cell(row=row, column=1, value="✅ Sin tickets sin identificar para esta semana")

    ws.freeze_panes = "A5"

# ─── PESTAÑA 3: GASOLINA MOVIMIENTOS ─────────────────────────────────────────
def build_gasolina_movimientos(wb, semana):
    ws = wb.create_sheet("GAS - MOVIMIENTOS")
    ws.sheet_view.showGridLines = False
    widths = [14,14,12,14,28,28,30,12,12,14]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width = w

    title_row(ws, 1, f"MOVIMIENTOS DETALLADOS DE GASOLINA — SEMANA {semana}", len(widths))

    headers = ["# Ticket","Folio","Fecha","Gasolinera","Placa","Conductor","Obra / Destino",
               "Litros","$/Litro","Importe"]
    for c,h in enumerate(headers,1): ws.cell(row=2, column=c, value=h)
    style_header(ws, 2, len(headers))

    db = get_db(); cur = db.cursor()
    cur.execute("""
        SELECT ticket, folio, fecha, gasolinera, COALESCE(placa,'SIN PLACA') as placa,
               COALESCE(conductor,'') as conductor, COALESCE(obra,'') as obra,
               ROUND(litros,3) as litros, ROUND(precio_litro,2) as precio,
               ROUND(importe,2) as importe
        FROM fenix_gas_tickets_reales WHERE semana=?
        ORDER BY gasolinera, fecha, ticket
    """, (semana,))
    rows = cur.fetchall(); db.close()

    for idx,r in enumerate(rows):
        row = idx+3
        style_row(ws, row, len(headers), idx)
        vals = [r['ticket'],r['folio'],r['fecha'],r['gasolinera'],r['placa'],
                r['conductor'],r['obra'],r['litros'],r['precio'],r['importe']]
        for c,v in enumerate(vals,1):
            cell = ws.cell(row=row, column=c, value=v)
            if c==8: cell.number_format = NUM_L
            if c in [9,10]: cell.number_format = NUM_MXN
            if r['placa']=='SIN PLACA' or r['placa']=='NO REPORTADA':
                cell.font = AMB_FONT

    add_filter(ws, 2, len(headers))
    ws.freeze_panes = "A3"

# ─── PESTAÑA 4: DIESEL ────────────────────────────────────────────────────────
def build_diesel(wb, semana):
    ws = wb.create_sheet("DIESEL")
    ws.sheet_view.showGridLines = False
    widths = [12,14,25,28,12,14,20]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width = w

    title_row(ws, 1, f"CONSUMO DE DIESEL — SEMANA {semana}", len(widths))

    headers = ["Semana","Fecha","Maquinaria/Equipo","Obra","Litros","Folio","Observaciones"]
    for c,h in enumerate(headers,1): ws.cell(row=2, column=c, value=h)
    style_header(ws, 2, len(headers))

    db = get_db(); cur = db.cursor()
    cur.execute("""
        SELECT m.semana, m.fecha,
               COALESCE(e.descripcion, m.destino_nombre, 'Sin equipo') as equipo,
               COALESCE(o.nombre, 'Sin asignar') as obra,
               ROUND(m.litros,2) as litros, m.folio_vale, m.observaciones
        FROM fenix_movimientos_combustible m
        LEFT JOIN fenix_equipos e ON e.id = m.equipo_id
        LEFT JOIN fenix_obras o ON o.id = m.obra_id
        WHERE m.semana=? AND m.tipo_combustible='Diesel' AND m.tipo_movimiento='CONSUMO'
        ORDER BY m.fecha, equipo
    """, (semana,))
    rows = cur.fetchall(); db.close()

    for idx,r in enumerate(rows):
        row = idx+3
        style_row(ws, row, len(headers), idx)
        vals = [r['semana'],r['fecha'],r['equipo'],r['obra'],r['litros'],r['folio_vale'],r['observaciones']]
        for c,v in enumerate(vals,1):
            cell = ws.cell(row=row, column=c, value=v)
            if c==5: cell.number_format = NUM_L

    add_filter(ws, 2, len(headers))
    ws.freeze_panes = "A3"

# ─── PESTAÑA 5: ACARREOS ─────────────────────────────────────────────────────
def build_acarreos(wb, semana):
    ws = wb.create_sheet("ACARREOS")
    ws.sheet_view.showGridLines = False
    widths = [12,14,14,28,22,14,14,14]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width = w

    title_row(ws, 1, f"ACARREOS — SEMANA {semana}", len(widths))

    headers = ["Semana","Fecha","Placa","Obra","Sindicato","Categoría","Material","Subtotal $"]
    for c,h in enumerate(headers,1): ws.cell(row=2, column=c, value=h)
    style_header(ws, 2, len(headers))

    db = get_db(); cur = db.cursor()
    cur.execute("""
        SELECT v.semana, v.fecha, v.placa,
               COALESCE(o.nombre,'—') as obra,
               COALESCE(s.nombre,'—') as sindicato,
               COALESCE(v.categoria,'—') as categoria,
               COALESCE(v.material,'—') as material,
               ROUND(v.subtotal,2) as subtotal
        FROM fenix_viajes_acarreo v
        LEFT JOIN fenix_obras o ON o.id=v.obra_id
        LEFT JOIN fenix_sindicatos s ON s.id=v.sindicato_id
        WHERE v.semana=? ORDER BY v.fecha, v.placa
    """, (semana,))
    rows = cur.fetchall(); db.close()

    for idx,r in enumerate(rows):
        row = idx+3
        style_row(ws, row, len(headers), idx)
        vals = [r['semana'],r['fecha'],r['placa'],r['obra'],r['sindicato'],
                r['categoria'],r['material'],r['subtotal']]
        for c,v in enumerate(vals,1):
            cell = ws.cell(row=row, column=c, value=v)
            if c==8: cell.number_format = NUM_MXN; cell.font = AMB_FONT

    add_filter(ws, 2, len(headers))
    ws.freeze_panes = "A3"

# ─── MAIN ─────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--semana", type=int, required=True)
    args = parser.parse_args()
    semana = args.semana

    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    print(f"Generando reporte para Semana {semana}...")
    build_resumen(wb, semana)
    build_gasolina_conciliacion(wb, semana)
    build_gasolina_movimientos(wb, semana)
    build_diesel(wb, semana)
    build_acarreos(wb, semana)

    fname = f"Reporte_Fenix_Sem{semana}_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    fpath = os.path.join(os.path.dirname(os.path.abspath(__file__)), fname)
    wb.save(fpath)
    print(f"[OK] Reporte guardado: {fpath}")
    return fpath

if __name__ == "__main__":
    main()
